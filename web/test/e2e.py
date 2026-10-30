"""网页版从头到尾点一遍（真浏览器 + 假谷歌服务器），出的表跟电脑版一格一格对。

跑：python web/test/e2e.py [截图放哪]
要：playwright（chromium）、openpyxl、Pillow、xlwt；可选 pillow-heif（测 HEIC）
"""
import functools
import http.server
import io
import json
import os
import shutil
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
import zipfile
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, 'tests'))
from bangong import engine, tables  # noqa: E402
import mock_server  # noqa: E402
from test_kb import company_book  # noqa: E402
from test_formats import to_xls  # noqa: E402

K1, K2 = 'AIzaKEY1_mock', 'AIzaKEY2_mock'
SHOTS = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix='luruseng-shots-')
os.makedirs(SHOTS, exist_ok=True)
WORK = tempfile.mkdtemp(prefix='luruseng-e2e-')
CORS = {'access-control-allow-origin': '*', 'access-control-allow-methods': 'POST, OPTIONS',
        'access-control-allow-headers': 'content-type, x-goog-api-key'}
passed = []


def step(fn):
    @functools.wraps(fn)
    def run(*a, **kw):
        t = time.time()
        fn(*a, **kw)
        passed.append(fn.__name__)
        print(f'  ✓ {fn.__name__}  ({time.time() - t:.1f}s)', flush=True)
    return run


def make_images(folder, n, start=0):
    os.makedirs(folder, exist_ok=True)
    out = []
    for i in range(start, start + n):
        p = os.path.join(folder, f'img_{i:02d}.jpg')
        Image.new('RGB', (800 + i, 600), (i * 7 % 255, 100, 150)).save(p)
        os.utime(p, (1_700_000_000 + i, 1_700_000_000 + i))
        out.append(p)
    return out


def serve(root):
    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *a):
            pass
    Quiet.extensions_map = {**http.server.SimpleHTTPRequestHandler.extensions_map, '.js': 'text/javascript',
                            '.mjs': 'text/javascript', '.webmanifest': 'application/manifest+json'}
    srv = http.server.ThreadingHTTPServer(('127.0.0.1', 0), functools.partial(Quiet, directory=root))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def python_book(images, spec_name=None):
    """电脑版拿同样的照片、同样的假服务器出一本表：拿来跟网页版的一格一格对。"""
    d = tempfile.mkdtemp(prefix='luruseng-py-')
    for p in images:
        shutil.copy2(p, d)
    kw = {}
    if spec_name:
        cat = os.path.join(d, '_cat')
        kw['spec'] = tables.TableSpec(spec_name, os.path.join(cat, 'db.sqlite'), os.path.join(cat, 'feed'))
    res = engine.Engine(d, [K1], log=lambda m: None, **kw).run()
    assert res.ok == len(images), res
    return res.xlsx


def cells(path):
    ws = load_workbook(path).worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in row:
            fill = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type == 'solid' else None
            font = (c.font.name, float(c.font.sz or 11), bool(c.font.b), c.font.color.rgb if c.font.color is not None and c.font.color.type == 'rgb' else None)
            v = c.value
            out[c.coordinate] = (v, fill, font, c.alignment.horizontal, bool(c.alignment.wrap_text), c.number_format if v is not None else None,
                                 bool(c.border.left.style))
    widths = {}
    for k, v in ws.column_dimensions.items():            # <col min max> 一段算好几列
        for i in range(v.min or column_index_from_string(k), (v.max or column_index_from_string(k)) + 1):
            if v.width:
                widths[get_column_letter(i)] = round(v.width, 2)
    meta = dict(title=ws.title, merged=sorted(str(m) for m in ws.merged_cells.ranges), freeze=ws.freeze_panes,
                filter=ws.auto_filter.ref, orient=ws.page_setup.orientation, fitw=ws.page_setup.fitToWidth,
                fith=ws.page_setup.fitToHeight, fit=ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None,
                titles=ws.print_title_rows, height1=ws.row_dimensions[1].height)
    return out, widths, meta


def same_book(js_path, py_path, skip=()):
    a, wa, ma = cells(js_path)
    b, wb, mb = cells(py_path)
    diffs = []
    for k in sorted(set(a) | set(b)):
        x, y = a.get(k), b.get(k)
        if x is None or y is None:            # 一边没有这个格子：只要另一边是空白就算一样
            z = x or y
            if z[0] is not None or z[1] is not None:
                diffs.append((k, x, y))
            continue
        if x[0] != y[0] and isinstance(x[0], str) and isinstance(y[0], str) and any(s in x[0] for s in skip):
            continue
        if x != y:
            diffs.append((k, x, y))
    for k in set(wa) | set(wb):
        if abs(wa.get(k, 0) - wb.get(k, 0)) > 0.01:
            diffs.append(('width ' + k, wa.get(k), wb.get(k)))
    for k in ma:
        if ma[k] != mb[k] and not (k == 'titles' and str(ma[k]).replace('$', '') == str(mb[k]).replace('$', '')):
            diffs.append((k, ma[k], mb[k]))
    assert not diffs, '网页版和电脑版的表不一样：\n' + '\n'.join(map(str, diffs[:30]))


def main():
    site = serve(ROOT)
    base = f'http://127.0.0.1:{site.server_address[1]}/web/'
    mock = mock_server.start()
    mock_url = f'http://127.0.0.1:{mock.server_address[1]}'
    engine.API_BASE = mock_url + '/v1beta'
    engine.RETRY_WAITS = [0.01, 0.01, 0.01]
    mock_server.STATE.update(delay=0)
    errors = []

    def gemini(route, request):
        if request.method == 'OPTIONS':
            route.fulfill(status=204, headers=CORS)
            return
        u = request.url.split('googleapis.com', 1)[1]
        req = urllib.request.Request(mock_url + u, data=request.post_data_buffer, method='POST',
                                     headers={'Content-Type': 'application/json'})
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                status, body = r.status, r.read()
        except urllib.error.HTTPError as e:
            status, body = e.code, e.read()
        route.fulfill(status=status, headers={**CORS, 'content-type': 'application/json'}, body=body)

    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        ctx = browser.new_context(viewport={'width': 1280, 'height': 860}, accept_downloads=True, locale='zh-CN')
        ctx.grant_permissions(['clipboard-read', 'clipboard-write'], origin=base.rsplit('/web/', 1)[0])
        page = ctx.new_page()
        page.on('pageerror', lambda e: errors.append(f'pageerror: {e}'))
        page.on('console', lambda m: errors.append(f'console.{m.type}: {m.text}')
                if m.type == 'error' and 'Failed to load resource' not in m.text else None)
        def offline(pg):                     # 除了（假的）谷歌，别的外网一概不许碰：网页版不靠 CDN
            pg.route('https://**', lambda route, req: errors.append(f'外网请求：{req.url}') or route.abort())
            pg.route('https://generativelanguage.googleapis.com/**', gemini)
        offline(page)

        def logs():
            return page.eval_on_selector_all('#log p', 'ps => ps.map(p => p.textContent)')

        def wait_log(text, timeout=60000):
            page.wait_for_function('t => [...document.querySelectorAll("#log p")].some(p => p.textContent.includes(t))',
                                   arg=text, timeout=timeout)

        def last_logs(n=8):
            return '\n'.join(logs()[-n:])

        def add_photos(paths):
            page.set_input_files('#files', paths)

        def run_and_wait(paths, until='成功', timeout=60000):
            before = len(logs())
            add_photos(paths)
            page.click('#start')
            page.wait_for_function('([t, n]) => [...document.querySelectorAll("#log p")].slice(n).some(p => p.textContent.includes(t))',
                                   arg=[until, before], timeout=timeout)
            page.wait_for_function('() => !document.querySelector("#start").disabled', timeout=timeout)
            return logs()[before:]

        def download(selector):
            with page.expect_download() as d:
                page.click(selector)
            path = os.path.join(WORK, d.value.suggested_filename)
            d.value.save_as(path)
            return path

        def close_viewer():
            if page.is_visible('#dlg-view'):
                page.click('#vclose')

        @step
        def opens_clean():
            page.goto(base)
            page.wait_for_function('() => document.querySelectorAll("#log p").length >= 2')
            assert page.title().startswith('录入僧网页版')
            assert page.inner_text('#status').replace('\n', '') == '待识别 0 张', page.inner_text('#status')
            text = '\n'.join(logs())
            assert '第一次用' in text and '数据库还是空的' in text, text
            assert page.inner_text('#cathint') == '内置 · 专属优化'
            page.screenshot(path=os.path.join(SHOTS, '1-打开.png'))

        @step
        def no_key_no_start():
            add_photos(make_images(os.path.join(WORK, 'none'), 1, start=90))
            page.click('#start')
            wait_log('✗ 先填钥匙')
            page.click('#clear')
            assert page.inner_text('#status').replace('\n', '') == '待识别 0 张'

        imgs = make_images(os.path.join(WORK, 'a'), 12)

        @step
        def twelve_photos_match_desktop_book():
            page.fill('#key1', K1)
            out = run_and_wait(imgs, '成功 12 张')
            text = '\n'.join(out)
            assert '按拍摄顺序编号 1–12' in text and '每张认两遍对答案' in text, text
            assert sum(1 for m in out if m.startswith('✓ ')) >= 12, text
            page.wait_for_selector('#dlg-view[open]')            # 干完自动打开核对
            assert '第 1 张 / 共 12 张' in page.inner_text('#vinfo')
            panel = page.inner_text('#vpanel')
            assert 'No. 1' in panel and '800' in panel and 'VM-07' in panel, panel
            page.wait_for_function('() => { const i = document.querySelector("#vphoto img"); return i && i.naturalWidth > 0; }')
            page.screenshot(path=os.path.join(SHOTS, '2-核对.png'))
            page.keyboard.press('ArrowRight')                                    # 一打开就能用 ←→ 翻页
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 2 张")')
            page.keyboard.press('ArrowLeft')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 1 张")')
            js = download('#vxlsx')
            assert os.path.basename(js).startswith('识别结果_') and js.endswith('.xlsx')
            same_book(js, python_book(imgs))
            assert len(mock_server.STATE['log']) >= 12 * 2 * 2      # 网页 24 次 + 电脑版 24 次

        @step
        def viewer_pages_rotates_copies_zips():
            page.keyboard.press('ArrowRight')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 2 张")')
            assert '801' in page.inner_text('#vpanel')
            w0 = page.eval_on_selector('#vphoto img', 'i => [i.naturalWidth, i.naturalHeight]')
            page.click('#vcw')
            page.wait_for_function('w => { const i = document.querySelector("#vphoto img"); return i && i.naturalWidth === w[1] && i.naturalHeight === w[0]; }', arg=w0)
            page.fill('#vnum', '12')
            page.press('#vnum', 'Enter')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 12 张")')
            page.fill('#vnum', '40')
            page.click('#vgo')
            assert '没有第 40 张（这批是 1–12）' in page.inner_text('#vinfo')
            page.click('#vcopy')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("复制好了")')
            clip = page.evaluate('navigator.clipboard.readText()').split('\r\n')
            assert len(clip) == 12 and clip[0].split('\t')[:6] == ['2026-10-28', 'VM-07', 'TD-524613-5', '82664580', '800', 'T4'], clip[0]
            z = download('#vzip')
            names = sorted(zipfile.ZipFile(z).namelist())
            assert names[:12] == [f'{i:02d}.jpg' for i in range(1, 13)] and names[12].endswith('.xlsx'), names
            im = Image.open(io.BytesIO(zipfile.ZipFile(z).read('01.jpg'))).convert('RGB')
            reds = sum(1 for x in range(im.width) for y in range(0, im.height, 4)
                       if (lambda p: p[0] > 170 and p[1] < 60 and p[2] < 60)(im.getpixel((x, y))))
            assert im.size == (800, 600) and reds > 100, (im.size, reds)        # 印了红编号，没被放大
            close_viewer()

        @step
        def same_photos_again_are_skipped():
            out = run_and_wait(imgs, '收工')
            text = '\n'.join(out)
            assert text.count('以前干过') == 12 and '第 1 张' in text and '列表里没有新照片，收工。' in text, text

        @step
        def result_button_downloads_latest():
            path = download('#result')
            assert load_workbook(path).worksheets[0].max_row == 15

        @step
        def feed_database_xlsx_xls_csv_html():
            book = company_book(os.path.join(WORK, 'TOOL BROKEN RECORD.xlsx'))
            xls = to_xls(book, os.path.join(WORK, '老表.xls'))
            csv_ = os.path.join(WORK, '刀单.csv')
            with open(csv_, 'wb') as f:
                f.write('DATE,MACHINE,PART NUMBER,MO,OPR NO.,T NO.,DESCRIPTION,BTMP,TOTAL SCREENED,PASSED,FAILED,SHIFT,SETTER\n'
                        '28/10,VM-31,TD-1,82660000,150,T4,EM2.5,1,1,1,0,D,700/Ah Seng\n'.encode('gbk'))
            html = os.path.join(WORK, '系统导出.xls')
            with open(html, 'w', encoding='utf-8') as f:
                f.write('<html><body><table><tr><td>DATE</td><td>MACHINE</td><td>PART NUMBER</td><td>MO</td><td>T NO.</td>'
                        '<td>DESCRIPTION</td><td>SETTER</td></tr><tr><td>29/10</td><td>VM-32</td><td>TD-2</td><td>82660001</td>'
                        '<td>T5</td><td>EM3</td><td>701/Kumar</td></tr></table></body></html>')
            broken = os.path.join(WORK, '坏表.xls')
            with open(broken, 'wb') as f:
                f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'xx' * 50)
            pdf = os.path.join(WORK, 'scan.pdf')
            with open(pdf, 'wb') as f:
                f.write(b'%PDF-1.4')
            ours = [p for p in os.listdir(WORK) if p.startswith('识别结果_') and p.endswith('.xlsx')][0]
            before = len(logs())
            page.click('#feed')
            assert '一条都不吃' in page.inner_text('#feeddesc')
            page.set_input_files('#feedfiles', [os.path.join(WORK, ours), book, xls, csv_, html, broken, pdf])
            page.wait_for_function('n => [...document.querySelectorAll("#log p")].slice(n).some(p => p.textContent.startsWith("数据库："))', arg=before)
            out = '\n'.join(logs()[before:])
            assert f'✓ 吃进 {ours}：读到 12 行记录，新增 12 行' in out, out
            assert '✓ 吃进 TOOL BROKEN RECORD.xlsx：读到 8 行记录，新增 8 行' in out, out
            assert '✓ 吃进 老表.xls：读到 8 行记录，新增 0 行' in out, out
            assert '✓ 吃进 刀单.csv：读到 1 行记录，新增 1 行' in out, out
            assert '✓ 吃进 系统导出.xls：读到 1 行记录，新增 1 行' in out, out
            assert '✗ 坏表.xls：这本表坏了' in out and '✗ scan.pdf：只认' in out, out
            assert '数据库：22 条记录' in out, out

        @step
        def database_fixes_confusable_letters():
            mock_server.STATE['modes'] = {K1: 'confusable'}
            try:
                run_and_wait(make_images(os.path.join(WORK, 'c'), 1, start=20), '成功 1 张')
            finally:
                mock_server.STATE['modes'] = {}
            page.wait_for_selector('#dlg-view[open]')
            fixed = page.eval_on_selector_all('#vpanel dd.fixed', 'ds => ds.map(d => d.textContent)')
            assert 'VM-07' in fixed and 'TD-524613-5' in fixed, fixed
            assert 'MACHINE VM-O7→VM-07（按数据库改）' in page.inner_text('#vpanel')
            page.screenshot(path=os.path.join(SHOTS, '3-按数据库改.png'))
            close_viewer()

        @step
        def names_dialog():
            page.click('#names')
            page.fill('#namestxt', '639/Lee Ah Kow\n# 注释\n232/Tan Mei Ling')
            page.click('#namesok')
            wait_log('✓ 名单存好了：2 人')
            assert page.inner_text('#nhead') == '名单：2 人'

        cat_imgs = make_images(os.path.join(WORK, 'q'), 2, start=30)

        @step
        def own_category_matches_desktop_book():
            page.click('#newcat')
            page.fill('#newname', '刀单')
            page.click('#newok')
            assert page.inner_text('#newerr') == '已经有「刀单」了'
            page.fill('#newname', '质检单')
            page.click('#newok')
            wait_log('「质检单」还没放过表')
            assert page.input_value('#cat') == '质检单' and page.inner_text('#cathint') == '自建 · 还没放过表'
            assert page.is_hidden('#names')
            out = run_and_wait(cat_imgs, '成功 2 张')
            assert any('类别「质检单」（还没定列：认全部内容）' in m for m in out), out
            page.wait_for_selector('#dlg-view[open]')
            assert 'M/C No.' in page.inner_text('#vpanel')
            js = download('#vxlsx')
            assert js.endswith('_质检单.xlsx'), js
            same_book(js, python_book(cat_imgs, '质检单'))
            close_viewer()
            before = len(logs())
            page.click('#feed')
            assert page.input_value('#feedcat') == '质检单'
            page.set_input_files('#feedfiles', [js])
            page.wait_for_function('n => [...document.querySelectorAll("#log p")].slice(n).some(p => p.textContent.startsWith("「质检单」数据库："))', arg=before)
            out = '\n'.join(logs()[before:])
            assert '✓ 放进去' in out and '新增 2 行' in out and '新列：Date、M/C No.、Qty、Remark' in out, out
            assert page.inner_text('#cathint') == '自建 · 4 列 · 2 条记录'
            out = run_and_wait(make_images(os.path.join(WORK, 'q2'), 1, start=40), '成功 1 张')
            assert any('类别「质检单」（4 列）' in m for m in out), out
            close_viewer()
            page.select_option('#cat', '刀单')
            wait_log('数据库：')

        @step
        def bad_key_stops_and_keeps_photos():
            mock_server.STATE['modes'] = {K1: 'invalid'}
            try:
                out = run_and_wait(make_images(os.path.join(WORK, 'bad'), 2, start=50), '停工')
            finally:
                mock_server.STATE['modes'] = {}
            assert any('■ 停工：钥匙1 · HTTP 400：API key not valid' in m for m in out), out
            assert page.inner_text('#status').replace('\n', '') == '待识别 2 张'     # 照片还在列表里
            out = run_and_wait([], '成功 2 张')                                   # 钥匙好了再点开始：接着干
            assert any('成功 2 张' in m for m in out)
            close_viewer()

        @step
        def stop_button_marks_unfinished():
            mock_server.STATE['delay'] = 0.25
            try:
                add_photos(make_images(os.path.join(WORK, 's'), 12, start=60))
                page.click('#start')
                page.wait_for_function('() => /^5\\/12/.test(document.querySelector("#status").textContent)', timeout=30000)
                page.click('#stop')
                wait_log('■ 已停止，没干完的下次接着干', timeout=60000)
                page.wait_for_function('() => !document.querySelector("#start").disabled')
            finally:
                mock_server.STATE['delay'] = 0
            left = page.inner_text('#droptxt')
            assert left.startswith('7 张：img_65.jpg'), left                     # 没干完的还在列表里
            page.click('#check')
            page.wait_for_selector('#dlg-view[open]')
            page.fill('#vnum', '6')
            page.press('#vnum', 'Enter')
            page.wait_for_function('() => document.querySelector("#vpanel").textContent.includes("No. 6")')
            assert '未完成：已停止' in page.inner_text('#vpanel'), page.inner_text('#vpanel')
            close_viewer()
            page.click('#clear')

        @step
        def photos_of_all_kinds():
            d = os.path.join(WORK, 'kinds')
            os.makedirs(d, exist_ok=True)
            big = os.path.join(d, 'big.png')
            Image.new('RGBA', (4032, 3024), (0, 0, 0, 0)).save(big)
            rot = os.path.join(d, 'rot.jpg')
            im = Image.new('RGB', (900, 600), 'white')
            ex = im.getexif()
            ex[0x0112] = 6                                  # 横着拍的：看的时候要转 90°
            im.save(rot, exif=ex)
            paths = [big, rot]
            try:
                import pillow_heif
                pillow_heif.register_heif_opener()
                heic = os.path.join(d, 'iphone.heic')
                Image.new('RGB', (1000, 700), (200, 220, 240)).save(heic, format='HEIF')
                paths.append(heic)
            except ImportError:
                print('    （没装 pillow-heif，HEIC 不测）')
            out = run_and_wait(paths, f'成功 {len(paths)} 张', timeout=120000)
            page.wait_for_selector('#dlg-view[open]')
            js = download('#vxlsx')
            ws = load_workbook(js).worksheets[0]
            widths = sorted(ws.cell(row=r, column=6).value for r in range(4, 4 + len(paths)))
            want = sorted([2048, 600] + ([1000] if len(paths) == 3 else []))
            assert widths == want, (widths, out)
            close_viewer()

        @step
        def title_on_the_right_turns_upright_number_in_top_right():
            d = os.path.join(WORK, 'side')
            os.makedirs(d, exist_ok=True)
            p = os.path.join(d, 'side.jpg')
            Image.new('RGB', (900, 600), 'white').save(p)
            mock_server.STATE['modes'] = {K1: 'side_right'}
            try:
                run_and_wait([p], '成功 1 张')
            finally:
                mock_server.STATE['modes'] = {}
            page.wait_for_selector('#dlg-view[open]')
            page.wait_for_function('() => { const i = document.querySelector("#vphoto img"); return i && i.complete && i.naturalWidth > 0; }')
            fits = page.evaluate('''() => { const b = document.querySelector("#vphoto").getBoundingClientRect();
                const i = document.querySelector("#vphoto img"), r = i.getBoundingClientRect();
                const s = Math.min(r.width / i.naturalWidth, r.height / i.naturalHeight);   // object-fit: contain 实际画多大
                return [i.naturalWidth * s <= b.width + 1, i.naturalHeight * s <= b.height + 1, r.top >= b.top - 1, r.bottom <= b.bottom + 1]; }''')
            assert all(fits), fits                                               # 竖的照片整张放得下，不被裁掉
            z = download('#vzip')
            im = Image.open(io.BytesIO(zipfile.ZipFile(z).read('01.jpg'))).convert('RGB')
            assert im.size == (600, 900), im.size                                 # 标题在右边：转成竖的
            red = [(x, y) for x in range(0, im.width, 2) for y in range(0, im.height, 2)
                   if (lambda q: q[0] > 170 and q[1] < 70 and q[2] < 70)(im.getpixel((x, y)))]
            corner = [1 for x, y in red if x > im.width * 0.6 and y < im.height * 0.3]
            assert red and len(corner) / len(red) > 0.95, (len(red), len(corner))   # 编号印在右上角，字是正的
            close_viewer()

        @step
        def settings_backup_restore_and_forget_keys():
            page.click('#gear')
            page.wait_for_selector('#dlg-set[open]')
            assert '录入僧网页版 v' in page.inner_text('#setver')
            assert '数据库 22 条' in page.inner_text('#usage'), page.inner_text('#usage')
            page.screenshot(path=os.path.join(SHOTS, '4-设置.png'))
            path = download('#backup')
            data = json.load(open(path, encoding='utf-8'))
            assert data['app'] == 'luruseng-web' and len(data['kb']['records']) == 22 and '质检单' in data['cats']
            assert K1 not in json.dumps(data)                                     # 钥匙不进备份
            page.set_input_files('#restorefile', path)
            wait_log('✓ 恢复好了（合并进来，原来的一条没删）：数据库新增 0 条')
            page.click('#gear')
            page.wait_for_selector('#dlg-set[open]')
            page.click('#setrows label:nth-child(2) .switch')                    # 关掉「记住钥匙」
            page.click('#dlg-set [data-close]')
            page.reload()
            page.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
            assert page.input_value('#key1') == ''
            page.fill('#key1', K1)
            page.click('#gear')
            page.wait_for_selector('#dlg-set[open]')
            page.click('#setrows label:nth-child(2) .switch')
            page.click('#dlg-set [data-close]')
            page.reload()
            page.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
            assert page.input_value('#key1') == K1

        @step
        def delete_a_batch():
            page.click('#check')
            page.wait_for_selector('#dlg-view[open]')
            n0 = page.eval_on_selector('#vruns', 's => s.options.length')
            page.click('#vdel')
            page.click('#dlg-ask [data-close]')
            assert page.eval_on_selector('#vruns', 's => s.options.length') == n0
            page.click('#vdel')
            page.click('#askok')
            page.wait_for_function('n => document.querySelector("#vruns").options.length === n - 1', arg=n0)
            close_viewer()

        @step
        def phone_layout():
            m = browser.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True)
            p = m.new_page()
            p.on('pageerror', lambda e: errors.append(f'phone pageerror: {e}'))
            offline(p)
            p.goto(base)
            p.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
            sw = p.evaluate('() => [document.documentElement.scrollWidth, window.innerWidth]')
            assert sw[0] <= sw[1], sw                                             # 不横着滚
            p.screenshot(path=os.path.join(SHOTS, '5-手机.png'), full_page=True)
            p.set_input_files('#files', imgs[:3])
            p.fill('#key1', K1)
            p.click('#start')
            p.wait_for_selector('#dlg-view[open]', timeout=60000)
            p.wait_for_function('() => { const i = document.querySelector("#vphoto img"); return i && i.naturalWidth > 0; }')
            sw = p.evaluate('() => [document.documentElement.scrollWidth, window.innerWidth]')
            assert sw[0] <= sw[1], sw
            p.screenshot(path=os.path.join(SHOTS, '6-手机核对.png'))
            m.close()

        for fn in (opens_clean, no_key_no_start, twelve_photos_match_desktop_book, viewer_pages_rotates_copies_zips,
                   same_photos_again_are_skipped, result_button_downloads_latest, feed_database_xlsx_xls_csv_html,
                   database_fixes_confusable_letters, names_dialog, own_category_matches_desktop_book,
                   bad_key_stops_and_keeps_photos, stop_button_marks_unfinished, photos_of_all_kinds,
                   title_on_the_right_turns_upright_number_in_top_right, settings_backup_restore_and_forget_keys, delete_a_batch, phone_layout):
            try:
                fn()
            except Exception:
                print(f'  ✗ {fn.__name__}', flush=True)
                print(last_logs(12))
                page.screenshot(path=os.path.join(SHOTS, f'失败-{fn.__name__}.png'))
                raise
        page.screenshot(path=os.path.join(SHOTS, '7-最后.png'))
        browser.close()
    bad = [e for e in errors if 'favicon' not in e]
    assert not bad, '\n'.join(bad)
    print(f'全过：{len(passed)} 项 · 截图在 {SHOTS}')


if __name__ == '__main__':
    main()
