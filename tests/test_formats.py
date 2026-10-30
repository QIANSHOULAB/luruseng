"""放入数据库也认老的 .xls 和 .csv（Excel 中文版存的 GBK、UTF-8 都行），投喂文件夹也一样。"""
import os
import sys
from datetime import datetime

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bangong import kb, sheets, tables  # noqa: E402
from test_kb import company_book  # noqa: E402


def to_xls(xlsx_path, xls_path):
    """测试用：把 openpyxl 做的表另存成老的 .xls（xlwt 只在测试里用，不进 exe）。"""
    xlwt = pytest.importorskip('xlwt')
    wb = load_workbook(xlsx_path)
    out = xlwt.Workbook()
    date_style = xlwt.easyxf(num_format_str='d-mmm')
    for ws in wb.worksheets:
        sh = out.add_sheet(ws.title)
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            for c, v in enumerate(row):
                if v is None:
                    continue
                if isinstance(v, datetime):
                    sh.write(r, c, v, date_style)
                else:
                    sh.write(r, c, v)
    out.save(str(xls_path))
    return str(xls_path)


def test_daodan_from_old_xls_same_as_xlsx(tmp_path):
    x = company_book(tmp_path / 'c.xlsx')
    a = kb.KnowledgeBase(str(tmp_path / 'a.sqlite'))
    b = kb.KnowledgeBase(str(tmp_path / 'b.sqlite'))
    ra = a.feed(str(x))
    rb = b.feed(to_xls(x, tmp_path / 'c.xls'))
    assert not rb.error and rb.added == ra.added == 8 and rb.people == ra.people
    sa, sb = a.load(), b.load()
    assert dict(sa.roster()) == dict(sb.roster())
    assert sorted(sa.count['machine']) == sorted(sb.count['machine'])


def test_csv_in_gbk_and_utf8(tmp_path):
    body = '日期,机台,数量,检验员\n28/10,VM-07,5,张三\n29/10,VM-14,7,李四\n'
    (tmp_path / 'gbk.csv').write_bytes(body.encode('gbk'))                 # Excel 中文版「另存为 CSV」就是这个
    (tmp_path / 'utf8.csv').write_bytes(body.encode('utf-8-sig'))
    (tmp_path / 'semi.csv').write_text(body.replace(',', ';'), encoding='utf-8')   # 有的电脑用分号
    for name in ('gbk.csv', 'utf8.csv', 'semi.csv'):
        db = tables.TableDB(str(tmp_path / f'{name}.sqlite'))
        r = db.feed(str(tmp_path / name))
        assert not r.error and r.added == 2, (name, r.error)
        assert db.columns() == ['日期', '机台', '数量', '检验员']
        assert db.load().records[1]['检验员'] == '李四'


def test_daodan_csv(tmp_path):
    lines = ['DATE,MACHINE,PART NUMBER,MO,OPR NO.,T NO.,DESCRIPTION,BTMP,TOTAL SCREENED,PASSED,FAILED,SHIFT,SETTER',
             '28/10,VM-07,TD-524613-5,82664580,150,T4,EM2.5,1,1,1,0,D,639/Lee Ah Kow']
    p = tmp_path / '刀单.csv'
    p.write_bytes('\n'.join(lines).encode('gbk'))
    db = kb.KnowledgeBase(str(tmp_path / 'db.sqlite'))
    r = db.feed(str(p))
    assert not r.error and r.added == 1 and dict(db.load().roster())['639'] == 'Lee Ah Kow'


def test_feed_folder_eats_xls_and_csv(tmp_path):
    feed = tmp_path / '投喂'
    feed.mkdir()
    (feed / 'a.csv').write_text('日期,机台\n28/10,VM-07\n', encoding='utf-8')
    (feed / '~$a.xlsx').write_bytes(b'lock')                                 # Excel 开着时留的临时文件：不吃
    (feed / 'notes.txt').write_text('不是表', encoding='utf-8')
    db = tables.TableDB(str(tmp_path / 'db.sqlite'))
    got = db.feed_folder(str(feed))
    assert [os.path.basename(r.path) for r in got] == ['a.csv'] and got[0].added == 1
    assert db.feed_folder(str(feed)) == []                                     # 没变就不重吃


def test_fake_xls_from_other_systems(tmp_path):
    """别的系统「导出 Excel」出来的 .xls 常常是假的：其实是 .xlsx、文字表、网页。"""
    import shutil
    x = company_book(tmp_path / 'c.xlsx')
    shutil.copy(x, tmp_path / '改了后缀.xls')                                  # 其实是 .xlsx
    db = kb.KnowledgeBase(str(tmp_path / 'a.sqlite'))
    r = db.feed(str(tmp_path / '改了后缀.xls'))
    assert not r.error and r.added == 8

    (tmp_path / 'tsv.xls').write_bytes('日期\t机台\n28/10\tVM-07\n29/10\tVM-14\n'.encode('gbk'))   # 其实是 Tab 文字表
    (tmp_path / 'u16.xls').write_bytes('日期\t机台\n28/10\tVM-07\n'.encode('utf-16'))   # Excel「Unicode 文本」
    t = tables.TableDB(str(tmp_path / 't.sqlite'))
    assert t.feed(str(tmp_path / 'tsv.xls')).added == 2 and t.columns() == ['日期', '机台']
    t2 = tables.TableDB(str(tmp_path / 't2.sqlite'))
    assert t2.feed(str(tmp_path / 'u16.xls')).added == 1 and t2.columns() == ['日期', '机台']

    (tmp_path / 'web.xls').write_text('<html><table><tr><td>日期</td></tr></table></html>', encoding='utf-8')
    assert '网页' in t.feed(str(tmp_path / 'web.xls')).error                   # 网页：说清楚怎么办
    (tmp_path / 'junk.xls').write_bytes(bytes(range(256)) * 4)
    assert '不是 Excel' in t.feed(str(tmp_path / 'junk.xls')).error
    (tmp_path / 'empty.xls').write_bytes(b'')
    assert '不是 Excel' in t.feed(str(tmp_path / 'empty.xls')).error


def test_old_xls_named_xlsx_and_csv_starting_with_pk(tmp_path):
    x = company_book(tmp_path / 'c.xlsx')
    old = to_xls(x, tmp_path / 'c.xls')
    os.replace(old, tmp_path / '其实是老格式.xlsx')                           # 老 .xls 改了后缀
    db = kb.KnowledgeBase(str(tmp_path / 'a.sqlite'))
    assert db.feed(str(tmp_path / '其实是老格式.xlsx')).added == 8
    (tmp_path / 'pkg.csv').write_text('PKG NO,QTY\nP-1,3\n', encoding='utf-8')  # 第一格 PK 开头：还是 CSV
    t = tables.TableDB(str(tmp_path / 't.sqlite'))
    assert t.feed(str(tmp_path / 'pkg.csv')).added == 1 and t.columns() == ['PKG NO', 'QTY']


def test_locked_and_broken_xls_say_what_to_do(tmp_path, monkeypatch):
    xlrd = pytest.importorskip('xlrd')
    p = tmp_path / 'a.xls'
    p.write_bytes(sheets.OLE2 + b'x' * 600)
    t = tables.TableDB(str(tmp_path / 't.sqlite'))
    assert t.feed(str(p)).error == sheets.BROKEN

    def locked(*a, **k):
        raise xlrd.XLRDError('Workbook is encrypted')
    monkeypatch.setattr(xlrd, 'open_workbook', locked)
    assert t.feed(str(p)).error == sheets.LOCKED


def test_supported_names():
    assert all(sheets.supported(n) for n in ('a.xlsx', 'b.XLSM', 'c.xls', 'd.CSV'))
    assert not any(sheets.supported(n) for n in ('a.pdf', 'b.jpg', 'c.docx', 'd'))
    assert not sheets.touch_ok('~$a.xlsx') and sheets.touch_ok('a.xls')
