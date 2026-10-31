"""发数据给方丈：开着才记、才发；记录、照片、改过的格子一样不少；断网不丢，发过的不重发。"""
import base64
import io
import json
import os
import shutil
import sys
import threading
import time
import urllib.error
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bangong import engine, report, settings  # noqa: E402
import mock_server  # noqa: E402

K1 = 'AIzaKEY1_mock'


@pytest.fixture(scope='module')
def server():
    srv = mock_server.start()
    old = engine.API_BASE, engine.RETRY_WAITS
    engine.API_BASE = f'http://127.0.0.1:{srv.server_address[1]}/v1beta'
    engine.RETRY_WAITS = [0.01, 0.01, 0.01]
    yield srv
    srv.shutdown()
    engine.API_BASE, engine.RETRY_WAITS = old


@pytest.fixture(autouse=True)
def reset(server):
    mock_server.STATE.update(modes={}, log=[], inflight=0, peak=0, delay=0, once={}, calls={})


def make_images(folder, sizes):
    os.makedirs(folder, exist_ok=True)
    for i, (w, h) in enumerate(sizes):
        p = os.path.join(folder, f'img_{i:02d}.jpg')
        Image.new('RGB', (w, h), (i * 40 % 255, 120, 160)).save(p)
        os.utime(p, (1_700_000_000 + i, 1_700_000_000 + i))


def run(folder, **kw):
    logs = []
    eng = engine.Engine(str(folder), [K1], log=logs.append, **kw)
    return eng.run(), logs, eng


def only_run(folder):
    h = os.path.join(folder, engine.OUT_DIR, report.HOME)
    runs = [d for d in os.listdir(h) if d != report.SAMPLES]
    assert len(runs) == 1, runs
    return runs[0], os.path.join(h, runs[0])


def test_collect_writes_reads_table_look_and_photos(tmp_path):
    make_images(tmp_path, [(800, 600), (801, 600), (2400, 1800)])
    mock_server.STATE['modes'] = {K1: 'flip'}             # 第 1、3 遍 D，第 2 遍 N：两遍对不上，认第三遍
    res, logs, eng = run(tmp_path, collect=True, reads=2)
    assert res.ok == 3
    run_id, d = only_run(tmp_path)
    assert run_id == eng.run_name and res.xlsx.endswith(f'识别结果_{run_id}.xlsx')
    rec = json.load(open(os.path.join(d, report.REPORT), encoding='utf-8'))
    assert rec['version'] == engine.VERSION and rec['category'] == '刀单' and rec['reads_per_photo'] == 2
    assert rec['xlsx'] == os.path.basename(res.xlsx) and rec['result']['ok'] == 3 and rec['columns'] == engine.FIELDS
    assert [p['no'] for p in rec['photos']] == [1, 2, 3]
    p1 = rec['photos'][0]
    assert p1['file'] == 'img_00.jpg' and len(p1['hash']) == 64 and p1['side'] == 'top' and p1['ms'] >= 0
    assert [r['img'] for r in p1['reads']] == ['原图', '加强', '原图']                  # 三遍都记着，哪遍用的哪张图
    assert [r['records'][0]['shift'] for r in p1['reads']] == ['D', 'N', 'D']          # 模型原话
    assert p1['look']['w'] == 800 and p1['look']['h'] == 600 and p1['look']['bytes'] > 0
    assert {'light', 'contrast', 'sharp'} <= set(p1['look'])
    row = p1['table'][0]
    assert row['values']['MO'] == '82664580' and row['values']['SHIFT'] == 'D'          # 表里最后写的
    assert 'SHIFT' in row['red'] and '几遍认得不一样' in row['note']                    # 当时标红的格子
    pics = sorted(os.listdir(os.path.join(d, report.PICS)))
    assert pics == ['01.jpg', '02.jpg', '03.jpg']                                       # 跟编号照片一个编号
    big = Image.open(os.path.join(d, report.PICS, pics[2]))
    assert max(big.size) == report.SIDE                                                 # 大照片缩到 1600
    assert os.path.isfile(os.path.join(d, report.PRISTINE))                            # 刚认完时的表留了原样
    assert load_workbook(os.path.join(d, report.PRISTINE)).worksheets[0]['E4'].value == 82664580


def test_switch_off_leaves_no_trace(tmp_path):
    make_images(tmp_path, [(800, 600)])
    res, _, eng = run(tmp_path, reads=1)
    assert res.ok == 1 and not eng.trace
    assert not os.path.exists(os.path.join(tmp_path, engine.OUT_DIR, report.HOME))


def test_failed_photo_is_reported_with_its_error(tmp_path):
    make_images(tmp_path, [(800, 600)])
    mock_server.STATE['modes'] = {K1: 'garbage'}
    res, _, _ = run(tmp_path, collect=True, reads=1)
    assert res.failed == 1
    _, d = only_run(tmp_path)
    p = json.load(open(os.path.join(d, report.REPORT), encoding='utf-8'))['photos'][0]
    assert p['error'] and p['reads'][0]['error'] and p['look']['w'] == 800
    assert os.listdir(os.path.join(d, report.PICS))                                     # 认不出来的照片更要看


def edit_cell(xlsx, ref, value):
    wb = load_workbook(xlsx)
    wb.worksheets[0][ref] = value
    wb.save(xlsx)


def test_corrections_found_after_the_table_is_saved(tmp_path):
    make_images(tmp_path, [(800, 600), (801, 600)])
    res, _, _ = run(tmp_path, collect=True, reads=1)
    run_id, d = only_run(tmp_path)
    later = time.time() + 3600
    assert report.scan_corrections(str(tmp_path), now=later) == 0                      # 没人动过
    edit_cell(res.xlsx, 'E5', 12345678)                                                 # 第 2 张的 MO 改了
    assert report.scan_corrections(str(tmp_path), now=time.time()) == 0                # 刚存完：等一分钟再比
    assert report.scan_corrections(str(tmp_path), now=later) == 1
    fixes = [f for f in os.listdir(d) if f.startswith(report.FIX_PREFIX)]
    got = json.load(open(os.path.join(d, fixes[0]), encoding='utf-8'))
    assert got['changes'] == [{'no': 2, 'k': 0, 'field': 'MO', 'was': '82664580', 'now': '12345678',
                               'red': False, 'fixed': False}]
    assert got['count']['cells'] == 26 and got['count']['changed_not_red'] == 1 and got['run'] == run_id
    assert report.scan_corrections(str(tmp_path), now=later) == 0                      # 没再改：不重记
    os.utime(res.xlsx, (time.time() - 10, time.time() - 10))                           # 又存了一次但内容没变
    assert report.scan_corrections(str(tmp_path), now=later) == 0
    edit_cell(res.xlsx, 'N4', '232')
    os.utime(res.xlsx, (time.time() + 5, time.time() + 5))                             # 新存的一版（时间戳不同）
    assert report.scan_corrections(str(tmp_path), now=later) == 1
    assert len([f for f in os.listdir(d) if f.startswith(report.FIX_PREFIX)]) == 2


def test_broken_table_is_not_taken_as_everything_deleted(tmp_path):
    make_images(tmp_path, [(800, 600)])
    res, _, _ = run(tmp_path, collect=True, reads=1)
    with open(res.xlsx, 'wb') as f:
        f.write(b'not a workbook')
    assert report.scan_corrections(str(tmp_path), now=time.time() + 3600) == 0


def test_diff_counts_red_hits_and_misses():
    before = {1: [{'values': {'MO': '1', 'SHIFT': 'X', 'T NO.': 'T4'}, 'red': ['SHIFT'], 'fixed': ['T NO.'], 'note': ''}],
              2: [{'values': {'MO': '2', 'SHIFT': 'D', 'T NO.': 'T5'}, 'red': [], 'fixed': [], 'note': ''}]}
    after = {1: [{'values': {'MO': '1', 'SHIFT': 'N', 'T NO.': 'T4'}, 'red': [], 'fixed': [], 'note': ''},
                 {'values': {'MO': '9', 'SHIFT': 'D', 'T NO.': 'T1'}, 'red': [], 'fixed': [], 'note': ''}],
             2: [{'values': {'MO': ' 3 ', 'SHIFT': 'D', 'T NO.': 'T5'}, 'red': [], 'fixed': [], 'note': ''}]}
    got = report.diff(before, after)
    assert [(c['no'], c['field'], c['now']) for c in got['changes']] == [(1, 'SHIFT', 'N'), (2, 'MO', ' 3 ')]
    assert got['added'] == [{'no': 1, 'values': {'MO': '9', 'SHIFT': 'D', 'T NO.': 'T1'}}]
    c = got['count']
    assert (c['cells'], c['red'], c['changed'], c['changed_red'], c['changed_not_red'], c['added']) == (6, 1, 2, 1, 1, 1)


def collected(tmp_path, n=2):
    make_images(tmp_path, [(800 + i, 600) for i in range(n)])
    res, _, _ = run(tmp_path, collect=True, reads=1)
    edit_cell(res.xlsx, 'E4', 1)
    report.scan_corrections(str(tmp_path), now=time.time() + 3600)
    return res


def test_send_order_marks_and_cleans_up(tmp_path):
    collected(tmp_path)
    run_id, d = only_run(tmp_path)
    got = []

    def poster(body):
        got.append(body)
        return {'ok': True}
    r = report.send(str(tmp_path), 'PC_abc123', poster=poster)
    names = [b['name'] for b in got]
    assert names[0] == report.REPORT and names[1].endswith('.jpg') and names[2].endswith('.jpg')
    assert names[3].startswith(report.FIX_PREFIX) and len(names) == 4
    assert all(b['token'] == report.TOKEN and b['machine'] == 'PC_abc123' and b['run'] == run_id for b in got)
    assert got[0]['kind'] == 'json' and got[0]['payload']['run'] == run_id
    assert got[1]['kind'] == 'photo' and Image.open(io.BytesIO(base64.b64decode(got[1]['data']))).size == (800, 600)
    assert (r.sent, r.left, r.error) == (4, 0, '')
    assert not os.listdir(os.path.join(d, report.PICS))                    # 照片副本发完就删
    assert os.path.isfile(os.path.join(d, report.REPORT))                  # 记录和原样留着（以后还要比）
    got.clear()
    assert report.send(str(tmp_path), 'PC_abc123', poster=poster).sent == 0 and not got   # 发过的不重发


def test_network_down_keeps_the_rest_for_later(tmp_path):
    collected(tmp_path)
    calls = []

    def flaky(body):
        calls.append(body['name'])
        if len(calls) == 2:
            raise urllib.error.URLError('断网了')
        return {'ok': True}
    r = report.send(str(tmp_path), 'PC', poster=flaky)
    assert r.sent == 1 and r.error.startswith('网络') and r.left == 3
    rest = []
    r2 = report.send(str(tmp_path), 'PC', poster=lambda b: rest.append(b['name']) or {'ok': True})
    assert rest[0] == calls[1] and len(rest) == 3 and r2.left == 0               # 从断的那份接着发


def test_refusals(tmp_path):
    collected(tmp_path)
    r = report.send(str(tmp_path), 'PC', poster=lambda b: {'ok': False, 'error': 'token'})
    assert r.error == 'token' and r.sent == 0 and r.left == 4                     # 不收：一份都不算发了
    seen = []

    def too_big_first(b):
        seen.append(b['name'])
        return {'ok': False, 'error': 'size'} if len(seen) == 1 else {'ok': True, 'dup': len(seen) == 2}
    r = report.send(str(tmp_path), 'PC', poster=too_big_first)
    assert r.error == '' and r.left == 0 and r.sent == 3 and len(seen) == 4        # 太大的跳过，重复的也算送到


def test_time_budget_stops_early(tmp_path):
    collected(tmp_path)
    t = [0.0]

    def clock():
        t[0] += 50
        return t[0]
    r = report.send(str(tmp_path), 'PC', poster=lambda b: {'ok': True}, limit_s=120, clock=clock)
    assert 0 < r.sent < 4 and r.left == 4 - r.sent and not r.error


class AppsScriptLike(BaseHTTPRequestHandler):
    """跟谷歌一样：POST 先干活，回 302 叫你去另一个地址用 GET 拿结果。"""
    got = []

    def log_message(self, *a):
        pass

    def do_POST(self):
        n = int(self.headers['Content-Length'])
        AppsScriptLike.got.append((self.headers.get('Content-Type'), json.loads(self.rfile.read(n).decode('utf-8'))))
        self.send_response(302)
        self.send_header('Location', '/echo?user_content_key=abc')
        self.end_headers()

    def do_GET(self):
        data = json.dumps({'ok': True, 'via': self.path}).encode()
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(data)))
        self.end_headers()
        self.wfile.write(data)


def test_post_follows_googles_redirect():
    srv = ThreadingHTTPServer(('127.0.0.1', 0), AppsScriptLike)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        url = f'http://127.0.0.1:{srv.server_address[1]}/macros/s/x/exec'
        r = report.post({'token': report.TOKEN, 'machine': '车间电脑', 'name': '记录.json', 'kind': 'json',
                         'payload': {'说': '方丈好'}}, url=url)
        assert r['ok'] and r['via'].startswith('/echo')
        ctype, body = AppsScriptLike.got[-1]
        assert 'charset=utf-8' in ctype and body['machine'] == '车间电脑' and body['payload'] == {'说': '方丈好'}
    finally:
        srv.shutdown()


def test_ping_says_why_not():
    srv = ThreadingHTTPServer(('127.0.0.1', 0), AppsScriptLike)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        assert report.ping(f'http://127.0.0.1:{srv.server_address[1]}/exec') == ''
    finally:
        srv.shutdown()
    port = srv.server_address[1]
    assert report.ping(f'http://127.0.0.1:{port}/exec', timeout=2) != ''              # 连不上：说出原因


def test_samples_queue_once_and_go_last(tmp_path):
    src = tmp_path / '改对的刀单.xlsx'
    src.write_bytes(b'PK\x03\x04 pretend workbook')
    p1 = report.queue_sample(str(tmp_path), str(src), '刀单')
    p2 = report.queue_sample(str(tmp_path), str(src), '刀单')
    assert p1 == p2 and len(os.listdir(os.path.dirname(p1))) == 1
    data = json.load(open(p1, encoding='utf-8'))
    assert base64.b64decode(data['b64']) == src.read_bytes() and data['category'] == '刀单'
    assert report.pending(str(tmp_path)) == [(report.SAMPLES, os.path.basename(p1), p1)]


def test_feed_folder_tables_are_queued_as_samples(tmp_path):
    work = tmp_path / '待识别'
    make_images(work, [(800, 600)])
    res, _, _ = run(work, collect=True, reads=1)
    feed = tmp_path / '投喂'
    feed.mkdir()
    shutil.copy(res.xlsx, feed / '改好的.xlsx')                          # 刀单格式的表：数据库吃得下
    make_images(work, [(900, 600)])
    run(work, collect=True, reads=1, kb_path=str(tmp_path / 'kb.sqlite'), feed_dir=str(feed))
    samples = os.listdir(os.path.join(work, engine.OUT_DIR, report.HOME, report.SAMPLES))
    assert len(samples) == 1 and '改好的' in samples[0]


def test_no_room_means_records_but_no_photos(tmp_path, monkeypatch):
    monkeypatch.setattr(report, 'PHOTO_ROOM', 0)
    make_images(tmp_path, [(800, 600)])
    run(tmp_path, collect=True, reads=1)
    _, d = only_run(tmp_path)
    assert os.path.isfile(os.path.join(d, report.REPORT)) and not os.path.isdir(os.path.join(d, report.PICS))


def test_machine_name_is_stable(monkeypatch):
    monkeypatch.setenv('COMPUTERNAME', 'WORKSHOP PC/2')
    cfg = {}
    a = report.machine(cfg)
    assert a == f'WORKSHOP_PC_2_{cfg["report_id"]}' and len(cfg['report_id']) == 6
    assert report.machine(cfg) == a


def test_switch_is_last_row_and_off_by_default():
    assert settings.ROWS[-1][:2] == ('report', '发数据给方丈')
    assert settings.NEW['report'] is False and settings.OLD['report'] is False


def test_settings_switch_and_background_sending(tmp_path, monkeypatch):
    try:
        import tkinter as tk
        root = tk.Tk()
    except Exception:
        pytest.skip('没有显示器，跳过窗口测试')
    import types
    from bangong import app, macui as ui
    monkeypatch.delenv('LURUSENG_SCALE', raising=False)
    monkeypatch.setattr(app, 'app_dir', lambda: str(tmp_path / 'app'))
    monkeypatch.setattr(app, 'CONFIG', str(tmp_path / 'config.json'))
    scaling0 = root.tk.call('tk', 'scaling')
    try:
        a = app.App(root)
        a.folder.set(str(tmp_path / '待识别'))
        a.open_settings()
        root.update()
        card = a.settings_win
        assert card.vars['report'].get() is False and card.note.get() == ''
        monkeypatch.setattr(app.report, 'ping', lambda: '')
        card.vars['report'].set(True)
        card.changed('report', True)
        cfg = json.load(open(tmp_path / 'config.json', encoding='utf-8'))
        assert cfg['report'] is True and len(cfg['report_id']) == 6
        assert card.note.get().startswith('发给方丈：') and '发数据给方丈：开了' in a.text.get('1.0', 'end')
        for _ in range(40):                                             # 当场问一声网盘通不通
            root.update()
            a.pump()
            if '方丈网盘连上了' in a.text.get('1.0', 'end'):
                break
            time.sleep(0.05)
        assert '方丈网盘连上了' in a.text.get('1.0', 'end') and card.note.get().startswith('✓ 方丈网盘连上了 · 发给方丈')

        seen = {}
        monkeypatch.setattr(app.engine, 'Engine', lambda *x, **kw: seen.update(kw) or types.SimpleNamespace(run=lambda: None))
        a.key1.set('AIzaX')
        a.start()
        a.worker.join(5)
        assert seen['collect'] is True                                  # 开着：认照片时把过程记下来

        calls = []
        monkeypatch.setattr(app.report, 'scan_corrections', lambda folder: calls.append(('scan', folder)) or 0)
        monkeypatch.setattr(app.report, 'send', lambda folder, who: calls.append(('send', who)) or report.Sent(sent=3))
        a.next_send = a.next_scan = 0
        a._tick()
        for _ in range(50):
            root.update()
            a.pump()
            if not a.sending:
                break
            time.sleep(0.05)
        assert [c[0] for c in calls] == ['scan', 'send'] and calls[1][1].endswith('_' + cfg['report_id'])
        assert a.reported == 3 and a.next_send > time.time() + 500 and a.send_fails == 0

        monkeypatch.setattr(app.report, 'send', lambda folder, who: report.Sent(error='网络：断了'))
        a.next_send = 0
        a._tick()
        for _ in range(50):
            root.update()
            a.pump()
            if not a.sending:
                break
            time.sleep(0.05)
        assert a.send_fails == 1 and 100 < a.next_send - time.time() <= 120   # 断网：两分钟后再试

        card.changed('report', False)
        assert card.note.get() == '' and '关了' in a.text.get('1.0', 'end')
        calls.clear()
        a.next_send = 0
        a._tick()
        assert not calls and not a.sending                              # 关了就不发
    finally:
        root.tk.call('tk', 'scaling', scaling0)
        root.destroy()
        ui.S = 1.0
