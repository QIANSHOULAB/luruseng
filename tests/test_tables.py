"""自建类别：建类别、放入数据库（表头定列）、认全部内容、按列认、按库改写法、核对窗口、跟刀单互不串味。"""
import os
import sys

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bangong import engine, tables, viewer  # noqa: E402
from bangong.kb import KnowledgeBase  # noqa: E402
import mock_server  # noqa: E402
from test_engine import K1, make_images, server, reset  # noqa: E402,F401


def book(path, rows, header=('日期', '机台', '数量', '备注'), title=None):
    wb = Workbook()
    ws = wb.active
    if title:
        ws.append([title])
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return str(path)


# ---------------------------------------------------------------- 类别文件夹
def test_create_and_list_categories(tmp_path):
    app = str(tmp_path)
    assert tables.names(app) == ['刀单']
    assert tables.check_name('', app) and tables.check_name('刀单', app) and tables.check_name('a/b', app)
    assert tables.check_name('CON', app) and tables.check_name('x' * 21, app)
    assert tables.check_name('质检单', app) == ''
    d = tables.create(app, '质检单')
    assert os.path.isfile(os.path.join(d, tables.DB_FILE)) and os.path.isdir(os.path.join(d, tables.FEED_DIR))
    tables.create(app, '送货单')
    assert tables.names(app) == ['刀单', '质检单', '送货单']
    assert '已经有' in tables.check_name('质检单', app)


# ---------------------------------------------------------------- 放入数据库：表头定列
def test_first_table_sets_columns_and_later_ones_append(tmp_path):
    db = tables.TableDB(str(tmp_path / 'db.sqlite'))
    r = db.feed(book(tmp_path / 'a.xlsx', [('28/10', 'VM-07', 5, ''), ('29/10', 'VM-14', 7, '急')]))
    assert not r.error and r.added == 2 and db.columns() == ['日期', '机台', '数量', '备注']
    assert db.feed(str(tmp_path / 'a.xlsx')).added == 0                         # 同一本再放：不重复
    r = db.feed(book(tmp_path / 'b.xlsx', [('30/10', 'VM-07', 3, '', '张三')], header=('日期', '机台', '数量', '备注', '签名')))
    assert r.added == 1 and db.columns() == ['日期', '机台', '数量', '备注', '签名']
    assert '3 条记录' in db.summary('质检单')
    assert db.feed(str(tmp_path / 'x.xls')).error
    assert db.feed(book(tmp_path / 'empty.xlsx', [], header=('只有一格',))).error


def test_reads_bangong_own_output_without_no_and_notes(tmp_path):
    p = tmp_path / 'out.xlsx'
    tb = tables.TableBook(str(p), '质检单', ['日期', '机台', '数量'])
    tb.add(1, [{'日期': '28/10', '机台': 'VM-07', '数量': '5', 'unsure': []}])
    tb.add(2, None, '超时')
    tb.add(3, [{'日期': '29/10', '机台': 'VM-□4', '数量': '7', 'unsure': []}])
    tb.save()
    cols, recs, used = tables.read_book(str(p))
    assert cols == ['日期', '机台', '数量']                                   # No. 和「核对提示」不算列
    assert [r for _, r in recs] == [{'日期': '28/10', '机台': 'VM-07', '数量': '5'},
                                    {'日期': '29/10', '机台': 'VM-□4', '数量': '7'}]   # 失败的空行不进库


# ---------------------------------------------------------------- 认：先认全部内容，放进库，再按列认
def test_open_mode_then_columns_then_snapping(tmp_path):
    app = tmp_path / 'app'
    tables.create(str(app), '质检单')
    photos = tmp_path / '待识别'
    make_images(photos, 3)
    mock_server.STATE['delay'] = 0

    spec = tables.TableSpec('质检单', tables.db_path(str(app), '质检单'), tables.feed_path(str(app), '质检单'))
    res = engine.Engine(str(photos), [K1], log=lambda m: None, spec=spec, reads=1).run()
    assert res.ok == 3 and res.xlsx.endswith('_质检单.xlsx')
    ws = load_workbook(res.xlsx).active
    assert [ws.cell(row=2, column=c).value for c in range(1, 6)] == ['No.', 'Date', 'M/C No.', 'Qty', 'Remark']
    assert ws.cell(row=1, column=1).value == '质检单'
    assert [ws.cell(row=r, column=1).value for r in (3, 4, 5)] == [1, 2, 3]

    # 人改对了这本，放进数据库：表头定成这个类别的列
    db = tables.TableDB(tables.db_path(str(app), '质检单'))
    fr = db.feed(res.xlsx)
    assert fr.added == 3 and db.columns() == ['Date', 'M/C No.', 'Qty', 'Remark']

    # 再认：按定好的列认；VM-O7 按库改成 VM-07（橙），Qty 是数字列
    make_images(photos, 2, start=10)
    mock_server.STATE['modes'] = {K1: 'cat_confusable'}
    spec = tables.TableSpec('质检单', tables.db_path(str(app), '质检单'), tables.feed_path(str(app), '质检单'))
    said = []
    res = engine.Engine(str(photos), [K1], log=said.append, spec=spec, reads=1).run()
    assert res.ok == 2 and any('4 列' in m for m in said)
    assert 'c2：M/C No.' in spec.prompt() and 'VM-07' in spec.prompt() and '以前录对的几条' in spec.prompt()
    ws = load_workbook(res.xlsx).active
    assert ws.cell(row=3, column=3).value == 'VM-07'
    assert ws.cell(row=3, column=3).fill.fgColor.rgb.endswith('FCE4D6')
    assert ws.cell(row=3, column=4).value == 810 and isinstance(ws.cell(row=3, column=4).value, int)
    assert 'M/C No. VM-O7→VM-07' in str(ws.cell(row=3, column=6).value)


def test_open_mode_two_reads_unify_columns_and_flag_disagreement(tmp_path):
    photos = tmp_path / '待识别'
    make_images(photos, 1)
    mock_server.STATE.update(delay=0, modes={K1: 'open_flip'})
    spec = tables.TableSpec('送货单', str(tmp_path / 'db.sqlite'))
    res = engine.Engine(str(photos), [K1], log=lambda m: None, spec=spec, reads=2).run()
    assert res.ok == 1 and res.requests == 3                       # 两遍对不上，第三遍当裁判
    ws = load_workbook(res.xlsx).active
    heads = [ws.cell(row=2, column=c).value for c in range(1, 7)]
    assert heads[:5] == ['No.', 'Date', 'M/C No.', 'Qty', 'Remark']    # 「date」「M/C No」「QTY」认成同一列
    assert ws.cell(row=3, column=4).value == '800'                # 三遍里两遍是 800
    assert ws.cell(row=3, column=4).fill.fgColor.rgb.endswith('FFC7CE')
    assert '几遍认得不一样' in str(ws.cell(row=3, column=6).value)


def test_unsure_column_goes_red(tmp_path):
    db = tables.TableDB(str(tmp_path / 'db.sqlite'))
    db.feed(book(tmp_path / 'a.xlsx', [('28/10', 'VM-07', 5, '')]))
    photos = tmp_path / '待识别'
    make_images(photos, 1)
    mock_server.STATE.update(delay=0, modes={K1: 'cat_unsure'})
    spec = tables.TableSpec('质检单', str(tmp_path / 'db.sqlite'))
    res = engine.Engine(str(photos), [K1], log=lambda m: None, spec=spec, reads=1).run()
    ws = load_workbook(res.xlsx).active
    assert ws.cell(row=3, column=4).fill.fgColor.rgb.endswith('FFC7CE') and '数量 模型拿不准' in str(ws.cell(row=3, column=6).value)


# ---------------------------------------------------------------- 互不串味
def test_daodan_db_refuses_other_forms(tmp_path):
    kb = KnowledgeBase(str(tmp_path / '数据库.sqlite'))
    r = kb.feed(book(tmp_path / 'qc.xlsx', [('28/10', 'VM-07', 5, ''), ('639', 'Lee Ah Kow', 1, '')]))
    assert r.error and '不像刀单' in r.error
    assert kb.load().n == 0 and kb.load().roster() == []


def test_category_runs_never_touch_daodan_db(tmp_path):
    photos = tmp_path / '待识别'
    make_images(photos, 2)
    mock_server.STATE['delay'] = 0
    dao = tmp_path / '数据库.sqlite'
    KnowledgeBase(str(dao))
    before = dao.read_bytes()
    spec = tables.TableSpec('质检单', str(tmp_path / '类别' / '质检单' / '数据库.sqlite'))
    engine.Engine(str(photos), [K1], log=lambda m: None, spec=spec, reads=1).run()
    assert dao.read_bytes() == before


# ---------------------------------------------------------------- 核对窗口认得自建类别的表
def test_viewer_reads_category_book(tmp_path):
    p = tmp_path / '识别结果_20261030_120000_质检单.xlsx'
    tb = tables.TableBook(str(p), '质检单', ['日期', '机台', '数量'])
    tb.add(1, [{'日期': '28/10', '机台': 'VM-□7', '数量': '5', 'unsure': []}])
    tb.add(2, [{'日期': '29/10', '机台': 'VM-14', '数量': '7', 'unsure': []},
               {'日期': '29/10', '机台': 'VM-15', '数量': '1', 'unsure': []}])
    tb.save()
    rows = viewer.load_rows(str(p))
    assert rows[1][0]['labels'] == ['日期', '机台', '数量'] and rows[1][0]['values'] == ['28/10', 'VM-□7', '5']
    assert rows[1][0]['red'] == {1} and '看不清' in rows[1][0]['note']
    assert len(rows[2]) == 2


# ---------------------------------------------------------------- 窗口里走一遍：新建类别 → 认 → 放入数据库
def test_window_flow(tmp_path, monkeypatch):
    import time
    import pytest
    try:
        import tkinter as tk
        root = tk.Tk()
    except Exception:
        pytest.skip('没有显示器，跳过窗口测试')
    from bangong import app
    monkeypatch.setattr(app, 'app_dir', lambda: str(tmp_path / 'app'))
    monkeypatch.setattr(app, 'CONFIG', str(tmp_path / 'config.json'))
    monkeypatch.setattr(app, 'open_path', lambda p: None)
    mock_server.STATE['delay'] = 0
    try:
        a = app.App(root)
        a.cfg['auto_open'] = False
        assert a.category.get() == '刀单' and '专属优化' in a.cat_hint.get()
        a.new_category()
        a.dialog.name.set('质检单')
        a.dialog.ok()
        assert a.category.get() == '质检单' and '还没放过表' in a.cat_hint.get()
        assert list(a.cat_combo.box.cget('values')) == ['刀单', '质检单']
        photos = tmp_path / '待识别'
        make_images(photos, 2)
        a.folder.set(str(photos))
        a.key1.set(K1)
        a.start()

        def wait(cond):
            end = time.time() + 30
            while not cond() and time.time() < end:
                root.update()
                time.sleep(0.05)
            root.update()
        wait(lambda: not (a.worker and a.worker.is_alive()) and a.start_btn.cget('state') == 'normal')
        out = [f for f in os.listdir(photos / engine.OUT_DIR) if f.endswith('_质检单.xlsx')]
        assert len(out) == 1 and a.last_run.get('ok') == 2
        a.do_feed('质检单', [str(photos / engine.OUT_DIR / out[0])])
        wait(lambda: not a.feeding)
        assert '4 列' in a.cat_hint.get() and '2 条记录' in a.cat_hint.get()
        assert app.load_config()['category'] == '质检单'
    finally:
        root.destroy()


def test_update_card_and_contact(tmp_path, monkeypatch):
    """别人的电脑：有新版弹卡片（更新了啥 + 现在更新 / 以后都自动更新 / 以后再说），右上角留蓝色小按钮；⚙ 里能联系制作人。"""
    import pytest
    try:
        import tkinter as tk
        root = tk.Tk()
    except Exception:
        pytest.skip('没有显示器，跳过窗口测试')
    from bangong import app, updater
    monkeypatch.setattr(app, 'app_dir', lambda: str(tmp_path / 'app'))
    monkeypatch.setattr(app, 'CONFIG', str(tmp_path / 'config.json'))
    monkeypatch.setattr(updater, 'notes', lambda since=None, timeout=15: '- 后来拿到的说明')
    try:
        a = app.App(root)
        assert a.cfg['auto_update'] is False                         # 新装：不自动更新
        a.offer_update('99.0.0', '- 大更新')
        root.update()
        card = a.dialog
        assert isinstance(card, app.UpdateCard) and card.notes.get() == '• 大更新'
        assert card.title() == '录入僧 v99.0.0 搞出来了'
        assert a.update_pill.winfo_ismapped() and a.update_pill.cget('text') == '新版本 v99.0.0'
        card.later()
        assert a.update_dismissed == '99.0.0' and not card.winfo_exists()
        a.reopen_update()                                          # 点右上角蓝色小按钮：卡片回来，说明自己去拿
        import time
        for _ in range(60):
            root.update()
            if a.dialog.notes.get() == '• 后来拿到的说明':
                break
            time.sleep(0.05)
        assert a.dialog.notes.get() == '• 后来拿到的说明'
        a.dialog.always()                                          # 以后都自动更新
        root.update()
        assert a.cfg['auto_update'] is True and app.load_config()['auto_update'] is True
        a.open_settings()
        root.update()
        a.settings_win.copy_mail()
        assert root.clipboard_get() == updater.CONTACT == 'oilsongchengmin@gmail.com'
    finally:
        root.destroy()


def test_key_tutorial_buttons(tmp_path, monkeypatch):
    """钥匙框后面：「免费领钥匙」「再领一把」点了打开网站上的教程页；没填钥匙就点开始，也告诉人去哪领。"""
    import pytest
    try:
        import tkinter as tk
        root = tk.Tk()
    except Exception:
        pytest.skip('没有显示器，跳过窗口测试')
    from bangong import app, updater
    monkeypatch.setattr(app, 'app_dir', lambda: str(tmp_path / 'app'))
    monkeypatch.setattr(app, 'CONFIG', str(tmp_path / 'config.json'))
    opened = []
    monkeypatch.setattr(app, 'open_path', opened.append)
    try:
        a = app.App(root)
        root.update()
        assert [b.cget('text') for b in a.key_btns] == ['免费领钥匙', '再领一把']
        for b in a.key_btns:
            b.command()
        assert opened == [updater.KEY_URL, updater.KEY_URL + '#key2']
        assert updater.KEY_URL == 'https://oilsongchengmin-svg.github.io/luruseng/key.html'
        a.key1.set('')
        a.key2.set('')
        a.start()
        root.update()
        assert '免费领钥匙' in a.text.get('1.0', 'end')
    finally:
        root.destroy()


def test_key_page_is_there():
    """网站上的教程页：按钮指过去的地方得真有（key.html 和 #key2），去谷歌领钥匙的网址对。"""
    page = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'key.html')
    html = open(page, encoding='utf-8').read()
    assert 'id="key2"' in html and 'https://aistudio.google.com/apikey' in html
    assert '老衲不打诳语' in html and '打死别绑银行卡' in html


def test_pretty_notes():
    from bangong.app import pretty_notes
    assert pretty_notes('v1.9.2\n- 改格子\n- 发邮件') == '• 改格子\n• 发邮件'           # 只跨一版：不写版本号
    assert pretty_notes('v1.9.2\n- a\nv1.9.1\n- b') == 'v1.9.2\n• a\nv1.9.1\n• b'   # 跨了几版：留着版本号
    assert pretty_notes('') == ''
