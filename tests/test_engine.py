import json
import os
import shutil
import sys
import threading
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bangong import engine  # noqa: E402
import mock_server  # noqa: E402

K1, K2 = 'AIzaKEY1_mock', 'AIzaKEY2_mock'


@pytest.fixture(scope='session')
def server():
    srv = mock_server.start()
    engine.API_BASE = f'http://127.0.0.1:{srv.server_address[1]}/v1beta'
    engine.RETRY_WAITS = [0.01, 0.01, 0.01]
    yield srv
    srv.shutdown()


@pytest.fixture(autouse=True)
def reset(server):
    mock_server.STATE.update(modes={}, log=[], inflight=0, peak=0, delay=0.05, once={})


def make_images(folder, n, start=0, size=(800, 600)):
    os.makedirs(folder, exist_ok=True)
    paths = []
    for i in range(start, start + n):
        p = os.path.join(folder, f'img_{i:02d}.jpg')
        Image.new('RGB', size, (i * 7 % 255, 100, 150)).save(p)
        os.utime(p, (1_700_000_000 + i, 1_700_000_000 + i))
        paths.append(p)
    return paths


class FakeClock:
    def __init__(self, t):
        self.t = t
        self.slept = []
        self.hook = None

    def __call__(self):
        return self.t

    def sleep(self, s):
        self.slept.append(s)
        self.t += s
        if self.hook:
            self.hook(self)


def run(folder, keys=(K1,), **kw):
    logs = []
    eng = engine.Engine(str(folder), list(keys), log=logs.append, **kw)
    res = eng.run()
    return res, logs, eng


def main_rows(xlsx):
    return [list(r) for r in load_workbook(xlsx)['识别结果'].iter_rows(values_only=True)]


def listing(xlsx):
    return [list(r) for r in load_workbook(xlsx)['清单'].iter_rows(values_only=True)][1:]


def test_twelve_images_three_batches(tmp_path):
    make_images(tmp_path, 12)
    res, logs, _ = run(tmp_path)
    log = mock_server.STATE['log']
    assert res.ok == 12 and res.failed == 0 and res.requests == 12
    assert len(log) == 12 and all(e['key'] == K1 and e['hdr'] is None for e in log)
    assert mock_server.STATE['peak'] <= 5
    assert all('temperature' not in e['gen'] and 'rows' in e['gen']['responseSchema']['properties'] for e in log)
    assert sorted(os.listdir(tmp_path / '已识别')) == [f'img_{i:02d}.jpg' for i in range(12)]
    assert not [f for f in os.listdir(tmp_path) if f.endswith('.jpg')]
    rows = main_rows(res.xlsx)
    titles = [r[0] for r in rows if r[0] and str(r[0]).startswith('第 ')]
    assert titles == [f'第 {i + 1} 张　img_{i:02d}.jpg' for i in range(12)]
    assert ['开始时间', '8:30', None] in rows
    assert [r[2] for r in listing(res.xlsx)] == ['成功'] * 12
    ledger = json.load(open(tmp_path / '识别结果' / engine.LEDGER, encoding='utf-8'))
    assert len(ledger) == 12
    assert (tmp_path / '识别结果' / '搬工日志.txt').exists()


def test_rerun_skips_done_even_renamed(tmp_path):
    make_images(tmp_path, 3)
    run(tmp_path)
    mock_server.STATE['log'] = []
    shutil.copy(tmp_path / '已识别' / 'img_01.jpg', tmp_path / '改了名.jpg')
    make_images(tmp_path, 1, start=50)
    res, logs, _ = run(tmp_path)
    assert res.ok == 1 and res.skipped == 1 and len(mock_server.STATE['log']) == 1
    assert any('改了名.jpg 以前干过' in m for m in logs)
    assert not [f for f in os.listdir(tmp_path) if f.endswith('.jpg')]


def test_nothing_to_do(tmp_path):
    res, logs, _ = run(tmp_path)
    assert res.requests == 0 and res.xlsx == '' and '没有新图片' in logs[-1]


def test_key1_out_for_day_switches_to_key2(tmp_path):
    make_images(tmp_path, 10)
    mock_server.STATE['modes'] = {K1: 'day'}
    res, logs, eng = run(tmp_path, keys=(K1, K2))
    keys = [e['key'] for e in mock_server.STATE['log']]
    assert res.ok == 10 and keys.count(K2) == 10 and keys.count(K1) <= 5
    assert eng.keys.dead_until[0] > eng.clock() + 60
    assert '钥匙2×10' in logs[-1]


def test_both_out_for_day_waits_until_reset_then_finishes(tmp_path):
    make_images(tmp_path, 7)
    clock = FakeClock(datetime(2026, 10, 29, 8, 0, tzinfo=timezone.utc).timestamp())   # 新加坡 16:00
    mock_server.STATE['modes'] = {K1: 'day', K2: 'day'}

    def reset_quota(c):
        if c.t >= datetime(2026, 10, 30, 7, 0, tzinfo=timezone.utc).timestamp():       # 太平洋午夜
            mock_server.STATE['modes'] = {}
    clock.hook = reset_quota
    res, logs, _ = run(tmp_path, keys=(K1, K2), clock=clock, sleep=clock.sleep)
    assert res.ok == 7 and res.failed == 0
    assert any('等到' in m and '额度刷新后接着干' in m for m in logs)
    assert max(clock.slept) > 20 * 3600


def test_minute_limit_waits_then_continues(tmp_path):
    make_images(tmp_path, 3)
    clock = FakeClock(1_800_000_000)
    mock_server.STATE['once'] = {K1: ('minute', 3)}
    res, logs, _ = run(tmp_path, clock=clock, sleep=clock.sleep)
    assert res.ok == 3 and res.failed == 0
    assert clock.slept and max(clock.slept) <= 60


def test_prepay_depleted_stops_at_once(tmp_path):
    make_images(tmp_path, 8)
    mock_server.STATE['modes'] = {K1: 'prepay', K2: 'ok'}
    res, logs, _ = run(tmp_path, keys=(K1, K2))
    assert res.stopped_reason and 'prepayment credits are depleted' in res.stopped_reason
    assert res.ok == 0 and len([f for f in os.listdir(tmp_path) if f.endswith('.jpg')]) == 8
    assert len(mock_server.STATE['log']) <= 5 and res.xlsx == ''


def test_invalid_key_stops(tmp_path):
    make_images(tmp_path, 2)
    mock_server.STATE['modes'] = {K1: 'invalid'}
    res, logs, _ = run(tmp_path)
    assert 'HTTP 400：API key not valid' in res.stopped_reason


def test_server_error_retried(tmp_path):
    make_images(tmp_path, 2)
    mock_server.STATE['once'] = {K1: ('500', 2)}
    res, logs, _ = run(tmp_path)
    assert res.ok == 2 and res.requests == 4


def test_bad_answer_fails_only_that_image_and_rerun_retries_it(tmp_path):
    make_images(tmp_path, 3)
    mock_server.STATE['delay'] = 0
    mock_server.STATE['once'] = {K1: ('garbage', 1)}
    res, logs, _ = run(tmp_path)
    assert res.ok == 2 and res.failed == 1
    left = [f for f in os.listdir(tmp_path) if f.endswith('.jpg')]
    assert len(left) == 1
    lst = listing(res.xlsx)
    assert sorted(r[2] for r in lst) == ['失败', '成功', '成功'] and any('JSON 解析不了' in (r[4] or '') for r in lst)
    res2, _, _ = run(tmp_path)
    assert res2.ok == 1 and res2.requests == 1


def test_stop_button(tmp_path):
    make_images(tmp_path, 15)
    mock_server.STATE['delay'] = 0.3
    stop = threading.Event()
    threading.Timer(0.15, stop.set).start()
    res, logs, _ = run(tmp_path, stop=stop)
    assert res.stopped_reason == '手动停止' and res.ok < 15
    res2, _, _ = run(tmp_path)
    assert res.ok + res2.ok == 15


def test_image_prep_heic_alpha_rotation_large(tmp_path):
    import pillow_heif
    big = tmp_path / 'big.png'
    Image.new('RGBA', (4032, 3024), (0, 0, 0, 0)).save(big)
    out = Image.open(__import__('io').BytesIO(engine.prepare_image(str(big))))
    assert out.size == (2048, 1536) and out.getpixel((5, 5)) == (255, 255, 255)
    heic = tmp_path / 'photo.heic'
    pillow_heif.from_pillow(Image.new('RGB', (1000, 500), 'red')).save(str(heic))
    assert Image.open(__import__('io').BytesIO(engine.prepare_image(str(heic)))).size == (1000, 500)
    rot = tmp_path / 'rot.jpg'
    im = Image.new('RGB', (600, 300), 'blue')
    exif = im.getexif()
    exif[0x0112] = 6
    im.save(rot, exif=exif)
    assert Image.open(__import__('io').BytesIO(engine.prepare_image(str(rot)))).size == (300, 600)


def test_pacific_midnight():
    f = lambda *a: datetime(*a, tzinfo=timezone.utc).timestamp()
    assert engine.next_pacific_midnight(f(2026, 10, 29, 8, 0)) == f(2026, 10, 30, 7, 0) + 30   # 夏令时：新加坡 15:00
    assert engine.next_pacific_midnight(f(2026, 11, 2, 8, 30)) == f(2026, 11, 3, 8, 0) + 30    # 冬令时：新加坡 16:00
    assert engine.next_pacific_midnight(f(2026, 11, 1, 5, 0)) == f(2026, 11, 1, 7, 0) + 30     # 切换当天
    assert engine.next_pacific_midnight(f(2027, 3, 14, 9, 0)) == f(2027, 3, 15, 7, 0) + 30     # 3 月切回夏令时


def test_parse_rows_variants():
    wrap = lambda t: {'candidates': [{'content': {'parts': [{'text': t}]}, 'finishReason': 'STOP'}]}
    assert engine.parse_rows(wrap('```json\n{"rows":[["a","b"],[]]}\n```')) == [['a', 'b']]
    assert engine.parse_rows(wrap('[["x"]]')) == [['x']]
    with pytest.raises(ValueError):
        engine.parse_rows({'promptFeedback': {'blockReason': 'SAFETY'}})


def test_schema_rejected_falls_back_to_plain_json(tmp_path):
    make_images(tmp_path, 6)
    mock_server.STATE['modes'] = {K1: 'noschema'}
    res, logs, _ = run(tmp_path)
    assert res.ok == 6 and res.failed == 0
    assert any('改用纯 JSON' in m for m in logs)
    assert sum(1 for e in mock_server.STATE['log'] if 'responseSchema' not in e['gen']) == 6


def test_model_prefix_and_thinking():
    assert engine.model_name('models/gemini-2.5-flash') == 'gemini-2.5-flash'
    assert engine.build_body(b'x', 'gemini-2.5-flash')['generationConfig']['thinkingConfig'] == {'thinkingBudget': 0}
    g3 = engine.build_body(b'x', 'gemini-3.5-flash-lite')['generationConfig']
    assert 'thinkingConfig' not in g3 and 'temperature' not in g3
    assert engine.build_body(b'x', 'gemini-2.5-flash-lite')['generationConfig']['temperature'] == 0.1


def test_excel_survives_formula_and_control_chars(tmp_path):
    book = engine.Book(str(tmp_path / 'x.xlsx'))
    book.add('a.jpg', [['=1+1', 'ok\x07bell', '\x00'], ['普通']])
    book.add('b.jpg', None, 'HTTP 500：boom\x1f')
    book.save()
    wb = load_workbook(tmp_path / 'x.xlsx')
    rows = [list(r) for r in wb['识别结果'].iter_rows(values_only=True)]
    assert ['=1+1', 'okbell', None] in rows
    assert wb['识别结果'].cell(row=2, column=1).data_type == 's'
    assert [r[2] for r in wb['清单'].iter_rows(values_only=True)][1:] == ['成功', '失败']
