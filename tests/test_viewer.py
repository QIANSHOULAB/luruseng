import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bangong import engine, viewer  # noqa: E402
import mock_server  # noqa: E402
from test_engine import K1, make_images, server, reset  # noqa: E402,F401


def make_run(tmp_path, n=4, once=None, modes=None):
    make_images(tmp_path, n)
    mock_server.STATE['delay'] = 0
    if once:
        mock_server.STATE['once'] = once
    if modes:
        mock_server.STATE['modes'] = modes
    res = engine.Engine(str(tmp_path), [K1], log=lambda m: None).run()
    return res


def test_runs_photos_and_rows(tmp_path):
    res = make_run(tmp_path, 4, modes={K1: 'mismatch'})
    runs = viewer.list_runs(str(tmp_path))
    assert len(runs) == 1
    run_id, d, xlsx = runs[0]
    assert d == res.photos_dir and xlsx == res.xlsx
    photos = viewer.load_photos(d)
    assert sorted(photos) == [1, 2, 3, 4] and photos[1].endswith('01.jpg')
    rows = viewer.load_rows(xlsx)
    assert sorted(rows) == [1, 2, 3, 4]
    r = rows[1][0]
    assert r['values'][1] == 'VM-07' and r['values'][0] == '28-Oct' and r['values'][12] == '639/Lee Ah Kow'
    assert r['red'] == {8, 9, 10}                        # TOTAL / PASSED / FAILED 标红
    assert 'PASSED 9 + FAILED 0 ≠ TOTAL 10' in r['note']
    assert viewer.run_label(run_id, 4).endswith('这批（4 张）')


def test_newest_run_first(tmp_path):
    make_run(tmp_path, 2)
    import time
    time.sleep(1.1)
    make_images(tmp_path, 1, start=60)
    engine.Engine(str(tmp_path), [K1], log=lambda m: None).run()
    runs = viewer.list_runs(str(tmp_path))
    assert len(runs) == 2 and runs[0][0] > runs[1][0]
    assert sorted(viewer.load_photos(runs[0][1])) == [1]


def _tk_root():
    try:
        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        return root
    except Exception:
        pytest.skip('没有显示器，跳过窗口测试')


def test_window_flip_and_jump(tmp_path):
    make_run(tmp_path, 4, once={K1: ('two', 1)})
    root = _tk_root()
    try:
        v = viewer.Viewer(root, str(tmp_path))
        root.update()
        assert v.cur == 1 and v.title() == '核对 · 第 1 张 / 共 4 张'
        v.step(1)
        assert v.cur == 2
        v.step(-5)
        assert v.cur == 1                                   # 到头了就停在 1
        v.num.set('4')
        v.jump()
        assert v.cur == 4 and v.num.get() == '4'
        v.step(1)
        assert v.cur == 4                                   # 到尾了就停在最后一张
        v.num.set('9')
        v.jump()
        assert v.cur == 4 and '没有第 9 张' in v.info.get()
        text = v.panel.get('1.0', 'end')
        assert 'No. 4' in text and 'SETTER' in text
        two = [n for n, recs in v.rows.items() if len(recs) == 2]
        v.show(two[0])
        assert '第 2 条' in v.panel.get('1.0', 'end')
        root.update()
        assert v._tk is not None                            # 照片真的画出来了
        v.destroy()
    finally:
        root.destroy()


def test_window_without_runs(tmp_path):
    root = _tk_root()
    try:
        v = viewer.Viewer(root, str(tmp_path))
        root.update()
        assert v.title() == '核对' and not v.runs
        v.destroy()
    finally:
        root.destroy()


# ---------------------------------------------------------------- 核对里直接改格子，表跟着改（跟网页版 core.test.mjs 对着写）
def _fill(ws, coord):
    c = ws[coord]
    return (c.fill.fgColor.rgb or '')[-6:].upper() if c.fill is not None and c.fill.fill_type == 'solid' else ''


def test_edit_cell_alts_confirm_sum_and_typos(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    monkeypatch.setattr(engine, 'READS', 2)                 # 两遍对答案：才有「几遍认得不一样」
    res = make_run(tmp_path, 1, modes={K1: 'chaos'})        # 几遍认出的机台都不一样：红，提示里有候选
    r = viewer.load_rows(res.xlsx)[1][0]
    assert 1 in r['red'] and r['row'] == 4 and r['builtin']
    alts = r['alts'][1]
    assert len(alts) >= 2 and all(a.startswith('VM-0') for a in alts), alts
    pick = alts[-1]
    out = viewer.edit_cell(res.xlsx, 4, 1, pick)
    assert out['value'] == pick and not out['still'] and out['checked']
    assert out['data']['machine'] == pick and out['data']['date'].endswith('-10-28') and out['data']['mo'] == '82664580'
    ws = load_workbook(res.xlsx).worksheets[0]
    assert ws['C4'].value == pick and _fill(ws, 'C4') == 'E2EFDA' and ws['C4'].font.color.rgb.endswith('375623')
    assert ws['O4'].value == '✓ 人核过' and ws['O4'].font.color.rgb.endswith('375623')
    assert _fill(ws, 'D4') == '' and _fill(ws, 'I4') == 'FFE699'          # 没动的格子照旧（BTMP 还是黄的）
    r = viewer.load_rows(res.xlsx)[1][0]
    assert r['human'] == {1} and r['red'] == set() and r['note'] == viewer.HUMAN_NOTE

    import time
    time.sleep(1.1)                                             # 再干一批（批次号按秒）：数量对不上，三格都红
    make_images(tmp_path, 1, start=7)
    mock_server.STATE['modes'] = {K1: 'mismatch'}
    res2 = engine.Engine(str(tmp_path), [K1], log=lambda m: None).run()
    assert res2.xlsx != res.xlsx
    x = res2.xlsx
    r = viewer.load_rows(x)[1][0]
    assert r['red'] == {8, 9, 10}
    out = viewer.edit_cell(x, 4, 8)                             # TOTAL「对的」：值不动，PASSED、FAILED 还红
    assert out['value'] == '10' and not out['still'] and not out['checked']
    r = viewer.load_rows(x)[1][0]
    assert r['red'] == {9, 10} and r['human'] == {8} and r['note'] == 'PASSED 9 + FAILED 0 ≠ TOTAL 10'
    out = viewer.edit_cell(x, 4, 10, '1o')                      # 打错了：还红，告诉你为什么
    assert out['still'] and out['why'] == 'FAILED 不是数字：1o' and not out['checked']
    assert viewer.load_rows(x)[1][0]['red'] == {9, 10}
    out = viewer.edit_cell(x, 4, 10, ' 1 ')                     # 9 + 1 = 10：三格一起不红
    assert out['value'] == '1' and out['checked'] and out['data']['failed'] == '1'
    ws = load_workbook(x).worksheets[0]
    assert ws['L4'].value == 1 and isinstance(ws['L4'].value, int)
    assert [_fill(ws, c) for c in ('J4', 'K4', 'L4')] == ['E2EFDA', '', 'E2EFDA']
    assert ws['K4'].font.color is None or not str(ws['K4'].font.color.rgb).endswith('9C0006')
    assert ws['O4'].value == '✓ 人核过'


def test_edit_cell_types_and_rules():
    from datetime import date, datetime
    d = date(2026, 10, 31)
    assert viewer.typed_value('date', '1/11', today=d) == datetime(2026, 11, 1)
    assert viewer.typed_value('date', 'hello', today=d) == 'hello'
    assert viewer.typed_value('machine', ' vm 07 ') == 'VM07'
    assert viewer.typed_value('part_number', 'td–524613－5') == 'TD-524613-5'
    assert viewer.typed_value('description', 'em ø6') == 'EM Φ6'
    assert viewer.typed_value('total_screened', '1,200') == 1200
    assert viewer.typed_value('shift', '夜班') == 'N'
    assert viewer.typed_value('setter', '639 / Lee Ah Kow') == '639/Lee Ah Kow'
    assert viewer.typed_value('btmp', '   ') == ''
    assert viewer.typed_value('数量', '12', False, 3) == 12 and viewer.typed_value('数量', '1.5', False, 3) == 1.5
    assert viewer.typed_value('备注', '  ok ', False, None) == 'ok'
    assert viewer.typed_problem('failed', '1o') == 'FAILED 不是数字：1o'
    assert viewer.typed_problem('shift', 'X') == 'SHIFT 不是 D/N：X'
    assert viewer.typed_problem('setter', '6□9/Lee') == 'SETTER 有看不清的字'
    assert viewer.typed_problem('btmp', '') is None and viewer.typed_problem('passed', 5) is None
    assert viewer.note_owner('Part No 几遍认得不一样：B1', ['Part', 'Part No']) == 'Part No'
    assert viewer.alts_from_note('MACHINE 几遍认得不一样：VM-01 / 空 / VM-02；SHIFT 不是 D/N：X', 'MACHINE') == ['VM-01', 'VM-02']
    assert viewer.cat_of('识别结果_20261031_080000.xlsx') == '刀单'
    assert viewer.cat_of('识别结果_20261031_080000_质检单_181502.xlsx') == '质检单'
    assert viewer.run_day('20261031_080000') == d


def test_edit_cell_category_book_and_locked(tmp_path, monkeypatch):
    from bangong import tables
    from openpyxl import load_workbook
    p = tmp_path / '识别结果_20261030_120000_质检单.xlsx'
    tb = tables.TableBook(str(p), '质检单', ['Part', 'Part No', '数量'], numeric={'数量'})
    tb.add(1, [{'Part': 'A', 'Part No': 'B1', '数量': '5', 'unsure': ['Part'], 'disputed': {'Part No': ['B1', 'B7']}}])
    tb.save()
    r = viewer.load_rows(str(p))[1][0]
    assert r['red'] == {0, 1} and r['alts'] == {1: ['B1', 'B7']} and not r['builtin']
    out = viewer.edit_cell(str(p), 3, 0)                        # 「Part」核过：「Part No」那句提示不能被误删
    assert not out['checked'] and viewer.load_rows(str(p))[1][0]['note'] == 'Part No 几遍认得不一样：B1 / B7'
    out = viewer.edit_cell(str(p), 3, 2, '12')                  # 数字列打数字：存成数字
    assert load_workbook(str(p)).worksheets[0]['D3'].value == 12 and not out['checked']
    before = p.read_bytes()
    real = os.replace

    def busy(a, b):
        raise PermissionError('in use')
    monkeypatch.setattr(viewer.os, 'replace', busy)              # Excel 开着这本：Windows 不让换
    with pytest.raises(viewer.Locked):
        viewer.edit_cell(str(p), 3, 1, 'B7')
    monkeypatch.setattr(viewer.os, 'replace', real)
    assert p.read_bytes() == before and not os.path.exists(str(p) + '.edit.tmp')
    (tmp_path / ('~$' + p.name)).write_bytes(b'x')              # Excel 崩过留下的假 ~$ 小文件：不碍事
    out = viewer.edit_cell(str(p), 3, 1, 'B7')
    assert out['checked'] and out['data'] == {'Part': 'A', 'Part No': 'B7', '数量': '12'}
    assert viewer.load_rows(str(p))[1][0]['note'] == viewer.HUMAN_NOTE
    assert viewer.cat_of(str(p)) == '质检单'


def test_checked_rows_go_into_databases(tmp_path):
    from bangong import tables
    from bangong.kb import KnowledgeBase
    kb = KnowledgeBase(str(tmp_path / 'kb.sqlite'))
    rec = {f: '' for f in engine.FIELDS}
    rec.update(date='2026-10-28', machine='VM-17', part_number='TD-1', mo='82660000', t_no='T4', setter='639/Lee Ah Kow')
    assert kb.put_checked('核对 20261031_080000 #1.0', rec)
    assert kb.put_checked('核对 20261031_080000 #1.0', dict(rec, machine='VM-18'))   # 同一条再改：换掉上次那份
    import sqlite3
    from contextlib import closing

    def q(sql):
        with closing(sqlite3.connect(str(tmp_path / 'kb.sqlite'))) as c:
            return c.execute(sql).fetchall()
    assert q('SELECT machine, sheet FROM records') == [('VM-18', '核对')]
    assert ('639', 'Lee Ah Kow') in kb.load().roster()
    assert not kb.put_checked('核对 20261031_080000 #1.0', None)                    # 又改红了：拿掉
    assert q('SELECT COUNT(*) FROM records') == [(0,)]
    db = tables.TableDB(str(tmp_path / 'cat' / 'db.sqlite'))
    assert db.put_checked('s', {'Date': '28/10', 'M/C No.': 'VM-07', 'Remark': ''})
    assert db.put_checked('t', {'date': '29/10', 'M/C No': 'VM-08', 'Qty': '5'})   # 列名按写法对，多的接在后面
    assert db.columns() == ['Date', 'M/C No.', 'Remark', 'Qty'] and db.count() == 2
    assert db.load().records[1] == {'Date': '29/10', 'M/C No.': 'VM-08', 'Qty': '5'}
    assert not db.put_checked('t', None) and db.count() == 1


def test_window_edit_cell(tmp_path, monkeypatch):
    monkeypatch.setattr(engine, 'READS', 2)
    res = make_run(tmp_path, 2, modes={K1: 'chaos'})
    root = _tk_root()
    fed = []
    try:
        v = viewer.Viewer(root, str(tmp_path), on_checked=lambda *a: fed.append(a) or True, busy=lambda: '')
        root.update()
        assert v.cat == '刀单' and v.today is not None
        v.open_cell(0, 1)
        root.update()
        assert v.editing == (1, 0, 1) and v.editor is not None and v.editor.winfo_exists()
        assert v.editor_var.get() == v.rows[1][0]['values'][1]
        v.open_cell(0, 1)                                           # 再点一下：收起
        assert v.editing is None and v.editor is None
        v.open_cell(0, 2)
        root.update()
        assert v.editor.winfo_children()[0].entry.bind('<Escape>')       # Esc 只收起小框（Windows 上没焦点按不了键，直接调）
        v.close_cell()
        assert v.editing is None and v.editor is None and v.winfo_exists()
        v.step(1)
        v.step(-1)
        assert v.cur == 1                                           # ←→ 照样翻页
        pick = v.rows[1][0]['alts'][1][-1]
        v.open_cell(0, 1)
        v.save_cell(0, 1, pick)
        root.update()
        text = v.panel.get('1.0', 'end')
        assert pick in text and '✓ 人核过' in text and v.editing is None
        assert f'MACHINE 改成 {pick}' in v.info.get() and '进了样板资料' in v.info.get() and '表跟着改了' in v.info.get()
        cat, source, data = fed[-1]
        assert cat == '刀单' and source == f'核对 {v.run_id} #1.0' and data['machine'] == pick
        assert 'human' in v.panel.tag_names(f'{v.panel.tag_ranges("c0_1")[0]}')
        assert viewer.load_rows(res.xlsx)[1][0]['values'][1] == pick
        v.save_cell(0, 9, 'abc')                                    # PASSED 打成字母：还红，样板资料拿掉
        assert '还不对' in v.info.get() and fed[-1][2] is None
        real = os.replace
        monkeypatch.setattr(viewer.os, 'replace', lambda a, b: (_ for _ in ()).throw(PermissionError('in use')))
        v.save_cell(0, 9, '1')                                      # Excel 开着：一个字不动，告诉你关掉再点
        assert '表在 Excel 里开着' in v.info.get() and viewer.load_rows(res.xlsx)[1][0]['values'][9] == 'abc'
        monkeypatch.setattr(viewer.os, 'replace', real)
        v.busy = lambda: v.run_dir                                  # 这批还在干：先不让改
        v.open_cell(0, 9)
        assert v.editing is None and '还在干' in v.info.get()
        v.destroy()
    finally:
        root.destroy()


def test_app_opens_only_checker_then_feeds_checked_rows(tmp_path, monkeypatch):
    """干完只弹核对窗口（不再自己打开 Excel，要看表点「打开表」）；核完的那条自己进样板资料。"""
    import sqlite3
    import time
    root = _tk_root()
    from bangong import app
    monkeypatch.setattr(app, 'app_dir', lambda: str(tmp_path / 'app'))
    monkeypatch.setattr(app, 'CONFIG', str(tmp_path / 'config.json'))
    opened = []
    monkeypatch.setattr(app, 'open_path', lambda p: opened.append(p))
    monkeypatch.setattr(viewer, '_open', lambda p: opened.append(p))
    monkeypatch.setattr(engine, 'READS', 2)
    mock_server.STATE.update(delay=0.3, modes={K1: 'chaos'})
    a = None
    try:
        a = app.App(root)
        a.cfg['auto_open'] = True
        photos = tmp_path / '待识别'
        make_images(photos, 1)
        a.folder.set(str(photos))
        a.key1.set(K1)
        assert a.busy_dir() == ''
        a.start()
        end = time.time() + 10
        busy = ''
        while not busy and time.time() < end:                       # 正在干：这批的编号照片文件夹
            busy = a.busy_dir()
            time.sleep(0.02)
        assert busy and busy == a.eng.res.photos_dir and engine.DONE_DIR in busy
        end = time.time() + 30
        while (a.worker.is_alive() or a.viewer is None) and time.time() < end:
            root.update()
            time.sleep(0.05)
        root.update()
        assert a.viewer is not None and a.viewer.winfo_exists() and opened == []      # 只弹核对窗口
        assert a.busy_dir() == ''
        v = a.viewer
        pick = v.rows[1][0]['alts'][1][0]
        v.save_cell(0, 1, pick)
        assert '进了样板资料' in v.info.get(), v.info.get()
        db = sqlite3.connect(app.kb_path())
        try:
            assert db.execute("SELECT machine, sheet FROM records WHERE source LIKE '核对 %'").fetchall() == [(pick, '核对')]
        finally:
            db.close()
        v.open_table()
        assert opened == [v.xlsx]                                                     # 要看表：点「打开表」
        v.destroy()
    finally:
        if a is not None and a.worker is not None:
            a.stop.set()
            a.worker.join(30)
        root.destroy()


def test_window_btmp_copies_total(tmp_path):
    """换刀调试单上没有 BTMP：表里照抄 TOTAL SCREENED（公司那本的规矩），核对里不列这份抄的；改 TOTAL 它跟着变。"""
    from openpyxl import load_workbook
    res = make_run(tmp_path, 1, modes={K1: 'nobtmp'})
    r = viewer.load_rows(res.xlsx)[1][0]
    assert r['values'][viewer.BTMP] == '1' == r['values'][viewer.TOTAL] and r['labels'][viewer.BTMP] == 'BTMP'
    ws = load_workbook(res.xlsx).worksheets[0]
    assert ws['I4'].value == 1 and _fill(ws, 'I4') == 'FFE699'                 # 数字，还是黄的
    root = _tk_root()
    try:
        v = viewer.Viewer(root, str(tmp_path))
        root.update()
        text = v.panel.get('1.0', 'end')
        assert 'BTMP' not in text and 'TOTAL SCREENED' in text and 'SETTER' in text
        v.open_cell(0, 8)                                           # 后面的格子照样点得开、改得对
        assert v.editor_var.get() == '1'
        v.save_cell(0, 8, '3')                                      # 改 TOTAL：抄过去的 BTMP 跟着变
        r = viewer.load_rows(res.xlsx)[1][0]
        assert r['human'] == {8} and r['values'][viewer.BTMP] == '3' and 'BTMP' not in v.panel.get('1.0', 'end')
        out = viewer.edit_cell(res.xlsx, 4, 8, '3o')                # 打错了：抄过去的也清掉，别留个错数
        assert out['still'] and viewer.load_rows(res.xlsx)[1][0]['values'][viewer.BTMP] == ''
        v.destroy()
    finally:
        root.destroy()
    res2 = make_run(tmp_path / 'b', 1, modes={K1: 'mismatch'})      # 单子上真写了 BTMP、跟 TOTAL 不一样：照样列
    root = _tk_root()
    try:
        v = viewer.Viewer(root, str(tmp_path / 'b'))
        root.update()
        assert 'BTMP' in v.panel.get('1.0', 'end')
        v.save_cell(0, 8, '11')                                     # 真写的 BTMP 不跟着 TOTAL 变
        assert viewer.load_rows(res2.xlsx)[1][0]['values'][viewer.BTMP] == '1'
        v.destroy()
    finally:
        root.destroy()
