"""数据库：投喂公司 TOOL BROKEN RECORD（或核对过的录入僧结果表），以后认单子以它为准。

只认喂进来的表：录入僧自己认出来的东西不进库，免得认错的被学进去越学越歪。
"""
from __future__ import annotations

import hashlib
import os
import re
import sqlite3
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

from .engine import FIELDS

DB_FILE = '数据库.sqlite'
FEED_DIR = '投喂'
CAT_FIELDS = ['machine', 'part_number', 'description']          # 来回就那几个值的字段
LABEL = {'date': 'DATE', 'machine': 'MACHINE', 'part_number': 'PART NUMBER', 'mo': 'MO', 'opr_no': 'OPR NO.',
         't_no': 'T NO.', 'description': 'DESCRIPTION', 'btmp': 'BTMP', 'total_screened': 'TOTAL SCREENED',
         'passed': 'PASSED', 'failed': 'FAILED', 'shift': 'SHIFT', 'setter': 'SETTER'}
HEADER_KEYS = {
    'DATE': 'date', 'MACHINE': 'machine', 'M/C': 'machine', 'MACHINENO': 'machine',
    'PARTNUMBER': 'part_number', 'PARTNO': 'part_number', 'MO': 'mo', 'MONO': 'mo', 'W/O': 'mo',
    'OPRNO': 'opr_no', 'OPR': 'opr_no', 'OPNO': 'opr_no', 'TNO': 't_no', 'TOOLNO': 't_no',
    'DESCRIPTION': 'description', 'TOOLDESCRIPTION': 'description', 'BTMP': 'btmp',
    'TOTALSCREENED': 'total_screened', 'PASSED': 'passed', 'FAILED': 'failed', 'SHIFT': 'shift', 'SETTER': 'setter',
}
# 容易看错的字：两边都换成同一个再比，比如 VM-O7 和 VM-07 就是同一个
CONFUSE = str.maketrans({'O': '0', 'Q': '0', 'D': '0', 'I': '1', 'L': '1', '|': '1', 'S': '5', 'B': '8',
                         'Z': '2', 'G': '6'})
SCHEMA = """
CREATE TABLE IF NOT EXISTS records (
    hash TEXT PRIMARY KEY, source TEXT, sheet TEXT, fed_at TEXT,
    date TEXT, machine TEXT, part_number TEXT, mo TEXT, opr_no TEXT, t_no TEXT, description TEXT,
    btmp TEXT, total_screened TEXT, passed TEXT, failed TEXT, shift TEXT, setter TEXT);
CREATE TABLE IF NOT EXISTS people (no TEXT PRIMARY KEY, name TEXT, source TEXT);
CREATE TABLE IF NOT EXISTS sources (path TEXT PRIMARY KEY, mtime REAL, size INTEGER, fed_at TEXT, added INTEGER);
"""


def label_key(v) -> str:
    return re.sub(r'[^A-Z/]', '', str(v or '').upper())


def nkey(v: str) -> str:
    """比对用的写法：大写、去空格、各种横杠统一。"""
    t = str(v or '').upper()
    t = re.sub(r'[‐-―_~]', '-', t)
    t = t.replace(',', '.')
    return re.sub(r'\s+', '', t)


def ckey(v: str) -> str:
    return nkey(v).translate(CONFUSE)


def lev1(a: str, b: str) -> bool:
    """两个字符串是否只差一个字（替换、多一个或少一个）。"""
    if a == b or abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) == 1
    if len(a) > len(b):
        a, b = b, a
    for i in range(len(b)):
        if a == b[:i] + b[i + 1:]:
            return True
    return False


def cell_text(v) -> str:
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r'\s+', ' ', str(v)).strip()


def record_hash(rec: dict) -> str:
    return hashlib.sha1('|'.join(nkey(rec.get(f, '')) for f in FIELDS).encode('utf-8')).hexdigest()


SETTER_RE = re.compile(r'^\s*(\d{2,6})\s*[/／]\s*([A-Za-z][A-Za-z .\'-]*[A-Za-z])\s*$')


@dataclass
class FeedResult:
    path: str
    rows: int = 0
    added: int = 0
    people: int = 0
    sheets: list = field(default_factory=list)
    error: str = ''

    def line(self) -> str:
        name = os.path.basename(self.path)
        if self.error:
            return f'✗ {name}：{self.error}'
        return (f'✓ 吃进 {name}：读到 {self.rows} 行记录，新增 {self.added} 行（其余以前吃过）'
                + (f'，新认识 {self.people} 人' if self.people else '')
                + (f'（{"、".join(self.sheets)}）' if self.sheets else ''))


# ---------------------------------------------------------------- 读表
def find_header(rows: list[tuple]) -> tuple[int, dict[int, str]] | None:
    """在前 15 行找表头：同时有 DATE 和 MACHINE 的那一行；下一行是小表头（T NO.、TOTAL SCREENED…）。"""
    for i, row in enumerate(rows[:15]):
        keys = [label_key(v) for v in row]
        if 'DATE' in keys and ('MACHINE' in keys or 'M/C' in keys):
            sub = rows[i + 1] if i + 1 < len(rows) else ()
            cols: dict[int, str] = {}
            for c in range(max(len(row), len(sub))):
                for v in ((sub[c] if c < len(sub) else None), (row[c] if c < len(row) else None)):
                    f = HEADER_KEYS.get(label_key(v))
                    if f and f not in cols.values():
                        cols[c] = f
                        break
            uses_sub = any(c < len(sub) and HEADER_KEYS.get(label_key(sub[c])) for c in cols)
            return (i + 2 if uses_sub else i + 1), cols
    return None


def read_people(rows: list[tuple]) -> list[tuple[str, str]]:
    """名字页：认「535/RAJ KUMAR」这种格子，或者一格工号紧跟一格名字。"""
    out = []
    for row in rows:
        cells = [cell_text(v) for v in row]
        for j, t in enumerate(cells):
            m = SETTER_RE.match(t)
            if m:
                out.append((m.group(1), m.group(2).strip()))
                continue
            if re.fullmatch(r'\d{2,6}', t):
                nxt = next((x for x in cells[j + 1:] if x), '')
                if re.fullmatch(r"[A-Za-z][A-Za-z .'-]{1,60}", nxt) and len(re.sub(r'[^A-Za-z]', '', nxt)) >= 3:
                    out.append((t, nxt.strip()))
    return out


def read_workbook(path: str) -> tuple[list[tuple[str, dict]], list[tuple[str, str]], list[str]]:
    """返回 ([(页名, 记录)], [(工号, 名字)], [读到记录的页名])。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    recs, people, used, loose = [], [], [], []
    has_header = False
    try:
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            hdr = find_header(rows)
            if not hdr:
                found = read_people(rows)
                if re.search(r'NAME|EMP|名单|名字|人员', ws.title, re.I):
                    people += found
                else:
                    loose += found
                continue
            has_header = True
            start, cols = hdr
            n = 0
            for row in rows[start:]:
                rec = {f: '' for f in FIELDS}
                for c, f in cols.items():
                    if c < len(row):
                        rec[f] = cell_text(row[c])
                core = sum(bool(rec[f]) for f in ('machine', 'part_number', 'mo', 't_no', 'description', 'setter'))
                if core >= 3:
                    recs.append((ws.title, rec))
                    n += 1
                    m = SETTER_RE.match(rec['setter'])
                    if m:
                        people.append((m.group(1), m.group(2).strip()))
            if n:
                used.append(ws.title)
    finally:
        wb.close()
    if has_header:                       # 真是刀单本子，别的页上的工号名字才算数（别的单子不往刀单里串）
        people += loose
    return recs, people, used


# ---------------------------------------------------------------- 数据库
class KnowledgeBase:
    def __init__(self, path: str):
        self.path = path
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with self._conn() as c:
            c.executescript(SCHEMA)

    def _conn(self):
        return sqlite3.connect(self.path, timeout=30)

    def feed(self, path: str) -> FeedResult:
        res = FeedResult(path)
        if not path.lower().endswith(('.xlsx', '.xlsm')):
            res.error = '只认 .xlsx / .xlsm（老的 .xls 请先在 Excel 里另存为 .xlsx）'
            return res
        try:
            recs, people, used = read_workbook(path)
        except Exception as e:
            res.error = f'读不了：{type(e).__name__}：{e}'
            return res
        if not recs and not people:
            res.error = ('这本不像刀单：没找到 DATE、MACHINE 那样的表头，一条没吃。'
                         '是别的单子的话，放进它自己的类别')
            return res
        now = datetime.now().isoformat(timespec='seconds')
        src = os.path.basename(path)
        with self._conn() as c:
            before = c.execute('SELECT COUNT(*) FROM records').fetchone()[0]
            c.executemany(
                f'INSERT OR IGNORE INTO records (hash, source, sheet, fed_at, {", ".join(FIELDS)}) '
                f'VALUES ({", ".join("?" * (4 + len(FIELDS)))})',
                [(record_hash(r), src, sheet, now, *[r[f] for f in FIELDS]) for sheet, r in recs])
            res.added = c.execute('SELECT COUNT(*) FROM records').fetchone()[0] - before
            p0 = c.execute('SELECT COUNT(*) FROM people').fetchone()[0]
            c.executemany('INSERT OR IGNORE INTO people (no, name, source) VALUES (?, ?, ?)',
                          [(n, nm, src) for n, nm in people])
            res.people = c.execute('SELECT COUNT(*) FROM people').fetchone()[0] - p0
            st = os.stat(path)
            c.execute('INSERT OR REPLACE INTO sources (path, mtime, size, fed_at, added) VALUES (?, ?, ?, ?, ?)',
                      (os.path.abspath(path), st.st_mtime, st.st_size, now, res.added))
        res.rows, res.sheets = len(recs), used
        return res

    def feed_folder(self, folder: str) -> list[FeedResult]:
        """投喂文件夹：新放进来的、或者改过的表才吃。"""
        if not folder or not os.path.isdir(folder):
            return []
        with self._conn() as c:
            seen = {p: (m, s) for p, m, s in c.execute('SELECT path, mtime, size FROM sources')}
        out = []
        for name in sorted(os.listdir(folder)):
            p = os.path.abspath(os.path.join(folder, name))
            if name.startswith('~$') or not name.lower().endswith(('.xlsx', '.xlsm')) or not os.path.isfile(p):
                continue
            st = os.stat(p)
            if seen.get(p) == (st.st_mtime, st.st_size):
                continue
            out.append(self.feed(p))
        return out

    def summary(self) -> str:
        with self._conn() as c:
            n = c.execute('SELECT COUNT(*) FROM records').fetchone()[0]
            if not n:
                return '数据库还是空的：点「放入数据库」放公司的 TOOL BROKEN RECORD'
            m = c.execute('SELECT COUNT(DISTINCT machine), COUNT(DISTINCT part_number), COUNT(DISTINCT description) '
                          "FROM records").fetchone()
            p = c.execute('SELECT COUNT(*) FROM people').fetchone()[0]
        return f'数据库：{n} 条记录 · 机台 {m[0]} 台 · 零件 {m[1]} 种 · 刀具描述 {m[2]} 种 · 认识 {p} 人'

    def load(self) -> 'Snapshot':
        with self._conn() as c:
            rows = c.execute(f'SELECT {", ".join(FIELDS)} FROM records').fetchall()
            people = c.execute('SELECT no, name FROM people').fetchall()
        return Snapshot([dict(zip(FIELDS, r)) for r in rows], people)


# ---------------------------------------------------------------- 拿数据库当准绳
class Snapshot:
    def __init__(self, records: list[dict], people: list[tuple[str, str]]):
        self.n = len(records)
        self.people = people
        self.count = {f: Counter() for f in CAT_FIELDS}
        spell = {f: defaultdict(Counter) for f in CAT_FIELDS}
        for r in records:
            for f in CAT_FIELDS:
                if r[f]:
                    spell[f][nkey(r[f])][r[f]] += 1
        # 每个值取最常见的写法当标准写法
        self.canon = {f: {k: c.most_common(1)[0][0] for k, c in spell[f].items()} for f in CAT_FIELDS}
        for f in CAT_FIELDS:
            for k, c in spell[f].items():
                self.count[f][self.canon[f][k]] = sum(c.values())
        self.by_ckey = {f: defaultdict(set) for f in CAT_FIELDS}
        for f in CAT_FIELDS:
            for k, v in self.canon[f].items():
                self.by_ckey[f][ckey(k)].add(v)

    def roster(self) -> list[tuple[str, str]]:
        return list(self.people)

    def hints(self, max_parts: int = 200) -> str:
        if not self.n:
            return ''
        lines = []
        for f, lim in (('machine', 100), ('description', 150), ('part_number', max_parts)):
            vals = [v for v, _ in self.count[f].most_common(lim)]
            if vals:
                lines.append(f'{LABEL[f]}：' + '、'.join(vals))
        return ('公司记录里出现过的值（数据库 {} 条记录）：单子上的字跟其中某个很像时，一律以这里的写法为准；'
                '明显不一样的（新机台、新零件、新刀）照单子原样抄。\n'.format(self.n) + '\n'.join(lines))

    def _snap(self, f: str, v: str) -> str | None:
        k = nkey(v)
        cands = self.by_ckey[f].get(ckey(v), set())
        if len(cands) == 1:
            return next(iter(cands))
        if len(k) >= 4:
            near = {c for kk, c in self.canon[f].items() if lev1(kk, k)}
            if len(near) == 1:
                c = next(iter(near))
                if self.count[f][c] >= 2:
                    return c
        return None

    def fix(self, rec: dict, unsure=()) -> tuple[dict, dict, dict]:
        """返回 (改好的记录, {字段: (原来, 改成, 原因)}, {字段: 标红原因})。
        认出来的值跟库里某个值只差一个容易看错的字，就改成库里的写法。
        库里没见过的机台、刀具描述：模型认得清就不标，模型自己拿不准（unsure）才标红。
        不按零件对历史（Song 定的：零件、刀、MO 之间的老关系不拿来标红；没见过的零件也不标红）。"""
        rec = dict(rec)
        fixed, flags = {}, {}
        if not self.n:
            return rec, fixed, flags
        for f in ('machine', 'part_number', 'description'):
            v = rec.get(f, '')
            if not v:
                continue
            k = nkey(v)
            if k in self.canon[f]:
                rec[f] = self.canon[f][k]
                continue
            c = self._snap(f, v)
            if c:
                fixed[f] = (v, c, '按数据库改')
                rec[f] = c
            elif f != 'part_number' and f in unsure:
                flags[f] = f'{LABEL[f]} 数据库里没见过，模型也拿不准：{v}'
        return rec, fixed, flags
