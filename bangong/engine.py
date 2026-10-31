"""录入僧（原名搬工）— 文件夹里的照片按拍摄顺序编号，5 张一批送 Gemini，认出刀具破损记录（TOOL BROKEN RECORD），边干边写 Excel。"""
from __future__ import annotations

import base64
import difflib
import hashlib
import io
import json
import os
import re
import threading
import time
from collections import Counter
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone

from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps, ImageStat

try:
    import pillow_heif  # iPhone 的 HEIC
    pillow_heif.register_heif_opener()
except Exception:  # pragma: no cover
    pass

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VERSION = '1.10.0'
APP_NAME = '录入僧'
SUBTITLE = '施主去玩，贫僧填表'
TITLE = '入我千手院，凡事尽释手；千僧替牛马，施主自风流。'      # 窗口标题栏只写这一句（Song 定的）
OLD_NAME = '搬工'                   # 改名前的名字：旧账本、旧配置照样认
API_BASE = os.environ.get('BANGONG_API', 'https://generativelanguage.googleapis.com/v1beta').rstrip('/')
DEFAULT_MODEL = 'gemini-3.5-flash-lite'
MODELS = ['gemini-3.5-flash-lite', 'gemini-3.8-flash', 'gemini-2.5-flash-lite', 'gemini-2.5-flash', 'gemini-2.5-pro']
IMG_EXT = {'.jpg', '.jpeg', '.png', '.webp', '.heic', '.heif', '.bmp', '.tif', '.tiff'}
BATCH = 5
MAX_SIDE = 2048
JPEG_Q = 92
DONE_DIR = '已识别'
ORIG_DIR = '原图'
DUP_DIR = '重复'
NAMES_FILE = '名单.txt'
STAMP_SIDE = 2560                   # 编号照片最长边（够放大看手写）
OUT_DIR = '识别结果'
LEDGER = '.录入僧账本.json'
OLD_LEDGER = '.搬工账本.json'
MINUTE_WAIT = 60
TIMEOUT = 180
RETRY_WAITS = [10, 30, 60]          # 5xx / 断网重试间隔
MAX_MINUTE_WAITS = 30               # 同一张连续撞每分钟上限的次数上限
READS = 2                           # 每张认两遍对答案；两遍对不上再认第三遍，按多数定，还是对不上的格子标红

# 公司那本 TOOL BROKEN RECORD 的 13 列（顺序一格不差）
FIELDS = ['date', 'machine', 'part_number', 'mo', 'opr_no', 't_no', 'description',
          'btmp', 'total_screened', 'passed', 'failed', 'shift', 'setter']
NUM_FIELDS = {'btmp', 'total_screened', 'passed', 'failed'}
NEVER_RED = {'date'}                # 日期错了无所谓，不标红

PROMPT = """你是工厂的刀具破损记录（TOOL BROKEN RECORD）录入员。从图片里找出每一条刀具破损记录，每条记录填下面 13 个字段。
图片可能是一张单子（一条记录），也可能是一页登记表（好几行，每行一条记录）。
最常见的是「TOOL REQUEST FORM 换刀调试单」，一张就是一条记录，各格对应：
  M/C No. 机器号码 → machine；Part Number 产品型号 → part_number；MO Number → mo；Opr No. 步骤 → opr_no；Date & Time 日期时间 → date；
  Shift 班次：只看手写的记号。Day 或 Night 被圈住、打了勾，或者旁边手写了「白」「早」「D」→ D；手写了「夜」「晚」「N」→ N。Day、Night 这两个词是印好的，没有任何手写记号就填空字符串，不要猜；
  Tool Information 刀具信息：手写的一串。所有 T 加数字的都是刀号 → t_no（写了几把刀，比如 T6.T7，就填 "T6, T7"，一把都不能漏），刀号不要写进 description；剩下的原样抄进 description，例如 Φ4 FMM X2、FACEMILL Φ40 R0.8。直径符号 Φ（手写常是一个圆圈中间一竖，也可能写成 ø）一定照写成 Φ，不要漏；印好的「磨损 Wear / 折断 Broken」被圈哪个都不用管；
  Total Quantity Screened 筛选数量 → total_screened；Quantity Passed 合格数量 → passed；Quantity Failed 不合格数量 → failed；
  Setter Emp No. 调试员工号 → setter（通常只写了工号）；单子上没有 BTMP，btmp 填空；QC 签名和 OK/NG 不要。

字段（括号里是单子上可能出现的写法和真实例子）：
- date：日期（DATE）。照单子上写的抄，例如 28/10、28-10-26、1-Oct。新加坡写法，日在前、月在后。
- machine：机台（MACHINE、M/C）。例如 VM-07、VM-14、4X-9、4X-16、5X-3、HM-3。
- part_number：零件号（PART NUMBER、PART NO.）。例如 TD-524613-5、AB-5313706C、AB-144125AE、TD-HAR8308G-05、TD-722-771-00-C。零件号中间没有空格。
- mo：工单号（MO、W/O）。通常是 8 位数字，例如 82664580、82665285。
- opr_no：工序号（OPR NO.、OP）。例如 110、130、150、170。
- t_no：刀号（T NO.、TOOL NO.）。例如 T4、T9、T58、T343。
- description：刀具描述（DESCRIPTION、Tool Information）。例如 EM2.5、EM6、BU10、SD3、DR09、FC6.4、Φ1.5 FEM、FACEMILL Φ40 R0.8。
- btmp：BTMP，数字。
- total_screened：检查总数（TOTAL SCREENED），数字。
- passed：合格数（PASSED），数字。
- failed：不合格数（FAILED），数字。
- shift：班次（SHIFT），只有 D（白班）或 N（夜班）；单子上没标就填空字符串。
- setter：设定员（SETTER、Setter Emp No.）。照单子写的抄：只写了工号就只填工号（例如 232），写了「工号/名字」就照抄（例如 639/Lee Ah Kow）。

规则：
- 照单子上写的抄，不猜、不补、不纠错；某个字实在看不清就用 □ 代替那个字。
- 单子上没写的字段填空字符串。
- 印刷的栏目名（DATE、MACHINE 之类）不是内容，不要抄进值里。
- 分清 0 和 O、1 和 I、5 和 S、8 和 B：数字字段只会是数字；机台和零件号照上面例子的样子判断。
- 被划掉、涂掉作废的内容不要。
- 图片歪斜或横着拍的，按正确方向读。
- unsure：每条记录里你拿不准的字段名（字迹潦草、被挡住、反光、可能看错的），只列真拿不准的；全都看得清就给空数组。
- 一条记录都没有就输出空数组。
- title_side：单子的标题（最上面那行印刷的大字，比如 TOOL REQUEST FORM 换刀调试单、TOOL BROKEN RECORD）在这张照片的哪一边，照片原样看，不要先在脑子里转正：top（上边）、right（右边）、bottom（下边）、left（左边）。
- 只输出 JSON：{"title_side": "top", "records": [{"date": "", "machine": "", ..., "unsure": []}]}，不要任何解释。"""

SCHEMA = {
    'type': 'OBJECT',
    'properties': {'title_side': {'type': 'STRING'}, 'records': {'type': 'ARRAY', 'items': {
        'type': 'OBJECT',
        'properties': {**{f: {'type': 'STRING'} for f in FIELDS},
                       'unsure': {'type': 'ARRAY', 'items': {'type': 'STRING'}}},
        'required': FIELDS + ['unsure'],
    }}},
    'required': ['records', 'title_side'],
}

# ---------------------------------------------------------------- 错误
class ApiError(Exception):
    def __init__(self, status: int, message: str, detail: str = ''):
        super().__init__(message)
        self.status, self.message, self.detail = status, message, detail

    def raw(self) -> str:
        return f'HTTP {self.status}：{self.message}' if self.status else self.message


class Fatal(Exception):
    """整轮停工（钥匙无效、模型不存在、预付余额用完、全部钥匙都不能用）。"""


class Stopped(Exception):
    pass


# ---------------------------------------------------------------- 时间
def _nth_sunday(year: int, month: int, n: int) -> datetime:
    d = datetime(year, month, 1)
    d += timedelta(days=(6 - d.weekday()) % 7)
    return d + timedelta(weeks=n - 1)


def pacific_offset(utc: datetime) -> timedelta:
    """美国太平洋时间偏移（夏令时 3 月第二个周日 2:00 至 11 月第一个周日 2:00）。不依赖 tzdata。"""
    y = utc.year
    start = _nth_sunday(y, 3, 2).replace(hour=10)   # 2:00 PST = 10:00 UTC
    end = _nth_sunday(y, 11, 1).replace(hour=9)     # 2:00 PDT = 09:00 UTC
    naive = utc.replace(tzinfo=None)
    return timedelta(hours=-7) if start <= naive < end else timedelta(hours=-8)


def next_pacific_midnight(now: float) -> float:
    utc = datetime.fromtimestamp(now, timezone.utc)
    off = pacific_offset(utc)
    local = (utc + off).replace(tzinfo=None)
    midnight = datetime(local.year, local.month, local.day) + timedelta(days=1)
    # 午夜那一刻的偏移（跨夏令时切换日也对）
    guess = midnight - off
    off2 = pacific_offset(guess.replace(tzinfo=timezone.utc))
    return (midnight - off2).replace(tzinfo=timezone.utc).timestamp() + 30   # 多等 30 秒，免得卡在边上


# ---------------------------------------------------------------- 钥匙
class Keys:
    def __init__(self, keys: list[str], clock):
        self.keys = [k.strip() for k in keys if k and k.strip()]
        self.dead_until = [0.0] * len(self.keys)
        self.used = [0] * len(self.keys)
        self.lock = threading.Lock()
        self.clock = clock

    def pick(self):
        """返回 (序号, 钥匙)；都在冷却就返回 (None, 最早恢复时间)。"""
        with self.lock:
            now = self.clock()
            for i, k in enumerate(self.keys):
                if self.dead_until[i] <= now:
                    return i, k
            return None, min(self.dead_until)

    def park(self, i: int, until: float):
        with self.lock:
            self.dead_until[i] = max(self.dead_until[i], until)


# ---------------------------------------------------------------- 图片
def file_hash(path: str) -> str:
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def _load_rgb(path: str) -> Image.Image:
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        if im.mode in ('RGBA', 'LA', 'P'):
            im = im.convert('RGBA')
            bg = Image.new('RGB', im.size, 'white')
            bg.paste(im, mask=im.split()[-1])
            return bg
        return im.convert('RGB') if im.mode != 'RGB' else im.copy()


def _shrink(im: Image.Image, side: int) -> Image.Image:
    w, h = im.size
    if max(w, h) <= side:
        return im
    s = side / max(w, h)
    return im.resize((max(1, round(w * s)), max(1, round(h * s))), Image.LANCZOS)


def prepare_image(path: str, enhance: bool = False) -> bytes:
    """enhance=True 是第二遍用的：拉一下对比度、锐化一点，让两遍看到的不完全一样，错不到一块去。"""
    im = _shrink(_load_rgb(path), MAX_SIDE)
    if enhance:
        im = ImageOps.autocontrast(im, cutoff=1).filter(ImageFilter.UnsharpMask(radius=2, percent=80, threshold=3))
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=JPEG_Q, optimize=True)
    return buf.getvalue()


def _font(size: int):
    for name in ('arialbd.ttf', 'Arial Bold.ttf', 'DejaVuSans-Bold.ttf', 'LiberationSans-Bold.ttf'):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # pragma: no cover
        return ImageFont.load_default()


TURN = {'right': Image.Transpose.ROTATE_90, 'bottom': Image.Transpose.ROTATE_180, 'left': Image.Transpose.ROTATE_270}


def upright(im: Image.Image, side: str) -> Image.Image:
    """标题在右边就逆时针转 90°，在左边就顺时针转 90°，在下边就转 180°：转完标题在上，字是正的。"""
    t = TURN.get(side or '')
    return im.transpose(t) if t is not None else im


def stamp_number(src: str, no: int, out: str, side: str = '') -> str:
    """在照片最空的那个角写上大红编号（白底红框），给人核对用；四个角都有字就在顶上加一条白边写。
    先按单子标题的位置把照片转正（人看着舒服）。只给人看：送去识别的是原图，编号不会被当成内容。"""
    im = upright(_shrink(_load_rgb(src), STAMP_SIDE), side)
    w, h = im.size
    short = min(w, h)
    text = str(no)
    font = _font(max(28, int(short * 0.11)))
    l, t, r, b = ImageDraw.Draw(im).textbbox((0, 0), text, font=font)
    tw, th = r - l, b - t
    pad = max(8, int(th * 0.35))
    bw, bh = tw + 2 * pad, th + 2 * pad
    m = max(6, int(short * 0.02))
    gray = im.convert('L')
    spots = [(w - bw - m, m), (m, m), (w - bw - m, h - bh - m), (m, h - bh - m)]   # 右上、左上、右下、左下，一样空就按这个顺序
    scored = []
    for i, (x, y) in enumerate(spots):
        if x < 0 or y < 0:
            continue
        st = ImageStat.Stat(gray.crop((x, y, x + bw, y + bh)))
        scored.append((round(st.stddev[0]), i, (x, y)))
    best = min(scored) if scored else (999, 0, (m, m))
    x, y = best[2]
    if best[0] > 30:                      # 四个角都不空：顶上加一条白边
        band = bh + 2 * m
        canvas = Image.new('RGB', (w, h + band), 'white')
        canvas.paste(im, (0, band))
        im, x, y = canvas, m, m
    d = ImageDraw.Draw(im)
    d.rounded_rectangle([x, y, x + bw, y + bh], radius=max(4, pad // 2), fill='white',
                        outline=(200, 0, 0), width=max(3, short // 250))
    d.text((x + pad - l, y + pad - t), text, font=font, fill=(200, 0, 0))
    im.save(out, 'JPEG', quality=88)
    return out


# ---------------------------------------------------------------- 名单
def load_names(path: str | None) -> list[tuple[str, str]]:
    """名单.txt：一行一个「工号/名字」，# 开头是注释。"""
    out = []
    if not path or not os.path.isfile(path):
        return out
    from .sheets import _decode                       # 记事本存的 GBK（ANSI）、UTF-8、UTF-16 都认
    with open(path, 'rb') as fh:
        text = _decode(fh.read())
    for line in text.splitlines():
        t = line.strip()
        if not t or t.startswith('#'):
            continue
        mm = re.match(r'^(\d+)\s*[/／\s]\s*(.+?)$', t)
        if mm:
            out.append((mm.group(1), re.sub(r'\s+', ' ', mm.group(2))))
    return out


def _sim(a: str, b: str) -> float:
    a = re.sub(r'[^A-Z\u4e00-\u9fff]', '', a.upper())
    b = re.sub(r'[^A-Z\u4e00-\u9fff]', '', b.upper())
    if not a or not b:
        return 0.0
    return difflib.SequenceMatcher(None, a, b).ratio()


def match_setter(raw: str, names: list[tuple[str, str]]) -> tuple[str, str, bool]:
    """按名单把 SETTER 补全/纠正成「工号/名字」。返回 (值, 提示, 要不要标红)。"""
    t = (raw or '').strip()
    if not names or not t:
        return raw, '', False
    mm = re.match(r'^\s*([0-9□]+)\s*[/／\-\s]?\s*(.*)$', t)
    num, nm = (mm.group(1), mm.group(2).strip()) if mm else ('', t)
    by_num = dict(names)
    if num.isdigit() and num in by_num:
        canon = f'{num}/{by_num[num]}'
        if not nm or _sim(nm, by_num[num]) >= 0.5:
            return canon, '', False
        return raw, f'SETTER 工号 {num} 在名单上是 {by_num[num]}，单子上名字是「{nm}」', True
    best = max(names, key=lambda x: _sim(nm, x[1]))
    if nm and _sim(nm, best[1]) >= 0.75:
        canon = f'{best[0]}/{best[1]}'
        if num == best[0]:
            return canon, '', False
        if '□' in num and re.fullmatch(num.replace('□', r'\d'), best[0]):      # 工号有个字看不清，名字对得上：按名单补
            return canon, f'SETTER 看不清的工号按名单补成 {best[0]}', False
        return canon, f'SETTER 按名字对上名单，工号补成 {best[0]}（单子上：{t}）', True
    return raw, 'SETTER 不在名单里', True


# ---------------------------------------------------------------- API
def model_name(m: str) -> str:
    return re.sub(r'^models/', '', (m or '').strip()) or DEFAULT_MODEL


def build_prompt(names: list[tuple[str, str]] | None = None, hints: str = '') -> str:
    out = PROMPT
    if names:
        roster = '\n'.join(f'{n}/{nm}' for n, nm in names)
        out += ('\n\n设定员名单（只用来帮你认潦草的工号和名字）：SETTER 还是照单子上写的抄，'
                '只写了工号就只填工号，不要补名字。\n' + roster)
    if hints:
        out += '\n\n' + hints
    return out


def build_body(jpeg: bytes, model: str, schema: bool = True, names: list[tuple[str, str]] | None = None,
               hints: str = '') -> dict:
    gen = {'responseMimeType': 'application/json'}
    if schema:
        gen['responseSchema'] = SCHEMA
    if re.match(r'^gemini-[12]\.', model):
        gen['temperature'] = 0.1          # Gemini 3 起官方已废弃采样参数（忽略，以后会报 400），只给老模型发
    if re.match(r'^gemini-2\.5-flash(?!-lite)', model):
        gen['thinkingConfig'] = {'thinkingBudget': 0}
    return {
        'systemInstruction': {'parts': [{'text': build_prompt(names, hints)}]},
        'contents': [{'role': 'user', 'parts': [
            {'text': '把这张图片里的刀具破损记录逐条填好。'},
            {'inline_data': {'mime_type': 'image/jpeg', 'data': base64.b64encode(jpeg).decode()}},
        ]}],
        'generationConfig': gen,
    }


def call_api(model: str, key: str, body: dict) -> dict:
    """Key 只走 URL 参数 ?key=。报错原样带回 HTTP 状态码和谷歌的 error.message。"""
    url = f'{API_BASE}/models/{urllib.parse.quote(model)}:generateContent?key={urllib.parse.quote(key)}'
    req = urllib.request.Request(url, data=json.dumps(body).encode(), headers={'Content-Type': 'application/json'}, method='POST')
    try:
        with urllib.request.urlopen(req, timeout=TIMEOUT) as r:
            raw = r.read().decode('utf-8', 'replace')
    except urllib.error.HTTPError as e:
        raw = e.read().decode('utf-8', 'replace')
        msg, detail = raw[:500] or (e.reason or ''), ''
        try:
            err = json.loads(raw).get('error', {})
            msg = err.get('message') or msg
            detail = json.dumps(err.get('details', ''), ensure_ascii=False)
        except Exception:
            pass
        raise ApiError(e.code, msg, detail) from None
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        raise ApiError(0, f'网络错误：{getattr(e, "reason", e)}') from None
    try:
        return json.loads(raw)
    except Exception:
        raise ApiError(0, '返回的不是 JSON：' + raw[:300]) from None


ALIASES = {
    'date': 'date', '日期': 'date',
    'machine': 'machine', 'm_c': 'machine', 'mc': 'machine', 'machine_no': 'machine',
    'part_number': 'part_number', 'part_no': 'part_number', 'partno': 'part_number', 'part': 'part_number',
    'mo': 'mo', 'mo_no': 'mo', 'w_o': 'mo', 'wo': 'mo',
    'opr_no': 'opr_no', 'opr': 'opr_no', 'op': 'opr_no', 'opr_number': 'opr_no',
    't_no': 't_no', 'tno': 't_no', 'tool_no': 't_no',
    'description': 'description', 'tool_description': 'description', 'desc': 'description',
    'btmp': 'btmp',
    'total_screened': 'total_screened', 'total': 'total_screened', 'screened': 'total_screened',
    'passed': 'passed', 'pass': 'passed', 'failed': 'failed', 'fail': 'failed',
    'shift': 'shift', 'setter': 'setter',
}


def norm_key(k) -> str:
    k = re.sub(r'[\s./\-]+', '_', str(k).strip().lower()).strip('_')
    return ALIASES.get(k, k)


def response_json(data: dict):
    cands = data.get('candidates') or []
    if not cands:
        br = (data.get('promptFeedback') or {}).get('blockReason')
        raise ValueError(f'promptFeedback.blockReason = {br}' if br else '返回里没有 candidates')
    c = cands[0]
    text = ''.join(p.get('text', '') for p in (c.get('content') or {}).get('parts', []) if not p.get('thought'))
    finish = c.get('finishReason', '')
    t = re.sub(r'^```(?:json)?\s*|\s*```$', '', text.strip())
    try:
        return json.loads(t)
    except Exception:
        raise ValueError(f'finishReason = {finish}，JSON 解析不了：{text[:200]}')


SIDES = {'top': 'top', 'up': 'top', 'upper': 'top', '上': 'top', '上边': 'top', '上面': 'top', '顶': 'top', '顶部': 'top',
         'right': 'right', '右': 'right', '右边': 'right', '右侧': 'right',
         'bottom': 'bottom', 'down': 'bottom', 'lower': 'bottom', '下': 'bottom', '下边': 'bottom', '下面': 'bottom', '底': 'bottom', '底部': 'bottom',
         'left': 'left', '左': 'left', '左边': 'left', '左侧': 'left'}


def norm_side(v) -> str:
    return SIDES.get(str(v or '').strip().lower(), '')


def parse_records(data: dict) -> list[dict]:
    return parse_response(data)[0]


def parse_response(data: dict) -> tuple[list[dict], str]:
    """返回 (记录, 标题在照片哪一边)。"""
    obj = response_json(data)
    side = norm_side(obj.get('title_side', '')) if isinstance(obj, dict) else ''
    recs = obj
    if isinstance(obj, dict):
        recs = obj.get('records', obj.get('rows', obj))
    if isinstance(recs, dict):
        recs = [recs]
    if not isinstance(recs, list):
        raise ValueError('JSON 里没有 records：' + json.dumps(obj, ensure_ascii=False)[:200])
    out = []
    for r in recs:
        rec = {f: '' for f in FIELDS}
        unsure = set()
        if isinstance(r, dict):
            for k, v in r.items():
                nk = norm_key(k)
                if nk == 'unsure':                  # 模型自己说拿不准的字段
                    items = v if isinstance(v, list) else re.split(r'[,，、;；\s]+', str(v or ''))
                    unsure = {norm_key(x) for x in items if x} & set(FIELDS)
                elif nk in rec and v is not None:
                    rec[nk] = re.sub(r'\s+', ' ', str(v)).strip()
        elif isinstance(r, list):
            for f, v in zip(FIELDS, r):
                rec[f] = '' if v is None else re.sub(r'\s+', ' ', str(v)).strip()
        if any(rec[f] for f in FIELDS):
            rec['unsure'] = sorted(unsure)
            out.append(rec)
    return out, side


MONTHS = {'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6, 'jul': 7, 'aug': 8,
          'sep': 9, 'sept': 9, 'oct': 10, 'nov': 11, 'dec': 12}


def _month(name: str):
    name = name.lower().rstrip('.')
    return MONTHS.get(name[:4]) or MONTHS.get(name[:3])


def parse_date(s: str, today: date | None = None) -> date | None:
    """新加坡写法：日在前。年份一律按今年（Song 定的：都是最近的记录，认成 1998 之类也改成今年，不往去年推）。"""
    today = today or date.today()
    t = (s or '').strip().lower()
    if not t:
        return None
    m = d = None
    if mm := re.fullmatch(r'(\d{4})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{1,2})', t):
        m, d = int(mm.group(2)), int(mm.group(3))
    elif mm := re.fullmatch(r'(\d{1,2})\s*[-/. ]?\s*([a-z]{3,9})\.?(?:\s*[-/., ]\s*(\d{2,4}))?', t):
        d, m = int(mm.group(1)), _month(mm.group(2))
    elif mm := re.fullmatch(r'([a-z]{3,9})\.?\s*(\d{1,2})(?:\s*,?\s*(\d{2,4}))?', t):
        m, d = _month(mm.group(1)), int(mm.group(2))
    elif mm := re.fullmatch(r'(\d{1,2})\s*[-/.]\s*(\d{1,2})(?:\s*[-/.]\s*(\d{2,4}))?', t):
        d, m = int(mm.group(1)), int(mm.group(2))
    if not m or not d:
        return None
    try:
        return date(today.year, m, d)
    except ValueError:
        return None


def to_int(s: str):
    t = (s or '').strip().replace(',', '')
    return int(t) if re.fullmatch(r'\d{1,12}', t) else None


SHIFT_WORDS = {'DAY': 'D', 'NIGHT': 'N', 'D/S': 'D', 'N/S': 'N', 'DS': 'D', 'NS': 'N', 'DAY SHIFT': 'D', 'NIGHT SHIFT': 'N',
               '白': 'D', '白班': 'D', '早': 'D', '早班': 'D', '日': 'D', '日班': 'D', '夜': 'N', '夜班': 'N', '晚': 'N', '晚班': 'N'}


def norm_shift(s: str) -> str:
    t = re.sub(r'\s+', ' ', (s or '').strip().upper())
    return SHIFT_WORDS.get(t, t)


# ---------------------------------------------------------------- 认完先整理格式（不改内容）
DASHES = re.compile(r'[‐‑‒–—―－ー一]')
DIAMETER = re.compile(r'[ØøΦφϕ⌀Фф∅]')
T_LEAD = re.compile(r'^\s*((?:T\s*\d+\s*[.,，、;；/&+]*\s*)+)(.*)$', re.I)


def t_numbers(s: str) -> list[str]:
    return ['T' + n for n in re.findall(r'T\s*(\d+)', s or '', re.I)]


def clean_record(rec: dict) -> dict:
    """只整理写法，不动内容：零件号/机台去空格、统一横杠；描述开头混进去的刀号挪回 T NO.；
    Φ 的各种写法统一成 Φ；班次的「白/夜/Day/Night」统一成 D/N。"""
    r = dict(rec)
    for f in ('part_number', 'machine'):
        r[f] = DASHES.sub('-', re.sub(r'\s+', '', r.get(f, ''))).upper()
    for f in ('mo', 'opr_no', 'btmp', 'total_screened', 'passed', 'failed'):
        r[f] = re.sub(r'\s+', '', r.get(f, ''))
    desc = DIAMETER.sub('Φ', r.get('description', ''))
    tn = t_numbers(r.get('t_no', ''))
    m = T_LEAD.match(desc)
    if m and t_numbers(m.group(1)):                 # 「T44 M3 TAP X2」：开头的 T44 是刀号，不是描述
        tn += [t for t in t_numbers(m.group(1)) if t not in tn]
        desc = m.group(2)
    desc = re.sub(r'^[\s.,，、:;；\-]+', '', desc)
    desc = re.sub(r'\s*Φ\s*', ' Φ', desc).strip()     # 「Φ 4」「Φ4」都写成「Φ4」，前面空一格
    r['description'] = re.sub(r'\s+', ' ', desc).upper()
    if tn:
        r['t_no'] = ', '.join(dict.fromkeys(tn))
    else:
        r['t_no'] = r.get('t_no', '').strip().upper()
    if r.get('shift'):
        r['shift'] = norm_shift(r['shift'])
    r['setter'] = re.sub(r'\s*/\s*', '/', r.get('setter', '').strip())
    return r


def vote_key(f: str, v: str) -> str:
    """两遍对答案时怎么算「一样」：大小写、空格不算；日期按日/月比；数字按数值比。"""
    t = (v or '').strip()
    if f == 'date':
        d = parse_date(t, date(2000, 1, 1))
        return f'{d.day}/{d.month}' if d else re.sub(r'\s+', '', t)
    if f in NUM_FIELDS or f in ('mo', 'opr_no'):
        n = to_int(t)
        return str(n) if n is not None else re.sub(r'\s+', '', t).upper()
    if f == 'shift':
        return norm_shift(t)
    return re.sub(r'\s+', '', t).upper()


def records_agree(a: list[dict], b: list[dict]) -> bool:
    return len(a) == len(b) and all(vote_key(f, x.get(f, '')) == vote_key(f, y.get(f, ''))
                                    for x, y in zip(a, b) for f in FIELDS)


def vote(reads: list[list[dict]]) -> list[dict]:
    """几遍的结果按多数定：条数先按多数定，再一格一格按多数定（打平就用第一遍的）。
    几遍认得不一样的格子记进 disputed，写表时标红。"""
    if len(reads) == 1:
        return [dict(r, disputed={}) for r in reads[0]]
    counts = Counter(len(r) for r in reads)
    target, n = counts.most_common(1)[0]
    if n == 1:
        target = len(reads[0])
    cand = [r for r in reads if len(r) == target]
    no_majority = n == 1
    out = []
    for i in range(target):
        rec, disputed = {}, {}
        for f in FIELDS:
            vals = [r[i].get(f, '') for r in cand]
            keys = [vote_key(f, v) for v in vals]
            c = Counter(keys)
            top = c.most_common(1)[0][1]
            best = keys[0] if c[keys[0]] == top else c.most_common(1)[0][0]
            rec[f] = vals[keys.index(best)]
            if len(c) > 1 or no_majority:
                disputed[f] = list(dict.fromkeys(v or '空' for v in vals))
        rec['unsure'] = sorted(set().union(*(set(r[i].get('unsure') or ()) for r in cand)))
        rec['disputed'] = disputed
        if len(cand) < len(reads):
            rec['counts'] = [len(r) for r in reads]
        out.append(rec)
    return out


def retry_delay(detail: str) -> float | None:
    """谷歌 429 里带的 RetryInfo.retryDelay（比如 "37s"）：照它说的等，不用死等满 60 秒。"""
    m = re.search(r'"retryDelay"\s*:\s*"(\d+(?:\.\d+)?)s"', detail or '')
    return float(m.group(1)) if m else None


def classify_429(e: ApiError) -> str:
    s = f'{e.message} {e.detail}'
    if re.search(r'prepay', s, re.I):
        return 'prepay'
    if re.search(r'PerDay|per day|daily', s, re.I):
        return 'day'
    return 'minute'


# ---------------------------------------------------------------- Excel
BLACK = Side(style='thin', color='808080')
BOX = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
YELLOW = PatternFill('solid', fgColor='FFE699')
RED_FILL = PatternFill('solid', fgColor='FFC7CE')
RED_FONT = Font(name='Calibri', size=11, color='9C0006')
GREY_FILL = PatternFill('solid', fgColor='EDEDED')
ORANGE_FILL = PatternFill('solid', fgColor='FCE4D6')      # 按数据库改过的格子
ORANGE_FONT = Font(name='Calibri', size=11, color='833C0B')
BASE_FONT = Font(name='Calibri', size=11)
BOLD = Font(name='Calibri', size=11, bold=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center')

# A = 照片编号；B..N = 公司那本的 13 列；O = 核对提示（复制 B..N 贴进公司表即可）
COL_NO, COL_FIRST, COL_NOTE = 1, 2, 15
WIDTHS = [6, 9, 10, 22, 11, 9, 7, 13, 7, 16, 9, 9, 7, 28, 40]
LEFT_FIELDS = {'part_number', 'mo', 'opr_no', 'setter'}
MONTH_TAB = ['', 'JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN', 'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC']


def put(ws, row: int, col: int, value):
    """写一个格子：去掉 Excel 不认的控制字符；= 开头的照文字存，不当公式。"""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub('', value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'
    return cell


class RecordBook:
    """长得跟公司那本 TOOL BROKEN RECORD 一样。边干边写：每批干完就存一次盘。"""

    def __init__(self, path: str, today: date | None = None, names: list[tuple[str, str]] | None = None, kb=None):
        self.path = path
        self.today = today
        self.names = names or []
        self.kb = kb
        self.wb = Workbook()
        ws = self.ws = self.wb.active
        ws.title = 'RECORD'
        self.named = False
        self.n_rows = 0
        self.photos = 0
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=14)
        t = put(ws, 1, 1, 'TOOL BROKEN RECORD')
        t.font = Font(name='Calibri', size=26)
        t.alignment = CENTER
        ws.row_dimensions[1].height = 40
        single = [(1, 'No.'), (2, 'DATE'), (3, 'MACHINE'), (4, 'PART NUMBER'), (5, 'MO'), (6, 'OPR NO.'),
                  (9, 'BTMP'), (13, 'SHIFT'), (14, 'SETTER'), (15, '核对提示（不用贴进公司表）')]
        for col, name in single:
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
            put(ws, 2, col, name)
        ws.merge_cells(start_row=2, start_column=7, end_row=2, end_column=8)
        put(ws, 2, 7, 'TOOL')
        put(ws, 3, 7, 'T NO.')
        put(ws, 3, 8, 'DESCRIPTION')
        ws.merge_cells(start_row=2, start_column=10, end_row=2, end_column=12)
        put(ws, 2, 10, 'CONTAINMENT/SCREENING DETAIL')
        put(ws, 3, 10, 'TOTAL SCREENED')
        put(ws, 3, 11, 'PASSED')
        put(ws, 3, 12, 'FAILED')
        for r in (1, 2, 3):
            for c in range(1, 15):
                cell = ws.cell(row=r, column=c)
                cell.fill = HEAD_FILL
                if r > 1:
                    cell.font = BOLD
                    cell.alignment = CENTER
                    cell.border = BOX
        for r in (2, 3):
            ws.cell(row=r, column=COL_NOTE).font = Font(name='Calibri', size=10, color='808080')
            ws.cell(row=r, column=COL_NOTE).alignment = CENTER
        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'B4'
        ws.page_setup.orientation = 'landscape'        # 打印：横向、一页宽
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = '1:3'

    def _next(self) -> int:
        self.n_rows += 1
        return 3 + self.n_rows

    def _name_sheet(self, d: date):
        if not self.named:
            self.ws.title = MONTH_TAB[d.month]
            self.named = True

    def _blank_row(self, no: int, note: str, fill):
        r = self._next()
        c = put(self.ws, r, COL_NO, no)
        c.alignment, c.border, c.font = CENTER, BOX, BOLD
        for col in range(COL_FIRST, COL_FIRST + len(FIELDS)):
            cell = self.ws.cell(row=r, column=col)
            cell.fill, cell.border = fill, BOX
        n = put(self.ws, r, COL_NOTE, note)
        n.font, n.alignment = RED_FONT if fill is RED_FILL else BASE_FONT, LEFT

    def add(self, no: int, records: list[dict] | None, error: str = ''):
        self.photos += 1
        if records is None:
            self._blank_row(no, f'✗ 识别失败：{error}（照片还在待识别里，下次点开始会重干）', RED_FILL)
            return
        if not records:
            self._blank_row(no, '这张照片里没认出刀具破损记录', RED_FILL)
            return
        for rec in records:
            self._record_row(no, rec)

    def unfinished(self, no: int, why: str):
        self._blank_row(no, f'未完成：{why}（照片还在待识别里，下次点开始会重干）', GREY_FILL)

    def _record_row(self, no: int, rec: dict):
        ws, r = self.ws, self._next()
        c = put(ws, r, COL_NO, no)
        c.alignment, c.border, c.font = CENTER, BOX, BOLD
        notes, red, orange = [], set(), set()
        values = {}
        unsure = set(rec.get('unsure') or ())       # 模型自己说拿不准的字段：认得清的不标，拿不准的才标
        disputed = dict(rec.get('disputed') or {})   # 几遍认得不一样的字段
        if self.kb is not None:                     # 先拿数据库当准绳：像的改成库里的写法（橙）
            rec, fixed, flags = self.kb.fix(rec, unsure)
            for f, (old, new, why) in fixed.items():
                orange.add(f)
                notes.append(f'{f.upper()} {old}→{new}（{why}）')
            for f, why in flags.items():
                red.add(f)
                notes.append(why)
        for f in FIELDS:
            raw = rec.get(f, '')
            v: object = raw
            if f == 'date' and raw:                 # 日期从不标红（Song 定的：错了无所谓），认不出就照原样写
                d = parse_date(raw, self.today)
                if d:
                    v = datetime(d.year, d.month, d.day)
                    self._name_sheet(d)
            elif f in NUM_FIELDS and raw:
                n = to_int(raw)
                if n is None:
                    red.add(f)
                    notes.append(f'{f.upper()} 不是数字：{raw}')
                else:
                    v = n
            elif f in ('mo', 'opr_no') and raw:
                n = to_int(raw)
                v = n if n is not None else raw
            elif f == 'shift' and raw:
                v = norm_shift(raw)
                if v not in ('D', 'N'):
                    red.add(f)
                    notes.append(f'SHIFT 不是 D/N：{raw}')
            elif f == 'setter' and raw:
                # Song 定的：工号和名字有一个就知道是谁，SETTER 照单子抄，不补、不按名单标红。
                # 只有工号里有看不清的字（□）而名单能唯一补上时才补。
                if '□' in raw and self.names:
                    mm = re.fullmatch(r'\s*([0-9□]+)(.*)', raw)
                    if mm and '□' in mm.group(1):
                        hit = [n for n, _ in self.names if re.fullmatch(mm.group(1).replace('□', r'\d'), n)]
                        if len(hit) == 1:
                            v = hit[0] + mm.group(2)
                            notes.append(f'SETTER 看不清的工号按名单补成 {hit[0]}')
            if isinstance(v, str) and '□' in v and f not in NEVER_RED:
                red.add(f)
                notes.append(f'{f.upper()} 有看不清的字')
            values[f] = v
        if rec.get('single'):
            notes.append('只认成了一遍，没能对答案')
        if rec.get('counts'):
            notes.append('几遍认出的条数不一样（' + ' / '.join(map(str, rec['counts'])) + ' 条），按多数定')
        for f in FIELDS:                            # 几遍认得不一样、又没被数据库改好的字段：标红
            if f in disputed and f not in orange and f not in NEVER_RED:
                red.add(f)
                notes.append(f'{f.upper()} 几遍认得不一样：' + ' / '.join(disputed[f]))
        for f in FIELDS:                            # 模型拿不准、又没被数据库改好的字段：标红
            if f in unsure and f not in red and f not in orange and f not in NEVER_RED:
                red.add(f)
                notes.append(f'{f.upper()} 模型拿不准')
        t, p, fl = (values.get(k) for k in ('total_screened', 'passed', 'failed'))
        if all(isinstance(x, int) for x in (t, p, fl)) and p + fl != t:
            red |= {'total_screened', 'passed', 'failed'}
            notes.append(f'PASSED {p} + FAILED {fl} ≠ TOTAL {t}')
        for i, f in enumerate(FIELDS):
            cell = put(ws, r, COL_FIRST + i, values[f] if values[f] != '' else None)
            cell.border, cell.font = BOX, BASE_FONT
            cell.alignment = LEFT if f in LEFT_FIELDS else CENTER
            if f == 'date' and isinstance(values[f], datetime):
                cell.number_format = 'd-mmm'
            if f == 'btmp':
                cell.fill = YELLOW
            if f in orange:
                cell.fill, cell.font = ORANGE_FILL, ORANGE_FONT
            if f in red:
                cell.fill, cell.font = RED_FILL, RED_FONT
        if notes:
            n = put(ws, r, COL_NOTE, '；'.join(notes))
            n.font, n.alignment = (RED_FONT if red else ORANGE_FONT), LEFT

    def save(self):
        last = max(3 + self.n_rows, 4)
        self.ws.auto_filter.ref = f'A3:N{last}'
        self.ws.print_area = f'A1:N{last}'
        tmp = self.path + '.tmp'
        self.wb.save(tmp)
        for _ in range(5):
            try:
                os.replace(tmp, self.path)
                return
            except PermissionError:      # Excel 正开着这个文件
                time.sleep(2)
        alt = self.path.replace('.xlsx', f'_{datetime.now():%H%M%S}.xlsx')
        os.replace(tmp, alt)
        self.path = alt


# ---------------------------------------------------------------- 主流程
@dataclass
class Result:
    ok: int = 0
    records: int = 0
    failed: int = 0
    skipped: int = 0
    requests: int = 0
    xlsx: str = ''
    photos_dir: str = ''
    stopped_reason: str = ''
    errors: list = field(default_factory=list)


def shot_time(path: str) -> float:
    """拍摄时间：先看照片里的 EXIF（拍照那一刻），没有就用文件修改时间。"""
    try:
        with Image.open(path) as im:
            ex = im.getexif()
            raw = ex.get_ifd(0x8769).get(36867) or ex.get(36867) or ex.get(306)
        if raw:
            return datetime.strptime(str(raw).strip()[:19], '%Y:%m:%d %H:%M:%S').timestamp()
    except Exception:
        pass
    return os.path.getmtime(path)


def scan(folder: str) -> list[str]:
    out = []
    for name in os.listdir(folder):
        p = os.path.join(folder, name)
        if os.path.isfile(p) and os.path.splitext(name)[1].lower() in IMG_EXT and os.path.getsize(p) > 0:
            out.append(p)
    out.sort(key=lambda p: (shot_time(p), os.path.basename(p).lower()))
    return out


def recycle(path: str) -> bool:
    """扔进 Windows 回收站（能从回收站捡回来，不是永久删除）。不是 Windows 或者扔不进去就返回 False。"""
    if os.name != 'nt':
        return False
    try:
        import ctypes
        from ctypes import wintypes

        class SHFILEOPSTRUCTW(ctypes.Structure):
            _pack_ = 1 if ctypes.sizeof(ctypes.c_void_p) == 4 else 8
            _fields_ = [('hwnd', wintypes.HWND), ('wFunc', wintypes.UINT), ('pFrom', wintypes.LPCWSTR),
                        ('pTo', wintypes.LPCWSTR), ('fFlags', ctypes.c_ushort), ('fAnyOperationsAborted', wintypes.BOOL),
                        ('hNameMappings', ctypes.c_void_p), ('lpszProgressTitle', wintypes.LPCWSTR)]
        fo_delete, flags = 3, 0x40 | 0x10 | 0x4 | 0x400     # 可撤销（回收站）| 不确认 | 不弹进度 | 不弹错误框
        op = SHFILEOPSTRUCTW(None, fo_delete, os.path.abspath(path) + '\0', None, flags, False, None, None)
        ok = ctypes.windll.shell32.SHFileOperationW(ctypes.byref(op)) == 0 and not op.fAnyOperationsAborted
        return ok and not os.path.exists(path)
    except Exception:
        return False


def unique_dest(d: str, name: str) -> str:
    base, ext = os.path.splitext(name)
    p, k = os.path.join(d, name), 1
    while os.path.exists(p):
        p = os.path.join(d, f'{base}_{k}{ext}')
        k += 1
    return p


@dataclass
class CopyResult:
    copied: list = field(default_factory=list)        # 复制进去的（新路径）
    already: list = field(default_factory=list)       # 待识别里已经有一模一样的
    done: list = field(default_factory=list)          # (文件名, 账本里那条)：以前干过的
    bad: list = field(default_factory=list)           # (文件名, 为什么没复制)


def copy_in(paths, folder: str) -> CopyResult:
    """「选照片…」：把挑的照片复制进待识别（Song 定的：复制，原照片原地不动；复制完还是人点「开始工作」）。
    待识别没有就当场建。待识别里已经有一模一样的、以前干过的（账本里有，改了名也认得）不再复制。"""
    import shutil
    res = CopyResult()
    os.makedirs(folder, exist_ok=True)
    ledger = {}
    for name in (LEDGER, OLD_LEDGER):
        try:
            with open(os.path.join(folder, OUT_DIR, name), encoding='utf-8') as fh:
                ledger = json.load(fh)
            break
        except Exception:
            continue
    here: dict[str, str] = {}                       # 待识别里已有照片的指纹（要用时才算）

    def known() -> dict[str, str]:
        if not here:
            for p in scan(folder):
                try:
                    here[file_hash(p)] = p
                except OSError:
                    pass
            here.setdefault('', '')
        return here

    for src in paths:
        name = os.path.basename(src)
        if os.path.splitext(name)[1].lower() not in IMG_EXT:
            res.bad.append((name, '不是照片'))
            continue
        try:
            if os.path.normcase(os.path.dirname(os.path.abspath(src))) == os.path.normcase(os.path.abspath(folder)):
                res.already.append(name)            # 本来就在待识别里
                continue
            h = file_hash(src)
            if h in ledger:
                res.done.append((name, ledger[h]))
                continue
            if h in known():
                res.already.append(name)
                continue
            dest = unique_dest(folder, name)
            shutil.copy2(src, dest)                 # copy2：连拍摄时间、修改时间一起带过去，编号顺序不乱
            here[h] = dest
            res.copied.append(dest)
        except OSError as e:
            res.bad.append((name, f'复制不了：{e.strerror or e}'))
    return res


class Engine:
    def __init__(self, folder: str, keys: list[str], model: str = DEFAULT_MODEL,
                 log=print, progress=None, stop: threading.Event | None = None,
                 clock=time.time, sleep=None, names_path: str | None = None,
                 kb_path: str | None = None, feed_dir: str | None = None, reads: int | None = None,
                 keep_originals: bool = True, discard=recycle, spec=None, collect: bool = False):
        self.folder = folder
        self.collect = collect      # ⚙「发数据给方丈」开着：每张怎么认的都记下来，干完写进 识别结果\.回传（见 report.py）
        self.collect_photos = False
        self.trace: dict[str, dict] = {}
        self.run_name = ''
        self.kb_line = ''
        self.spec = spec          # None = 刀单（下面这套专门的）；自建类别交给 tables.TableSpec
        self.keep_originals, self.discard = keep_originals, discard   # 不留原图：认完扔进回收站（扔不进去就照旧留着）
        self.reads = max(1, reads if reads is not None else READS)
        self.model = model_name(model)
        self.log = log
        self.progress = progress or (lambda done, total: None)
        self.stop = stop or threading.Event()
        self.clock = clock
        self._sleep = sleep
        self.keys = Keys(keys, clock)
        self.lock = threading.Lock()
        self.res = Result()
        self.fatal: str = ''
        self.use_schema = True
        self.names = load_names(names_path)
        self.kb_path, self.feed_dir = kb_path, feed_dir
        self.snapshot = None
        self.hints = ''

    # 可被「停止」打断的等待
    def wait(self, seconds: float):
        if self._sleep:
            self._sleep(seconds)
            if self.stop.is_set():
                raise Stopped()
            return
        end = time.time() + seconds
        while time.time() < end:
            if self.stop.wait(min(1.0, end - time.time())):
                raise Stopped()

    def wait_until(self, ts: float, why: str):
        secs = max(0.0, ts - self.clock())
        if secs > 90:
            back = datetime.fromtimestamp(ts).strftime('%m-%d %H:%M')
            self.log(f'⏸ {why}，等到 {back} 额度刷新后接着干（约 {secs / 3600:.1f} 小时）')
        elif secs >= 1:
            self.log(f'⏸ {why}，等 {int(secs)} 秒')
        self.wait(secs)

    def recognize(self, path: str) -> tuple[list[dict], str]:
        try:
            return self._recognize(path)
        except Fatal as e:
            if not self.fatal:
                self.fatal = str(e)
            self.stop.set()          # 同批其他几张别再傻等
            raise

    def _recognize(self, path: str) -> tuple[list[dict], str]:
        """认一张：先认两遍（第二遍用拉过对比度的图），对得上就收；对不上认第三遍，按多数定。
        顺便按多数定单子标题在照片哪一边（印编号的那份照片拿它转正）。"""
        jpeg = prepare_image(path)
        clean = self.spec.clean if self.spec else clean_record
        agree = self.spec.agree if self.spec else records_agree
        vote_ = self.spec.vote if self.spec else vote
        if self.reads == 1:
            recs, side = self._traced_read(jpeg, path, '原图')
            return vote_([[clean(r) for r in recs]]), side
        plan = [(jpeg, '原图'), (prepare_image(path, enhance=True), '加强')]
        reads, errs, sides = [], [], []
        while plan:
            img, what = plan.pop(0)
            try:
                recs, side = self._traced_read(img, path, what)
                reads.append([clean(r) for r in recs])
                sides.append(side)
            except (Stopped, Fatal):
                raise
            except Exception as e:              # 某一遍没认成：再补一遍，凑够两遍对答案
                errs.append(e)
                if len(errs) == 1 and len(reads) + len(plan) < 2:
                    plan.append((jpeg, '原图'))
                continue
            if not plan and len(reads) == 2 and not agree(reads[0], reads[1]):
                plan.append((jpeg, '原图'))           # 两遍对不上：第三遍当裁判
        if not reads:
            raise errs[0]
        out = vote_(reads)
        if len(reads) == 1:
            for r in out:
                r['single'] = True
        known = [x for x in sides if x]
        side = ''
        if known:
            c = Counter(known)
            top = c.most_common(1)[0][1]
            side = known[0] if c[known[0]] == top else c.most_common(1)[0][0]
        return out, side

    def _traced_read(self, img: bytes, path: str, what: str) -> tuple[list[dict], str]:
        """认一遍；开着「发数据给方丈」就把这一遍认成什么（模型原话，整理之前）、花了多久记下来。"""
        t0 = time.time()
        try:
            recs, side = self._read(img, path)
        except (Stopped, Fatal):
            raise
        except Exception as e:
            if self.collect:
                why = e.raw() if isinstance(e, ApiError) else f'{type(e).__name__}：{e}'
                self.trace.setdefault(path, {}).setdefault('reads', []).append(
                    {'img': what, 'ms': int((time.time() - t0) * 1000), 'error': why})
            raise
        if self.collect:
            self.trace.setdefault(path, {}).setdefault('reads', []).append(
                {'img': what, 'ms': int((time.time() - t0) * 1000), 'side': side, 'records': recs})
        return recs, side

    def _collect(self, no: int, path: str, width: int, side: str, err: str, t0: float):
        """一张认完（不管成没成）：记下照片本身的样子，再存一份缩小的给方丈。出什么错都不耽误干活。"""
        try:
            t = self.trace.setdefault(path, {})
            t.update(side=side, error=err, ms=int((time.time() - t0) * 1000))
            from . import report
            info, small = report.snapshot(path)
            t['look'] = info
            if self.collect_photos:
                pics = os.path.join(report.run_dir(self.folder, self.run_name), report.PICS)
                os.makedirs(pics, exist_ok=True)
                with open(os.path.join(pics, f'{no:0{width}d}.jpg'), 'wb') as fh:
                    fh.write(small)
        except Exception:
            pass

    def _body(self, jpeg: bytes, schema: bool) -> dict:
        if self.spec:
            return self.spec.body(jpeg, self.model, schema)
        return build_body(jpeg, self.model, schema, self.names, self.hints)

    def _read(self, jpeg: bytes, path: str) -> tuple[list[dict], str]:
        body = self._body(jpeg, self.use_schema)
        minute_hits = 0
        net_tries = 0
        while True:
            if self.stop.is_set():
                raise Stopped()
            if self.fatal:
                raise Fatal(self.fatal)
            i, key = self.keys.pick()
            if i is None:
                self.wait_until(key, '所有钥匙额度都用完了' if key - self.clock() > 120 else '每分钟次数到顶')
                continue
            try:
                with self.lock:
                    self.res.requests += 1
                data = call_api(self.model, key, body)
                self.keys.used[i] += 1
                return self.spec.parse(data) if self.spec else parse_response(data)
            except ApiError as e:
                if e.status == 429:
                    kind = classify_429(e)
                    if kind == 'prepay':
                        raise Fatal(f'钥匙{i + 1} · {e.raw()}（账单账户预付余额用完，等也没用）')
                    if kind == 'day':
                        until = next_pacific_midnight(self.clock())
                        self.keys.park(i, until)
                        self.log(f'钥匙{i + 1} 今天的额度用完了 · {e.raw()}')
                    else:
                        minute_hits += 1
                        if minute_hits > MAX_MINUTE_WAITS:
                            raise
                        d = retry_delay(e.detail)
                        wait = MINUTE_WAIT if d is None else min(MINUTE_WAIT, max(3.0, d + 1))
                        self.keys.park(i, self.clock() + wait)
                    continue
                if e.status == 400 and re.search(r'schema', e.message, re.I):
                    if self.use_schema:
                        self.use_schema = False
                        self.log(f'模型不认结构化输出格式，改用纯 JSON 提示重试 · {e.raw()}')
                    if 'responseSchema' in body['generationConfig']:
                        body = self._body(jpeg, False)
                        continue
                if e.status in (400, 401, 403) and re.search(r'API key|API_KEY|permission|PERMISSION', f'{e.message} {e.detail}'):
                    raise Fatal(f'钥匙{i + 1} · {e.raw()}')
                if e.status == 404:
                    raise Fatal(f'{self.model} · {e.raw()}')
                if e.status == 0 or e.status >= 500:
                    if net_tries < len(RETRY_WAITS):
                        w = RETRY_WAITS[net_tries]
                        net_tries += 1
                        self.log(f'{os.path.basename(path)} · {e.raw()} · {w} 秒后重试')
                        self.wait(w)
                        continue
                raise

    def run(self) -> Result:
        res = self.res
        if not self.keys.keys:
            res.stopped_reason = '没填钥匙'
            self.log('✗ 没填钥匙')
            return res
        os.makedirs(self.folder, exist_ok=True)
        out_dir = os.path.join(self.folder, OUT_DIR)
        done_root = os.path.join(self.folder, DONE_DIR)
        os.makedirs(out_dir, exist_ok=True)
        os.makedirs(done_root, exist_ok=True)
        logfile = os.path.join(out_dir, '录入僧日志.txt')
        outer = self.log

        def log(msg):
            outer(msg)
            try:
                with open(logfile, 'a', encoding='utf-8') as fh:
                    fh.write(f'{datetime.now():%Y-%m-%d %H:%M:%S}  {msg}\n')
            except OSError:
                pass
        self.log = log
        if self.spec:                               # 自建类别：吃它自己的投喂文件夹，拿它自己的库当准绳
            self.spec.prepare(self.log)
        elif self.kb_path:                          # 数据库：先吃投喂文件夹里新放的表，再拿来当准绳
            try:
                from .kb import KnowledgeBase
                kb = KnowledgeBase(self.kb_path)
                for fr in kb.feed_folder(self.feed_dir):
                    self.log(fr.line())
                    if self.collect and not fr.error:   # 投喂文件夹里吃进去的表：人改对的答案，也给方丈一份
                        try:
                            from . import report
                            report.queue_sample(self.folder, fr.path, '刀单')
                        except Exception:
                            pass
                self.snapshot = kb.load()
                self.hints = self.snapshot.hints()
                have = {n for n, _ in self.names}
                self.names = self.names + [(n, nm) for n, nm in self.snapshot.roster() if n not in have]
                self.kb_line = kb.summary()
                self.log(self.kb_line)
            except Exception as e:
                self.log(f'数据库打不开，这次不用它：{type(e).__name__}：{e}')
                self.snapshot, self.hints = None, ''
        ledger_path = os.path.join(out_dir, LEDGER)
        old_ledger = os.path.join(out_dir, OLD_LEDGER)
        if not os.path.exists(ledger_path) and os.path.exists(old_ledger):
            try:
                os.replace(old_ledger, ledger_path)      # 改名前干过的照片照样认得
            except OSError:
                ledger_path = old_ledger
        try:
            with open(ledger_path, encoding='utf-8') as fh:
                ledger = json.load(fh)
        except Exception:
            ledger = {}

        todo = []
        for p in scan(self.folder):
            h = file_hash(p)
            if h in ledger:          # 干过的（改了名也认得出）
                dup = os.path.join(done_root, DUP_DIR)
                os.makedirs(dup, exist_ok=True)
                os.replace(p, unique_dest(dup, os.path.basename(p)))
                res.skipped += 1
                self.log(f'↷ {os.path.basename(p)} 以前干过（{ledger[h].get("xlsx", "")} 第 {ledger[h].get("no", "?")} 张），挪进 已识别\\{DUP_DIR}')
            else:
                todo.append((p, h))
        total = len(todo)
        if not total:
            self.log('文件夹里没有新图片，收工。')
            return res

        run_id = datetime.now().strftime('%Y%m%d_%H%M%S')
        width = max(2, len(str(total)))
        done_dir = os.path.join(done_root, run_id)
        tail = f'_{self.spec.name}' if self.spec else ''
        res.xlsx = os.path.join(out_dir, f'识别结果_{run_id}{tail}.xlsx')
        res.photos_dir = done_dir
        self.run_name = f'{run_id}{tail}'
        started = time.time()
        if self.collect:
            try:
                from . import report
                self.collect_photos = report.room(self.folder)
            except Exception:
                self.collect = False
        book = self.spec.book(res.xlsx) if self.spec else RecordBook(res.xlsx, names=self.names, kb=self.snapshot)
        items = [(i + 1, p, h) for i, (p, h) in enumerate(todo)]      # 按拍摄顺序编号：第 1 张对应表里 No. 1
        if self.spec:
            what = f'，类别「{self.spec.name}」' + ('（还没定列：认全部内容）' if self.spec.open_mode
                                                 else f'（{len(self.spec.columns)} 列）')
        else:
            what = '，类别「刀单」' + (f'，名单 {len(self.names)} 人' if self.names else '，没找到名单')
        self.log(f'▶ 开工（录入僧 v{VERSION}）：{total} 张，按拍摄顺序编号 1–{total}，每批 {BATCH} 张，模型 {self.model}'
                 + ('，每张认两遍对答案' if self.reads >= 2 else '') + what)
        done = 0
        self.progress(0, total)

        def one(item):
            no, p, h = item
            t0 = time.time()
            err = ''
            try:
                recs, side = self.recognize(p)
            except (Stopped, Fatal):
                raise
            except ApiError as e:
                recs, side, err = None, '', e.raw()
            except Exception as e:
                recs, side, err = None, '', f'{type(e).__name__}：{e}'
            if self.collect:
                self._collect(no, p, width, side, err, t0)
            if recs is None:
                return no, p, h, None, err
            try:                                    # 认完了才印编号：送去识别的永远是原图
                os.makedirs(done_dir, exist_ok=True)
                stamp_number(p, no, os.path.join(done_dir, f'{no:0{width}d}.jpg'), side)
            except Exception as e:
                self.log(f'{no:0{width}d} 编号照片没印成：{type(e).__name__}：{e}')
            return no, p, h, recs, ''

        try:
            with ThreadPoolExecutor(BATCH) as pool:
                for b in range(0, total, BATCH):
                    batch = items[b:b + BATCH]
                    futures = [(it, pool.submit(one, it)) for it in batch]
                    halt = None
                    for (no, p, h), f in futures:          # 按编号顺序写进表
                        try:
                            _, _, _, recs, err = f.result()
                        except (Stopped, Fatal) as e:
                            halt = halt or e
                            why = '已停止' if isinstance(e, Stopped) and not self.fatal else (self.fatal or str(e))
                            book.unfinished(no, why)
                            if self.collect:
                                self.trace.setdefault(p, {})['error'] = f'未完成：{why}'
                            continue
                        name = os.path.basename(p)
                        tag = f'{no:0{width}d}'
                        book.add(no, recs, err)
                        if recs is not None:
                            res.ok += 1
                            res.records += len(recs)
                            tossed = (not self.keep_originals and os.path.exists(os.path.join(done_dir, f'{tag}.jpg'))
                                      and self.discard(p))     # 有编号照片了才扔原图
                            if not tossed:
                                orig_dir = os.path.join(done_dir, ORIG_DIR)
                                os.makedirs(orig_dir, exist_ok=True)
                                os.replace(p, unique_dest(orig_dir, f'{tag}_{name}'))
                            ledger[h] = {'file': name, 'no': no, 'xlsx': os.path.basename(book.path),
                                         'at': datetime.now().isoformat(timespec='seconds')}
                            self.log(f'✓ {tag} {name} · {len(recs)} 条')
                        else:
                            res.failed += 1
                            res.errors.append((name, err))
                            self.log(f'✗ {tag} {name} · {err}')
                        done += 1
                    book.save()
                    with open(ledger_path, 'w', encoding='utf-8') as fh:
                        json.dump(ledger, fh, ensure_ascii=False, indent=1)
                    self.progress(done, total)
                    if halt:
                        raise halt
        except Fatal as e:
            res.stopped_reason = str(e)
            self.log(f'■ 停工：{e}')
        except Stopped:
            if self.fatal:
                res.stopped_reason = self.fatal
                self.log(f'■ 停工：{self.fatal}')
            else:
                res.stopped_reason = '手动停止'
                self.log('■ 已停止，没干完的下次接着干')
        if res.ok == 0 and res.failed == 0:
            try:
                os.remove(res.xlsx)
            except OSError:
                pass
            res.xlsx = ''
        if self.collect and res.xlsx:
            try:                                    # 发数据给方丈：这批怎么认的、刚认完时的表，落在本机等着发
                from . import report
                report.finish_run(self, items, getattr(book, 'path', res.xlsx), started, self.kb_line)
            except Exception as e:
                self.log(f'给方丈的记录没写成（不耽误干活）：{type(e).__name__}：{e}')
        used =' · '.join(f'钥匙{i + 1}×{n}' for i, n in enumerate(self.keys.used) if n)
        self.log(f'{"✓" if not res.failed and not res.stopped_reason else "⚠"} 成功 {res.ok} 张（{res.records} 条记录）· 失败 {res.failed} 张 · '
                 f'{res.requests} 次请求{" · " + used if used else ""}' + (f'\n结果：{res.xlsx}' if res.xlsx else ''))
        return res
