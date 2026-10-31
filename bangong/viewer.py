"""核对窗口：一张一张看编号照片（第一张就是 1），右边列出表里同一号的记录，红的就是要看的。"""
from __future__ import annotations

import os
import re
import subprocess
import sys
import tkinter as tk
from datetime import date, datetime


from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageTk

from . import engine
from . import macui as ui
from .kb import cell_text

LABELS = ['DATE', 'MACHINE', 'PART NUMBER', 'MO', 'OPR NO.', 'T NO.', 'DESCRIPTION', 'BTMP',
          'TOTAL SCREENED', 'PASSED', 'FAILED', 'SHIFT', 'SETTER']
RUN_RE = re.compile(r'\d{8}_\d{6}')
BTMP = LABELS.index('BTMP')


# ---------------------------------------------------------------- 数据（不碰 Tk，方便测试）
def load_photos(run_dir: str) -> dict[int, str]:
    out = {}
    if os.path.isdir(run_dir):
        for f in os.listdir(run_dir):
            m = re.fullmatch(r'(\d+)\.jpg', f, re.I)
            if m:
                out[int(m.group(1))] = os.path.join(run_dir, f)
    return out


def find_xlsx(folder: str, run_id: str) -> str:
    od = os.path.join(folder, engine.OUT_DIR)
    if not os.path.isdir(od):
        return ''
    c = sorted(f for f in os.listdir(od) if f.startswith(f'识别结果_{run_id}') and f.endswith('.xlsx'))
    return os.path.join(od, c[0]) if c else ''


def list_runs(folder: str) -> list[tuple[str, str, str]]:
    """[(批次号, 编号照片文件夹, 那一批的 Excel)]，新的在前。"""
    done = os.path.join(folder, engine.DONE_DIR)
    out = []
    if os.path.isdir(done):
        for name in sorted(os.listdir(done), reverse=True):
            d = os.path.join(done, name)
            if os.path.isdir(d) and RUN_RE.fullmatch(name) and load_photos(d):
                out.append((name, d, find_xlsx(folder, name)))
    return out


RESULT_RE = re.compile(r'识别结果_(\d{8}_\d{6})(?:_(.+))?\.xlsx', re.I)


def list_results(folder: str) -> list[dict]:
    """「所有结果」：图片文件夹里每一批的表，新的在前。
    [{'run': 批次号, 'xlsx': 表, 'cat': 类别, 'photos': 编号照片几张, 'dir': 编号照片文件夹}]"""
    od = os.path.join(folder, engine.OUT_DIR)
    out = []
    if not os.path.isdir(od):
        return out
    for f in os.listdir(od):
        m = RESULT_RE.fullmatch(f)
        if not m or f.startswith('~$'):
            continue
        cat = re.sub(r'(?:^|_)\d{6}$', '', m.group(2) or '')      # Excel 开着时另存的「…_181502.xlsx」
        d = os.path.join(folder, engine.DONE_DIR, m.group(1))
        out.append({'run': m.group(1), 'xlsx': os.path.join(od, f), 'cat': cat or '刀单',
                    'photos': len(load_photos(d)), 'dir': d})
    out.sort(key=lambda r: (r['run'], r['xlsx']), reverse=True)
    return out


def fmt(v) -> str:
    if isinstance(v, datetime):
        return f'{v.day}-{v.strftime("%b")}'
    return '' if v is None else str(v)


HUMAN_NOTE = '✓ 人核过'                                   # 红格子全处理完：核对提示换成这句
HUMAN_FILL = PatternFill('solid', fgColor='E2EFDA')         # 核对里人改过、点过「对的」的格子：淡绿
HUMAN_FONT = Font(name='Calibri', size=11, color='375623')
NO_FILL = PatternFill(fill_type=None)
SUM = ('total_screened', 'passed', 'failed')


class Locked(Exception):
    """表在 Excel 里开着，写不进去。"""


def _layout(ws) -> tuple[bool, list[str], int, int, int]:
    """(是不是刀单, 列名, 第一列, 数据从第几行起, 核对提示在第几列)"""
    if str(ws.cell(row=1, column=1).value or '').strip() == 'TOOL BROKEN RECORD':
        return True, LABELS, 2, 4, 15
    heads = []                                      # 自建类别：第 2 行是表头，最后一列是核对提示
    c = 2
    while c <= ws.max_column:
        h = ws.cell(row=2, column=c).value
        if h is None or str(h).startswith('核对提示'):
            break
        heads.append(str(h))
        c += 1
    return False, heads, 2, 3, 2 + len(heads)


def _look(cell) -> str:
    rgb = cell.fill.fgColor.rgb if cell.fill is not None and cell.fill.fill_type == 'solid' else ''
    rgb = rgb.upper() if isinstance(rgb, str) else ''
    return 'red' if rgb.endswith('FFC7CE') else 'fixed' if rgb.endswith('FCE4D6') else 'human' if rgb.endswith('E2EFDA') else ''


def _names(builtin: bool, labels: list[str]) -> list[str]:
    """核对提示里怎么称呼每一格：刀单是 MACHINE、PART_NUMBER 这样，自建类别就是列名。"""
    return [f.upper() for f in engine.FIELDS] if builtin else list(labels)


def alts_from_note(note: str, name: str) -> list[str]:
    """红格子的候选：核对提示里「XX 几遍认得不一样：A / B」的 A、B（空的不算）。"""
    m = re.search(r'(?:^|；)' + re.escape(name) + r' 几遍认得不一样：([^；]+)', note or '')
    return [v.strip() for v in m.group(1).split(' / ') if v.strip() and v.strip() != '空'] if m else []


def _read_row(ws, r: int, layout) -> dict:
    builtin, labels, first, _, note_col = layout
    vals, raw, red, fixed, human = [], [], set(), set(), set()
    for i in range(len(labels)):
        cell = ws.cell(row=r, column=first + i)
        vals.append(fmt(cell.value))
        raw.append(cell_text(cell.value))
        look = _look(cell)
        {'red': red, 'fixed': fixed, 'human': human}.get(look, set()).add(i)
    note = ws.cell(row=r, column=note_col).value
    note = '' if note is None else str(note)
    names = _names(builtin, labels)
    alts = {i: a for i in sorted(red) if (a := alts_from_note(note, names[i]))}
    return {'labels': labels, 'values': vals, 'red': red, 'fixed': fixed, 'note': note,
            'human': human, 'raw': raw, 'alts': alts, 'row': r, 'builtin': builtin}


def load_rows(xlsx: str) -> dict[int, list[dict]]:
    """{编号: [{'labels': 列名, 'values': 各列, 'red': 标红的列下标, 'fixed': 按库改过的, 'note': 核对提示,
    'human': 人核过的, 'raw': 放进数据库的写法, 'alts': {标红的列: 几遍认出的写法}, 'row': 在表里第几行}]}
    刀单那本（TOOL BROKEN RECORD，第 4 行起）和自建类别的表（第 2 行表头，第 3 行起）都认。"""
    rows: dict[int, list[dict]] = {}
    if not xlsx or not os.path.isfile(xlsx):
        return rows
    try:
        ws = load_workbook(xlsx).worksheets[0]
    except Exception:
        return rows
    layout = _layout(ws)
    for r in range(layout[3], ws.max_row + 1):
        no = ws.cell(row=r, column=1).value
        if isinstance(no, int):
            rows.setdefault(no, []).append(_read_row(ws, r, layout))
    return rows


# ---------------------------------------------------------------- 核对里直接改格子（跟网页版 core.js 的 checkCell 同一套）
def typed_value(key: str, text: str, builtin: bool = True, orig=None, today: date | None = None):
    """人打的字变成表里该有的样子：规矩跟认出来时一样（日期按这批的年份、数字变数字、班次变 D/N、零件号去空格大写）。"""
    t = re.sub(r'\s+', ' ', str(text or '')).strip()
    if not t:
        return ''
    if builtin:
        if key == 'date':
            d = engine.parse_date(t, today)
            return datetime(d.year, d.month, d.day) if d else t
        if key in engine.NUM_FIELDS or key in ('mo', 'opr_no'):
            n = engine.to_int(t)
            return t if n is None else n
        if key == 'shift':
            return engine.norm_shift(t)
        if key in ('machine', 'part_number'):
            return engine.DASHES.sub('-', re.sub(r'\s+', '', t)).upper()
        if key in ('t_no', 'description'):
            return engine.DIAMETER.sub('Φ', t).upper()
        if key == 'setter':
            return re.sub(r'\s*/\s*', '/', t)
        return t
    if isinstance(orig, (int, float)) and not isinstance(orig, bool):
        n = t.replace(',', '')
        if re.fullmatch(r'-?\d+(?:\.\d+)?', n):
            return int(n) if re.fullmatch(r'-?\d+', n) else float(n)
    return t


def typed_problem(key: str, v) -> str | None:
    """人打进刀单格子里的字，跟认出来的一样查一遍：数字格要是数字、班次要是 D/N、不能有看不清的 □。"""
    if v is None or v == '':
        return None
    if key in engine.NUM_FIELDS and not (isinstance(v, int) and not isinstance(v, bool)):
        return f'{key.upper()} 不是数字：{v}'
    if key == 'shift' and v not in ('D', 'N'):
        return f'SHIFT 不是 D/N：{v}'
    if isinstance(v, str) and '□' in v and key not in engine.NEVER_RED:
        return f'{key.upper()} 有看不清的字'
    return None


def note_owner(piece: str, names: list[str]) -> str | None:
    """核对提示里的一句是哪一格的：开头写着格子名；名字有长有短，取最长对得上的。"""
    best = None
    for nm in names:
        if piece.startswith(nm + ' ') and (best is None or len(nm) > len(best)):
            best = nm
    return best


def _restyle(cell, look: str, builtin: bool, key: str):
    if look == 'red':
        cell.fill, cell.font = engine.RED_FILL, engine.RED_FONT
    elif look == 'human':
        cell.fill, cell.font = HUMAN_FILL, HUMAN_FONT
    elif look == 'fixed':
        cell.fill, cell.font = engine.ORANGE_FILL, engine.ORANGE_FONT
    else:
        cell.fill = engine.YELLOW if builtin and key == 'btmp' else NO_FILL
        cell.font = engine.BASE_FONT


def _isint(v) -> bool:
    return isinstance(v, int) and not isinstance(v, bool)


def edit_cell(xlsx: str, row: int, i: int, text: str | None = None, today: date | None = None) -> dict:
    """核对里改一格，直接写回表里：给了 text 就改成它，没给就是「对的」（照原样，只把红去掉）。
    改过的格子染淡绿；这一格的核对提示跟着删；数量对上了三格一起不红；打错了（数字格打了字母）还红着。
    红格子全处理完，核对提示换成「✓ 人核过」。表在 Excel 里开着就抛 Locked，一个字都不动。
    返回 {'value': 这格现在显示的字, 'still': 还红着没有, 'why': 为什么还红, 'checked': 这条核完没有,
    'data': 核完的话放进数据库的那份}。"""
    # 不看旁边有没有 ~$ 小文件：Excel 崩过会留下一个假的，那样就永远改不了了。真开着的话最后换文件那一下会被拒
    try:
        wb = load_workbook(xlsx)
    except PermissionError:
        raise Locked(xlsx)
    ws = wb.worksheets[0]
    layout = _layout(ws)
    builtin, labels, first, _, note_col = layout
    keys = engine.FIELDS if builtin else labels
    names = _names(builtin, labels)
    cells = [ws.cell(row=row, column=first + j) for j in range(len(labels))]
    before = [_look(c) for c in cells]
    red = {j for j, lk in enumerate(before) if lk == 'red'}
    human = {j for j, lk in enumerate(before) if lk == 'human'}
    cell, key, bad = cells[i], keys[i], None
    if text is not None:
        v = typed_value(key, text, builtin, cell.value, today)
        cell.value = None if v == '' else v
        if builtin and key == 'date':
            cell.number_format = 'd-mmm' if isinstance(v, datetime) else 'General'
        if builtin:
            bad = typed_problem(key, v)
    red.discard(i)
    human.add(i)
    note_cell = ws.cell(row=row, column=note_col)
    notes = [n for n in str(note_cell.value or '').split('；')
             if n and n != HUMAN_NOTE and note_owner(n, names) != names[i]]
    if bad:
        red.add(i)
        notes.append(bad)
    if builtin and key in SUM:                       # 数量：PASSED + FAILED 要等于 TOTAL
        idx = [engine.FIELDS.index(k) for k in SUM]
        t, p, f = (cells[j].value for j in idx)
        if all(_isint(x) for x in (t, p, f)):
            notes = [n for n in notes if not n.startswith('PASSED ')]
            if p + f == t:
                red -= set(idx)                      # 对上了：三格一起不红
            else:                                    # 还对不上：没核过的那几格红着
                red |= {j for j in idx if j not in human}
                notes.append(f'PASSED {p} + FAILED {f} ≠ TOTAL {t}')
    for j, c in enumerate(cells):
        look = 'red' if j in red else 'human' if j in human else ('fixed' if before[j] == 'fixed' else '')
        if look != before[j] or j == i:
            _restyle(c, look, builtin, keys[j])
    if red:
        note_cell.value = '；'.join(notes) or None
        note_cell.font = engine.RED_FONT
    else:
        note_cell.value = HUMAN_NOTE
        note_cell.font = HUMAN_FONT
    note_cell.alignment = engine.LEFT
    tmp = xlsx + '.edit.tmp'
    wb.save(tmp)
    try:
        os.replace(tmp, xlsx)
    except PermissionError:                          # Excel 开着这本：Windows 不让换
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise Locked(xlsx)
    data = {(engine.FIELDS[j] if builtin else labels[j]): cell_text(c.value) for j, c in enumerate(cells)}
    return {'value': fmt(cell.value), 'still': i in red, 'why': bad or (notes[-1] if i in red and notes else ''),
            'checked': not red, 'data': data}


def cat_of(xlsx: str) -> str:
    """这本表是哪个类别的：识别结果_批次号_类别.xlsx；没写类别就是刀单（Excel 开着时另存的 _181502 不算）。"""
    m = RESULT_RE.fullmatch(os.path.basename(xlsx or ''))
    cat = re.sub(r'(?:^|_)\d{6}$', '', (m.group(2) or '') if m else '')
    return cat or '刀单'


def run_day(run_id: str) -> date | None:
    """这批是哪天认的：改日期时按这批的年份。"""
    try:
        return datetime.strptime(run_id[:8], '%Y%m%d').date()
    except ValueError:
        return None


def rotate_photo(path: str, clockwise: bool) -> None:
    """自动转正万一转错了：把这张编号照片在硬盘上转 90°（原图不动）。"""
    with Image.open(path) as im:
        im = im.convert('RGB').transpose(Image.Transpose.ROTATE_270 if clockwise else Image.Transpose.ROTATE_90)
    im.save(path, 'JPEG', quality=90)


def run_label(run_id: str, n: int) -> str:
    d = datetime.strptime(run_id, '%Y%m%d_%H%M%S')
    return f'{d:%Y-%m-%d %H:%M} 这批（{n} 张）'


def _open(p: str):
    try:
        if os.name == 'nt':
            os.startfile(p)  # type: ignore[attr-defined]
        else:
            subprocess.Popen(['xdg-open' if sys.platform != 'darwin' else 'open', p])
    except Exception:
        pass


# ---------------------------------------------------------------- 窗口
class Viewer(tk.Toplevel):
    """on_checked(类别, 来源, 那条的数据或 None)：一条核完（红格子全处理完）就交给它放进样板资料，又改红了给 None 拿掉。
    busy()：正在干的那批的编号照片文件夹，那批的表干完会整本重写，先不让改。"""

    def __init__(self, master, folder: str, run_dir: str | None = None, start: int | None = None,
                 on_checked=None, busy=None):
        super().__init__(master)
        self.folder = folder
        self.on_checked, self.busy = on_checked, busy
        self.editing: tuple[int, int, int] | None = None      # 正在改哪一格：(编号, 第几条, 第几列)
        self.editor = None
        self.run_id = self.run_dir = self.cat = ''
        self.today: date | None = None
        self.geometry(f'{ui.px(1100)}x{ui.px(720)}')
        self.minsize(ui.px(640), ui.px(420))
        self.runs = list_runs(folder)
        self.photos: dict[int, str] = {}
        self.nos: list[int] = []
        self.rows: dict[int, list[dict]] = {}
        self.cur = 0
        self._pil = None
        self._tk = None
        self._after = None
        self.configure(bg=ui.BG)
        fam = ui.family(self)
        bf = (fam, 10)
        if not self.runs:
            self.title('核对')
            tk.Label(self, text='还没有编号照片：先点「开始工作」干一批。', bg=ui.BG, fg=ui.TEXT, font=bf,
                     padx=ui.px(30), pady=ui.px(30)).pack()
            return

        bar = ui.Card(self, pad=10, radius=14)
        bar.pack(fill='x', padx=9, pady=(12, 4))                 # 卡片自带影子的边：看着离窗口边 14、离照片 12
        top = bar.body
        self.run_var = tk.StringVar()
        holder = ui.Combo(top, self.run_var, [run_label(r, len(load_photos(d))) for r, d, _ in self.runs],
                          font=bf, bg=ui.CARD, width=236)
        holder.pack(side='left')
        self.run_box = holder.box
        self.run_box.configure(state='readonly')
        self.run_box.bind('<<ComboboxSelected>>', lambda e: self.load_run(self.run_box.current()))
        ui.Pill(top, radius=8, bg=ui.CARD, height=38, padx=16, text='◀ 上一张', command=lambda: self.step(-1), font=bf).pack(side='left', padx=(12, 6))
        self.num = tk.StringVar()
        box = ui.Field(top, self.num, font=bf, bg=ui.CARD)
        box.configure(width=ui.px(64))
        box.pack(side='left')
        self.entry = box.entry
        self.entry.configure(justify='center')
        self.entry.bind('<Return>', lambda e: self.jump())
        ui.Pill(top, radius=8, bg=ui.CARD, height=38, padx=16, text='看', command=self.jump, primary=True, font=(fam, 10, 'bold')).pack(side='left', padx=6)
        ui.Pill(top, radius=8, bg=ui.CARD, height=38, padx=16, text='下一张 ▶', command=lambda: self.step(1), font=bf).pack(side='left')
        ui.Pill(top, radius=8, bg=ui.CARD, height=38, padx=16, text='↺', command=lambda: self.turn(False), font=bf).pack(side='left', padx=(12, 6))
        ui.Pill(top, radius=8, bg=ui.CARD, height=38, padx=16, text='↻', command=lambda: self.turn(True), font=bf).pack(side='left')
        self.table_btn = ui.Pill(top, radius=8, bg=ui.CARD, height=38, padx=14, text='打开表', icon='table',
                                 command=self.open_table, font=bf)
        self.table_btn.pack(side='left', padx=(10, 0))
        self.info = tk.StringVar()
        tk.Label(top, textvariable=self.info, bg=ui.CARD, fg=ui.GREY, font=(fam, 9)).pack(side='left', padx=12)

        body = tk.Frame(self, bg=ui.BG)
        body.pack(fill='both', expand=True, padx=(14, 9), pady=(0, 6))
        body.columnconfigure(0, weight=1)
        body.rowconfigure(0, weight=1)
        self.img = tk.Label(body, background='#1C1C1E', cursor='hand2', bd=0)
        self.img.grid(row=0, column=0, sticky='nsew', pady=(2, 8))   # 跟右边卡片（自带影子的边）上下对齐
        self.img.bind('<Configure>', lambda e: self._later_render())
        self.img.bind('<Double-Button-1>', lambda e: self.cur in self.photos and _open(self.photos[self.cur]))
        side = ui.Card(body, pad=10, stretch=True)
        side.configure(width=ui.px(370))                         # 看着 360 宽
        side.grid(row=0, column=1, sticky='ns', padx=(7, 0))
        self.panel = tk.Text(side.body, width=40, wrap='word', relief='flat', bd=0, highlightthickness=0,
                             background=ui.CARD, foreground=ui.TEXT, padx=ui.px(6), pady=ui.px(4), font=bf,
                             tabs=(ui.px(130),))
        self.panel.pack(fill='both', expand=True)
        self.panel.tag_configure('h', font=(fam, 12, 'bold'))
        self.panel.tag_configure('k', foreground=ui.GREY)
        self.panel.tag_configure('red', foreground='#9C0006', background='#FFC7CE')
        self.panel.tag_configure('fixed', foreground='#833C0B', background='#FCE4D6')
        self.panel.tag_configure('note', foreground='#9C0006')
        self.panel.tag_configure('dim', foreground='#8E8E93')
        self.panel.tag_configure('human', foreground='#375623', background='#E2EFDA')
        self.panel.tag_configure('ok', foreground='#375623')
        self.panel.configure(cursor='arrow')                     # 格子上变小手：点一下就改
        self.panel.tag_bind('cell', '<Enter>', lambda e: self.panel.configure(cursor='hand2'))
        self.panel.tag_bind('cell', '<Leave>', lambda e: self.panel.configure(cursor='arrow'))
        for key in ('<Left>', '<Prior>'):
            self.bind(key, lambda e: self._key(-1))
        for key in ('<Right>', '<Next>'):
            self.bind(key, lambda e: self._key(1))

        i = 0
        if run_dir:
            for k, (_, d, _) in enumerate(self.runs):
                if os.path.normcase(os.path.abspath(d)) == os.path.normcase(os.path.abspath(run_dir)):
                    i = k
        self.load_run(i, start)
        self.focus_set()

    # ---------- 动作
    def _key(self, d: int):
        if not isinstance(self.focus_get(), tk.Entry):            # 在框里打字时 ←→ 是挪光标，不是翻页
            self.step(d)

    def load_run(self, i: int, start: int | None = None):
        run_id, d, xlsx = self.runs[i]
        self.xlsx = xlsx
        self.run_id, self.run_dir, self.cat, self.today = run_id, d, cat_of(xlsx), run_day(run_id)
        self.editing = None
        self.run_box.current(i)
        self.photos = load_photos(d)
        self.nos = sorted(self.photos)
        self.rows = load_rows(xlsx)
        self.show(start if start in self.photos else self.nos[0])

    def open_table(self):
        """看到哪批开哪批的表，不用去文件夹里找。"""
        if getattr(self, 'xlsx', '') and os.path.isfile(self.xlsx):
            _open(self.xlsx)
        else:
            self.info.set('这批的表找不到了（挪走或者删了？）')

    def step(self, delta: int):
        if not self.nos:
            return
        i = self.nos.index(self.cur) if self.cur in self.photos else 0
        self.show(self.nos[max(0, min(len(self.nos) - 1, i + delta))])

    def jump(self):
        t = self.num.get().strip()
        if not t.isdigit() or int(t) not in self.photos:
            self.info.set(f'没有第 {t} 张（这批是 {self.nos[0]}–{self.nos[-1]}）' if t else '先输入编号')
            return
        self.show(int(t))
        self.focus_set()

    def show(self, no: int):
        self.cur = no
        self.num.set(str(no))
        self.title(f'核对 · 第 {no} 张 / 共 {len(self.nos)} 张')
        self.info.set(f'第 {no} 张 / 共 {len(self.nos)} 张　←→ 翻页，双击看大图，点格子改')
        if self.editing and self.editing[0] != no:
            self.editing = None
        try:
            with Image.open(self.photos[no]) as im:
                self._pil = im.convert('RGB')
        except Exception as e:
            self._pil = None
            self.info.set(f'第 {no} 张打不开：{e}')
        self.render()
        self.fill_panel(no)

    def turn(self, clockwise: bool):
        """这张转得不对：转 90° 存回去，下次打开也是转好的。"""
        if self.cur in self.photos:
            try:
                rotate_photo(self.photos[self.cur], clockwise)
            except Exception as e:
                self.info.set(f'转不了：{e}')
                return
            self.show(self.cur)

    def _later_render(self):
        if self._after:
            self.after_cancel(self._after)
        self._after = self.after(80, self.render)

    def render(self):
        self._after = None
        if self._pil is None:
            self.img.configure(image='')
            return
        w = max(50, self.img.winfo_width() - 4)
        h = max(50, self.img.winfo_height() - 4)
        im = self._pil.copy()
        im.thumbnail((w, h), Image.LANCZOS)
        self._tk = ImageTk.PhotoImage(im)
        self.img.configure(image=self._tk)

    def fill_panel(self, no: int):
        p = self.panel
        y = p.yview()[0]
        p.configure(state='normal')
        p.delete('1.0', 'end')
        if self.editor is not None:
            self.editor.destroy()
            self.editor = None
        recs = self.rows.get(no, [])
        if not recs:
            p.insert('end', f'No. {no}\n\n', 'h')
            p.insert('end', '表里没有这一号的记录。\n', 'dim')
        for k, rec in enumerate(recs):
            p.insert('end', f'No. {no}' + (f'　第 {k + 1} 条' if len(recs) > 1 else '') + '\n', 'h')
            if not any(rec['values']):
                p.insert('end', rec['note'] + '\n\n', 'note')
                continue
            for i, (lab, val) in enumerate(zip(rec.get('labels', LABELS), rec['values'])):
                if (rec.get('builtin') and i == BTMP and not val and i not in rec['red']
                        and i not in rec.get('fixed', ())):
                    continue                            # 换刀调试单上本来就没有 BTMP：空着就不列（表里那一列照留）
                p.insert('end', lab + '\t', 'k')
                look = ('red' if i in rec['red'] else 'fixed' if i in rec.get('fixed', ()) else
                        'human' if i in rec.get('human', ()) else '')
                tag = f'c{k}_{i}'                       # 核对时直接改：点哪格改哪格，表跟着改
                p.insert('end', val or '—', (look, 'cell', tag) if look else ('cell', tag))
                p.insert('end', '\n', look or ())
                p.tag_bind(tag, '<Button-1>', lambda e, k=k, i=i: self.open_cell(k, i))
                if self.editing == (no, k, i):
                    p.window_create('end', window=self._editor(rec, k, i), padx=ui.px(2), pady=ui.px(4))
                    p.insert('end', '\n')
            if rec['note']:
                p.insert('end', '\n' + rec['note'] + '\n', 'ok' if rec['note'] == HUMAN_NOTE else 'note')
            p.insert('end', '\n')
        p.configure(state='disabled')
        p.yview_moveto(y)
        if self.editor is not None:
            p.see(str(self.editor))

    # ---------- 核对里直接改格子
    def _busy_here(self) -> bool:
        try:
            d = self.busy() if self.busy else ''
        except Exception:
            d = ''
        return bool(d) and os.path.normcase(os.path.abspath(d)) == os.path.normcase(os.path.abspath(self.run_dir))

    def open_cell(self, k: int, i: int):
        """点一格：底下展开一个小框，红格子还列出几遍认出来的不同写法，点一个就填上。再点一下收起。"""
        if self._busy_here():
            self.info.set('这批还在干：干完再改（干完存盘会把这本表整本重写）')
            return
        key = (self.cur, k, i)
        self.editing = None if self.editing == key else key
        self.fill_panel(self.cur)

    def close_cell(self):
        if self.editing:
            self.editing = None
            self.fill_panel(self.cur)
        self.focus_set()

    def _editor(self, rec: dict, k: int, i: int):
        p = self.panel
        fam = ui.family(self)
        f = tk.Frame(p, bg=ui.CARD, highlightthickness=1, highlightbackground=ui.LINE, highlightcolor=ui.LINE,
                     padx=ui.px(8), pady=ui.px(8))
        var = tk.StringVar(value=rec['values'][i])
        fld = ui.Field(f, var, font=(fam, 11), bg=ui.CARD)
        fld.configure(width=max(ui.px(220), p.winfo_width() - ui.px(48)))
        fld.pack(fill='x')
        red = i in rec['red']
        alts = rec.get('alts', {}).get(i, []) if red else []
        if alts:
            row = tk.Frame(f, bg=ui.CARD)
            row.pack(fill='x', pady=(ui.px(8), 0))
            for a in alts:
                ui.Pill(row, text=a, radius=8, height=32, padx=12, bg=ui.CARD, font=(fam, 10, 'bold'),
                        command=lambda a=a: self.save_cell(k, i, a)).pack(side='left', padx=(0, ui.px(6)))
        btns = tk.Frame(f, bg=ui.CARD)
        btns.pack(fill='x', pady=(ui.px(8), 0))
        ui.Pill(btns, text='改好了', primary=True, radius=8, height=32, padx=16, bg=ui.CARD, font=(fam, 10, 'bold'),
                command=lambda: self.save_cell(k, i, var.get())).pack(side='right')
        if red or i in rec.get('fixed', ()):
            ui.Pill(btns, text='对的', radius=8, height=32, padx=16, bg=ui.CARD, font=(fam, 10),
                    command=lambda: self.save_cell(k, i, None)).pack(side='right', padx=(0, ui.px(8)))
        ent = fld.entry
        ent.bind('<Return>', lambda e: (self.save_cell(k, i, var.get()), 'break')[1])
        ent.bind('<KP_Enter>', lambda e: (self.save_cell(k, i, var.get()), 'break')[1])
        ent.bind('<Escape>', lambda e: (self.close_cell(), 'break')[1])

        def grab():
            if ent.winfo_exists():
                ent.focus_set()
                ent.select_range(0, 'end')
                ent.icursor('end')
        self.after_idle(grab)
        self.editor, self.editor_var = f, var
        return f

    def save_cell(self, k: int, i: int, text: str | None):
        """「改好了」/ 回车（text 是框里的字）、点候选（text 是那个写法）、「对的」（text 是 None）：写回表里。"""
        if self._busy_here():
            self.info.set('这批还在干：干完再改（干完存盘会把这本表整本重写）')
            return
        no = self.cur
        rec = self.rows[no][k]
        label, before = rec['labels'][i], rec['values'][i]
        try:
            res = edit_cell(self.xlsx, rec['row'], i, text, today=self.today)
        except Locked:
            self.info.set('✗ 表在 Excel 里开着，写不进去：关掉 Excel 再点「改好了」')
            return
        except Exception as e:
            self.info.set(f'✗ 改不了：{type(e).__name__}：{e}')
            return
        if res['still']:
            msg = f'{label} {res["value"] or "空"} 还不对：{res["why"]}'
        elif text is None or res['value'] == before:
            msg = f'{label} 核过了'
        else:
            msg = f'{label} 改成 {res["value"] or "空"}'
        if self.on_checked:
            try:
                if self.on_checked(self.cat, f'核对 {self.run_id} #{no}.{k}', res['data'] if res['checked'] else None):
                    msg += '，这条核完进了样板资料'
            except Exception:
                pass                                        # 样板资料存不下，不耽误改表
        self.rows = load_rows(self.xlsx)
        self.editing = None
        self.fill_panel(no)
        self.info.set(msg + ('' if res['still'] else '，表跟着改了'))
        self.focus_set()
