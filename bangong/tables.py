"""自建类别：刀单以外的单子。

刀单那套（专门写的提示词、刀号/Φ/班次整理、名单、数据库）原样留给刀单，这里一样都不碰。
每个自建类别一个文件夹 `录入僧\\类别\\<名字>\\`，里面是它自己的 `数据库.sqlite` 和 `投喂\\`：
- 还没放过表：认全部内容。模型拿单子上印好的栏目名当列名，出一本表。
- 改对了的表「放入数据库」：表头就定成这个类别的列；每列出现过的值拿来当提示和纠错，
  几行改对过的记录放进提示词当例子；数字列只收数字。越放越准。
两遍对答案、拿不准标红、照片印编号这些通用的，跟刀单一样。
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import time
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import engine
from .kb import FeedResult, cell_text, ckey, lev1, nkey

BUILTIN = '刀单'                    # 内置的那一个：公司的 TOOL BROKEN RECORD
CAT_DIR = '类别'
DB_FILE = '数据库.sqlite'
FEED_DIR = '投喂'
NOTE_HEAD = '核对提示'
MAX_NAME = 20
RESERVED = {'CON', 'PRN', 'AUX', 'NUL', *(f'COM{i}' for i in range(1, 10)), *(f'LPT{i}' for i in range(1, 10))}
BAD_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
DATE_HEAD = re.compile(r'日期|时间|date|time', re.I)
NUM_RE = re.compile(r'-?\d+(?:\.\d+)?')


# ---------------------------------------------------------------- 类别（文件夹）
def root(app_dir: str) -> str:
    return os.path.join(app_dir, CAT_DIR)


def cat_dir(app_dir: str, name: str) -> str:
    return os.path.join(root(app_dir), name)


def names(app_dir: str) -> list[str]:
    """['刀单', 自建的按建的先后]。"""
    out = [BUILTIN]
    r = root(app_dir)
    try:
        subs = [d for d in os.listdir(r) if os.path.isdir(os.path.join(r, d)) and not d.startswith('.')]
    except OSError:
        subs = []
    subs.sort(key=lambda d: (os.path.getctime(os.path.join(r, d)), d))
    return out + [d for d in subs if d != BUILTIN]


def check_name(name: str, app_dir: str) -> str:
    """新类别的名字行不行：行就返回 ''，不行返回原因。"""
    n = (name or '').strip()
    if not n:
        return '名字不能空着'
    if len(n) > MAX_NAME:
        return f'名字太长了（最多 {MAX_NAME} 个字）'
    if BAD_CHARS.search(n) or n.endswith('.'):
        return '名字里不能有 \\ / : * ? " < > | 这些符号'
    if n.upper() in RESERVED:
        return 'Windows 不让用这个名字'
    if n.casefold() in {x.casefold() for x in names(app_dir)}:
        return f'已经有「{n}」了'
    return ''


def create(app_dir: str, name: str) -> str:
    d = cat_dir(app_dir, name.strip())
    os.makedirs(os.path.join(d, FEED_DIR), exist_ok=True)
    TableDB(os.path.join(d, DB_FILE))
    return d


def db_path(app_dir: str, name: str) -> str:
    return os.path.join(cat_dir(app_dir, name), DB_FILE)


def feed_path(app_dir: str, name: str) -> str:
    return os.path.join(cat_dir(app_dir, name), FEED_DIR)


# ---------------------------------------------------------------- 读表
def colkey(name: str) -> str:
    """列名比对用：大小写、空格、标点不算（「M/C No.」和「m/c no」是一列）。"""
    return re.sub(r'[\s.:：。,，、/\\\-_()（）\[\]【】#№]+', '', str(name or '').casefold())


def _headerish(texts: list[str]) -> bool:
    words = [t for t in texts if not re.fullmatch(r'[-\d.,/: ]+', t)]
    return len(texts) >= 2 and len(words) >= max(2, int(len(texts) * 0.6)) and all(len(t) <= 40 for t in texts)


def read_table(rows: list[tuple]) -> tuple[list[tuple[int, str]], list[tuple]] | None:
    """找表头：前 15 行里第一行「至少两格、大多是字、后面还有数据」的。录入僧自己出的表：
    第 1 行是标题（只有一格），第 2 行表头，A 列 No. 是照片编号、最后的「核对提示」不算内容。"""
    for i, row in enumerate(rows[:15]):
        cells = [cell_text(v) for v in row]
        texts = [(j, t) for j, t in enumerate(cells) if t]
        if not _headerish([t for _, t in texts]):
            continue
        if not any(any(cell_text(v) for v in r) for r in rows[i + 1:i + 60]):
            continue
        cols = []
        for j, t in texts:
            if t.startswith(NOTE_HEAD):
                continue
            if j == 0 and t.strip().lower().rstrip('.') in ('no', '序号'):
                continue
            cols.append((j, t))
        if cols:
            return cols, rows[i + 1:]
    return None


def read_book(path: str) -> tuple[list[str], list[tuple[str, dict]], list[str]]:
    """-> (列名, [(页名, {列名: 值})], 读到记录的页名)。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    columns: list[str] = []
    seen: dict[str, str] = {}
    recs, used = [], []
    try:
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            t = read_table(rows)
            if not t:
                continue
            cols, data = t
            for _, nm in cols:
                if colkey(nm) not in seen:
                    seen[colkey(nm)] = nm
                    columns.append(nm)
            heads = {colkey(nm) for _, nm in cols}
            n = 0
            for row in data:
                rec = {seen[colkey(nm)]: (cell_text(row[j]) if j < len(row) else '') for j, nm in cols}
                vals = [v for v in rec.values() if v]
                if not vals or {colkey(v) for v in vals} <= heads:      # 空行、又一行表头
                    continue
                recs.append((ws.title, rec))
                n += 1
            if n:
                used.append(ws.title)
    finally:
        wb.close()
    return columns, recs, used


# ---------------------------------------------------------------- 数据库
SCHEMA = """
CREATE TABLE IF NOT EXISTS columns (pos INTEGER PRIMARY KEY, name TEXT, key TEXT UNIQUE);
CREATE TABLE IF NOT EXISTS records (hash TEXT PRIMARY KEY, source TEXT, sheet TEXT, fed_at TEXT, data TEXT);
CREATE TABLE IF NOT EXISTS sources (path TEXT PRIMARY KEY, mtime REAL, size INTEGER, fed_at TEXT, added INTEGER);
"""


class TableDB:
    """一个自建类别的专属数据库。只认放进来的表（改对了的），录入僧自己认的不进库。"""

    def __init__(self, path: str):
        self.path = path
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with self._conn() as c:
            c.executescript(SCHEMA)

    def _conn(self):
        return sqlite3.connect(self.path, timeout=30)

    def columns(self) -> list[str]:
        with self._conn() as c:
            return [r[0] for r in c.execute('SELECT name FROM columns ORDER BY pos')]

    def feed(self, path: str) -> FeedResult:
        res = FeedResult(path)
        if not path.lower().endswith(('.xlsx', '.xlsm')):
            res.error = '只认 .xlsx / .xlsm（老的 .xls 请先在 Excel 里另存为 .xlsx）'
            return res
        try:
            cols, recs, used = read_book(path)
        except Exception as e:
            res.error = f'读不了：{type(e).__name__}：{e}'
            return res
        if not recs:
            res.error = '没找到表头和记录（表头要在前 15 行，至少两列）'
            return res
        now = datetime.now().isoformat(timespec='seconds')
        with self._conn() as c:
            have = {k: nm for nm, k in c.execute('SELECT name, key FROM columns')}
            pos = c.execute('SELECT COALESCE(MAX(pos), -1) FROM columns').fetchone()[0]
            new = []
            for nm in cols:                         # 第一本的表头定成这个类别的列；以后多出来的列接在后面
                if colkey(nm) not in have:
                    pos += 1
                    c.execute('INSERT INTO columns (pos, name, key) VALUES (?, ?, ?)', (pos, nm, colkey(nm)))
                    have[colkey(nm)] = nm
                    new.append(nm)
            order = [r[0] for r in c.execute('SELECT name FROM columns ORDER BY pos')]
            before = c.execute('SELECT COUNT(*) FROM records').fetchone()[0]
            rows = []
            for sheet, rec in recs:
                data = {have[colkey(k)]: v for k, v in rec.items() if v}
                h = hashlib.sha1('|'.join(nkey(data.get(k, '')) for k in order).encode('utf-8')).hexdigest()
                rows.append((h, os.path.basename(path), sheet, now, json.dumps(data, ensure_ascii=False)))
            c.executemany('INSERT OR IGNORE INTO records (hash, source, sheet, fed_at, data) VALUES (?, ?, ?, ?, ?)', rows)
            res.added = c.execute('SELECT COUNT(*) FROM records').fetchone()[0] - before
            st = os.stat(path)
            c.execute('INSERT OR REPLACE INTO sources (path, mtime, size, fed_at, added) VALUES (?, ?, ?, ?, ?)',
                      (os.path.abspath(path), st.st_mtime, st.st_size, now, res.added))
        res.rows, res.sheets = len(recs), used
        res.new_columns = new
        return res

    def feed_folder(self, folder: str) -> list[FeedResult]:
        if not folder or not os.path.isdir(folder):
            return []
        with self._conn() as c:
            seen = {p: (m, s) for p, m, s in c.execute('SELECT path, mtime, size FROM sources')}
        out = []
        for name in sorted(os.listdir(folder)):
            p = os.path.abspath(os.path.join(folder, name))
            if name.startswith('~$') or not name.lower().endswith(('.xlsx', '.xlsm')) or not os.path.isfile(p):
                continue
            st = os.stat(p)
            if seen.get(p) == (st.st_mtime, st.st_size):
                continue
            out.append(self.feed(p))
        return out

    def count(self) -> int:
        with self._conn() as c:
            return c.execute('SELECT COUNT(*) FROM records').fetchone()[0]

    def summary(self, name: str) -> str:
        cols, n = self.columns(), self.count()
        if not cols:
            return f'「{name}」还没放过表：先认全部内容，改对了点「放入数据库」，表头就定成这个类别的列'
        return f'「{name}」数据库：{n} 条记录 · {len(cols)} 列（{"、".join(cols[:8])}{"…" if len(cols) > 8 else ""}）'

    def load(self) -> 'TableSnapshot':
        with self._conn() as c:
            rows = [json.loads(r[0]) for r in c.execute('SELECT data FROM records ORDER BY rowid')]
        return TableSnapshot(self.columns(), rows)


class TableSnapshot:
    """拿这个类别自己的表当准绳：哪些列是数字、哪些列来回就那几个值（拿来提示和纠错）。"""

    def __init__(self, columns: list[str], records: list[dict]):
        self.columns = list(columns)
        self.records = records
        self.n = len(records)
        vals = {c: [r.get(c, '') for r in records if r.get(c, '')] for c in self.columns}
        self.datey = {c for c in self.columns if DATE_HEAD.search(c)}
        self.numeric = {c for c in self.columns if c not in self.datey and len(vals[c]) >= 3
                        and sum(bool(NUM_RE.fullmatch(v.replace(',', ''))) for v in vals[c]) >= 0.8 * len(vals[c])}
        spell = {c: defaultdict(Counter) for c in self.columns}
        for c in self.columns:
            for v in vals[c]:
                spell[c][nkey(v)][v] += 1
        self.canon = {c: {k: cnt.most_common(1)[0][0] for k, cnt in spell[c].items()} for c in self.columns}
        self.count = {c: Counter() for c in self.columns}
        for c in self.columns:
            for k, cnt in spell[c].items():
                self.count[c][self.canon[c][k]] = sum(cnt.values())
        # 来回就那几个值的列（机台、零件、名字这种）：值会重复出现，又不是数字、日期、长句子
        self.cat = {c for c in self.columns if c not in self.numeric and c not in self.datey and len(vals[c]) >= 3
                    and len(self.canon[c]) <= max(1, int(len(vals[c]) * 0.7))
                    and sum(len(v) for v in vals[c]) / len(vals[c]) <= 40}
        self.by_ckey = {c: defaultdict(set) for c in self.cat}
        for c in self.cat:
            for k, v in self.canon[c].items():
                self.by_ckey[c][ckey(k)].add(v)

    def column_lines(self) -> list[str]:
        out = []
        for i, c in enumerate(self.columns, start=1):
            ex = [v for v, _ in self.count[c].most_common(60 if c in self.cat else 3)]
            tip = '（数字）' if c in self.numeric else ''
            if ex:
                tip += '，例如 ' + '、'.join(ex[:60 if c in self.cat else 3])
            out.append(f'- c{i}：{c}{tip}')
        return out

    def hints(self) -> str:
        if not any(c in self.cat for c in self.columns):
            return ''
        return ('上面「例如」后面的值都是以前录对的：单子上的字跟其中某个很像时，一律以那里的写法为准；'
                '明显不一样的（新的）照单子原样抄。')

    def examples(self, k: int = 3) -> list[dict]:
        good = [r for r in self.records if sum(bool(r.get(c)) for c in self.columns) >= max(1, len(self.columns) // 2)]
        return good[-k:]

    def _snap(self, c: str, v: str) -> str | None:
        cands = self.by_ckey[c].get(ckey(v), set())
        if len(cands) == 1:
            return next(iter(cands))
        k = nkey(v)
        if len(k) >= 4:
            near = {x for kk, x in self.canon[c].items() if lev1(kk, k)}
            if len(near) == 1:
                x = next(iter(near))
                if self.count[c][x] >= 2:
                    return x
        return None

    def fix(self, rec: dict, unsure=()) -> tuple[dict, dict, dict]:
        """跟刀单同一招：认出来的值跟库里某个值只差一个容易看错的字（O/0、I/1、S/5…），就改成库里的写法。"""
        rec = dict(rec)
        fixed: dict = {}
        for c in self.cat:
            v = rec.get(c, '')
            if not v:
                continue
            k = nkey(v)
            if k in self.canon[c]:
                rec[c] = self.canon[c][k]
                continue
            x = self._snap(c, v)
            if x and x != v:
                fixed[c] = (v, x, '按数据库改')
                rec[c] = x
        return rec, fixed, {}


# ---------------------------------------------------------------- 提示词
RULES = """规则：
- 照单子上写的抄，不猜、不补、不纠错；某个字实在看不清就用 □ 代替那个字。
- 单子上没写的就填空字符串。
- 印刷的栏目名不是内容，不要抄进值里。
- 被划掉、涂掉作废的内容不要。
- 图片歪斜或横着拍的，按正确方向读。
- 一条记录都没有就输出空数组。
- title_side：单子的标题（最上面那行印刷的大字）在这张照片的哪一边，照片原样看，不要先在脑子里转正：top（上边）、right（右边）、bottom（下边）、left（左边）。"""

PROMPT_COLS = """你是工厂的录入员，要把「{name}」这种单子上的内容录进表格。从图片里找出每一条记录，每条记录填下面这些列：
{cols}
图片可能是一张单子（通常一条记录），也可能是一页登记表（好几行，每行一条记录）。

{rules}
- unsure：每条记录里你拿不准的列（写列的编号，比如 c2），字迹潦草、被挡住、反光、可能看错的；全都看得清就给空数组。
- 只输出 JSON：{{"title_side": "top", "records": [{{{keys}, "unsure": []}}]}}，不要任何解释。"""

PROMPT_OPEN = """你是工厂的录入员，要把「{name}」这种单子上的内容整理成表格。这种单子还没有定好的表头，由你来定：
- 先看单子上印好的栏目名（比如 Date、M/C No.、Qty、日期、数量），用它们当列名。栏目名中英文都有的只写英文那部分，没有英文就写中文；同一种单子每次用一样的列名，不要自己发明新说法。
- 每条记录一行，把每个栏目里手写或打印的内容填进对应的列。一张单子通常就是一条记录；一页登记表就是每行一条记录。
- 只有印好的说明文字、没有填内容的栏目也要列出来（值填空字符串），这样列是全的。

{rules}
- unsure：每条记录里你拿不准的列名，字迹潦草、被挡住、反光、可能看错的；全都看得清就给空数组。
- 只输出 JSON：{{"title_side": "top", "columns": ["列名1", "列名2"], "records": [{{"values": ["列名1的内容", "列名2的内容"], "unsure": []}}]}}，values 跟 columns 一一对应，不要任何解释。"""

S = {'type': 'STRING'}
SA = {'type': 'ARRAY', 'items': S}


def _gen(model: str, schema: dict | None) -> dict:
    gen = {'responseMimeType': 'application/json'}
    if schema:
        gen['responseSchema'] = schema
    if re.match(r'^gemini-[12]\.', model):
        gen['temperature'] = 0.1
    if re.match(r'^gemini-2\.5-flash(?!-lite)', model):
        gen['thinkingConfig'] = {'thinkingBudget': 0}
    return gen


# ---------------------------------------------------------------- 一个类别怎么认、怎么出表
class TableSpec:
    """交给 Engine 用：刀单以外的类别怎么认（提示词、解析、整理、对答案）和怎么出表。"""

    def __init__(self, name: str, db: str, feed: str = ''):
        self.name, self.db_path, self.feed_dir = name, db, feed
        self.snapshot: TableSnapshot | None = None
        self.columns: list[str] = []
        self.reg: dict[str, str] = {}               # 认全部内容时：列名 → 第一次见到的写法（几张照片列名统一）

    # 开工前：先吃投喂文件夹，再拿库当准绳
    def prepare(self, log):
        try:
            db = TableDB(self.db_path)
            for fr in db.feed_folder(self.feed_dir):
                log(fr.line())
            self.snapshot = db.load()
            self.columns = list(self.snapshot.columns)
            log(db.summary(self.name))
        except Exception as e:
            log(f'「{self.name}」的数据库打不开，这次不用它：{type(e).__name__}：{e}')
            self.snapshot, self.columns = None, []
        self.reg = {colkey(c): c for c in self.columns}

    @property
    def open_mode(self) -> bool:
        return not self.columns

    def prompt(self) -> str:
        if self.open_mode:
            return PROMPT_OPEN.format(name=self.name, rules=RULES)
        snap = self.snapshot
        lines = snap.column_lines() if snap is not None else [f'- c{i}：{c}' for i, c in enumerate(self.columns, 1)]
        keys = ', '.join(f'"c{i}": ""' for i in range(1, len(self.columns) + 1))
        out = PROMPT_COLS.format(name=self.name, cols='\n'.join(lines), rules=RULES, keys=keys)
        if snap is not None:
            if snap.hints():
                out += '\n\n' + snap.hints()
            ex = snap.examples()
            if ex:
                out += '\n\n以前录对的几条（只是让你看格式和写法，不要抄进结果）：\n' + '\n'.join(
                    json.dumps({f'c{i}': r.get(c, '') for i, c in enumerate(self.columns, 1)}, ensure_ascii=False)
                    for r in ex)
        return out

    def schema(self) -> dict:
        if self.open_mode:
            item = {'type': 'OBJECT', 'properties': {'values': SA, 'unsure': SA}, 'required': ['values', 'unsure']}
            return {'type': 'OBJECT', 'properties': {'title_side': S, 'columns': SA,
                                                     'records': {'type': 'ARRAY', 'items': item}},
                    'required': ['columns', 'records', 'title_side']}
        keys = [f'c{i}' for i in range(1, len(self.columns) + 1)]
        item = {'type': 'OBJECT', 'properties': {**{k: S for k in keys}, 'unsure': SA}, 'required': keys + ['unsure']}
        return {'type': 'OBJECT', 'properties': {'title_side': S, 'records': {'type': 'ARRAY', 'items': item}},
                'required': ['records', 'title_side']}

    def body(self, jpeg: bytes, model: str, schema: bool = True) -> dict:
        import base64
        return {
            'systemInstruction': {'parts': [{'text': self.prompt()}]},
            'contents': [{'role': 'user', 'parts': [
                {'text': f'把这张「{self.name}」上的内容逐条录好。'},
                {'inline_data': {'mime_type': 'image/jpeg', 'data': base64.b64encode(jpeg).decode()}},
            ]}],
            'generationConfig': _gen(model, self.schema() if schema else None),
        }

    def _name(self, raw: str) -> str:
        """认全部内容时，几张照片、几遍的列名统一成第一次见到的写法。"""
        raw = re.sub(r'\s+', ' ', str(raw or '')).strip()
        k = colkey(raw)
        if not k:
            return ''
        return self.reg.setdefault(k, raw)

    def parse(self, data: dict) -> tuple[list[dict], str]:
        obj = engine.response_json(data)
        side = engine.norm_side(obj.get('title_side', '')) if isinstance(obj, dict) else ''
        recs = obj.get('records', []) if isinstance(obj, dict) else obj
        if isinstance(recs, dict):
            recs = [recs]
        if not isinstance(recs, list):
            raise ValueError('JSON 里没有 records：' + json.dumps(obj, ensure_ascii=False)[:200])
        out = []
        if self.open_mode:
            cols = [self._name(c) for c in (obj.get('columns') or [])] if isinstance(obj, dict) else []
            for r in recs:
                rec: dict = {}
                if isinstance(r, dict) and isinstance(r.get('values'), list):
                    vals = r['values']
                    unsure = r.get('unsure') or []
                elif isinstance(r, dict):                       # 模型没按格式：{列名: 值}
                    vals, unsure = [], r.get('unsure') or []
                    for k, v in r.items():
                        if k != 'unsure' and self._name(k):
                            rec[self._name(k)] = _txt(v)
                elif isinstance(r, list):
                    vals, unsure = r, []
                else:
                    continue
                for c, v in zip(cols, vals):
                    if c:
                        rec[c] = _txt(v)
                for c in cols[len(vals):]:
                    if c:
                        rec.setdefault(c, '')
                if any(rec.values()):
                    names_ = {colkey(x): x for x in rec}
                    rec['unsure'] = sorted({names_[colkey(u)] for u in _items(unsure) if colkey(u) in names_})
                    out.append(rec)
            return out, side
        idx = {f'c{i}': c for i, c in enumerate(self.columns, 1)}
        by_key = {colkey(c): c for c in self.columns}
        for r in recs:
            rec = {c: '' for c in self.columns}
            unsure = set()
            if isinstance(r, dict):
                for k, v in r.items():
                    if k == 'unsure':
                        for u in _items(v):
                            c = idx.get(str(u).strip().lower()) or by_key.get(colkey(u))
                            if c:
                                unsure.add(c)
                        continue
                    c = idx.get(str(k).strip().lower()) or by_key.get(colkey(k))
                    if c and v is not None:
                        rec[c] = _txt(v)
            elif isinstance(r, list):
                for c, v in zip(self.columns, r):
                    rec[c] = _txt(v)
            if any(rec.values()):
                rec['unsure'] = sorted(unsure)
                out.append(rec)
        return out, side

    def numeric(self) -> set:
        return self.snapshot.numeric if self.snapshot is not None else set()

    def datey(self, c: str) -> bool:
        return bool(DATE_HEAD.search(c))

    def clean(self, rec: dict) -> dict:
        """只整理空白；数字列去掉中间的空格。不碰中文（「一」不当横杠）。"""
        out = {}
        for k, v in rec.items():
            if k in ('unsure', 'disputed', 'counts', 'single'):
                out[k] = v
            else:
                t = re.sub(r'\s+', ' ', str(v or '')).strip()
                out[k] = re.sub(r'\s+', '', t) if k in self.numeric() else t
        return out

    def vote_key(self, c: str, v: str) -> str:
        t = (v or '').strip()
        if self.datey(c):
            d = engine.parse_date(t, date(2000, 1, 1))
            if d:
                return f'{d.day}/{d.month}'
        n = engine.to_int(t)
        if n is not None:
            return str(n)
        return re.sub(r'\s+', '', t).upper()

    def _fields(self, reads: list[list[dict]]) -> list[str]:
        out = list(self.columns)
        for recs in reads:
            for r in recs:
                for k in r:
                    if k not in ('unsure', 'disputed', 'counts', 'single') and k not in out:
                        out.append(k)
        return out

    def agree(self, a: list[dict], b: list[dict]) -> bool:
        fields = self._fields([a, b])
        return len(a) == len(b) and all(self.vote_key(f, x.get(f, '')) == self.vote_key(f, y.get(f, ''))
                                        for x, y in zip(a, b) for f in fields)

    def vote(self, reads: list[list[dict]]) -> list[dict]:
        """跟刀单一样：条数按多数定，再一格一格按多数定；几遍不一样的格子记进 disputed（标红）。"""
        if len(reads) == 1:
            return [dict(r, disputed={}) for r in reads[0]]
        fields = self._fields(reads)
        counts = Counter(len(r) for r in reads)
        target, n = counts.most_common(1)[0]
        if n == 1:
            target = len(reads[0])
        cand = [r for r in reads if len(r) == target]
        out = []
        for i in range(target):
            rec, disputed = {}, {}
            for f in fields:
                vals = [r[i].get(f, '') for r in cand]
                keys = [self.vote_key(f, v) for v in vals]
                c = Counter(keys)
                top = c.most_common(1)[0][1]
                best = keys[0] if c[keys[0]] == top else c.most_common(1)[0][0]
                rec[f] = vals[keys.index(best)]
                if len(c) > 1 or n == 1:
                    disputed[f] = list(dict.fromkeys(v or '空' for v in vals))
            rec['unsure'] = sorted(set().union(*(set(r[i].get('unsure') or ()) for r in cand)))
            rec['disputed'] = disputed
            if len(cand) < len(reads):
                rec['counts'] = [len(r) for r in reads]
            out.append(rec)
        return out

    def book(self, path: str) -> 'TableBook':
        return TableBook(path, self.name, self.columns, self.snapshot, numeric=self.numeric())


def _txt(v) -> str:
    return '' if v is None else re.sub(r'\s+', ' ', str(v)).strip()


def _items(v) -> list[str]:
    if isinstance(v, list):
        return [str(x) for x in v if x]
    return [x for x in re.split(r'[,，、;；\s]+', str(v or '')) if x]


# ---------------------------------------------------------------- 出表
class TableBook:
    """自建类别的表：第 1 行标题（类别名），第 2 行表头（No. + 各列 + 核对提示），第 3 行起一条记录一行。
    认全部内容时列会越认越多：每次存盘整本重写，新列接在右边。"""

    def __init__(self, path: str, title: str, columns: list[str], snapshot: TableSnapshot | None = None,
                 numeric: set | None = None):
        self.path, self.title = path, title
        self.columns = list(columns)
        self.snapshot = snapshot
        self.numeric = set(numeric or ())
        self.rows: list[tuple] = []                 # (编号, {列: 值}, 红的列, 橙的列, 提示, 整行底色)
        self.photos = 0

    def _see(self, rec: dict):
        for k in rec:
            if k not in ('unsure', 'disputed', 'counts', 'single') and k not in self.columns:
                self.columns.append(k)

    def add(self, no: int, records: list[dict] | None, error: str = ''):
        self.photos += 1
        if records is None:
            self.rows.append((no, {}, set(), set(), f'✗ 识别失败：{error}（照片还在待识别里，下次点开始会重干）', 'red'))
            return
        if not records:
            self.rows.append((no, {}, set(), set(), f'这张照片里没认出「{self.title}」的记录', 'red'))
            return
        for rec in records:
            self._record(no, rec)

    def unfinished(self, no: int, why: str):
        self.rows.append((no, {}, set(), set(), f'未完成：{why}（照片还在待识别里，下次点开始会重干）', 'grey'))

    def _record(self, no: int, rec: dict):
        self._see(rec)
        unsure = set(rec.get('unsure') or ())
        disputed = dict(rec.get('disputed') or {})
        red, orange, notes = set(), set(), []
        if self.snapshot is not None and self.snapshot.n:
            rec, fixed, _ = self.snapshot.fix(rec, unsure)
            for c, (old, new, why) in fixed.items():
                orange.add(c)
                notes.append(f'{c} {old}→{new}（{why}）')
        values = {}
        for c in self.columns:
            raw = rec.get(c, '')
            v: object = raw
            never = DATE_HEAD.search(c) is not None        # 日期从不标红（跟刀单一样）
            if c in self.numeric and raw:
                t = raw.replace(',', '')
                if NUM_RE.fullmatch(t):
                    v = int(t) if re.fullmatch(r'-?\d+', t) else float(t)
                else:
                    red.add(c)
                    notes.append(f'{c} 不是数字：{raw}')
            if isinstance(v, str) and '□' in v and not never:
                red.add(c)
                notes.append(f'{c} 有看不清的字')
            if c in disputed and c not in orange and not never:
                red.add(c)
                notes.append(f'{c} 几遍认得不一样：' + ' / '.join(disputed[c]))
            elif c in unsure and c not in red and c not in orange and not never:
                red.add(c)
                notes.append(f'{c} 模型拿不准')
            values[c] = v
        if rec.get('single'):
            notes.append('只认成了一遍，没能对答案')
        if rec.get('counts'):
            notes.append('几遍认出的条数不一样（' + ' / '.join(map(str, rec['counts'])) + ' 条），按多数定')
        self.rows.append((no, values, red, orange, '；'.join(dict.fromkeys(notes)), ''))

    def save(self):
        E = engine
        wb = Workbook()
        ws = wb.active
        ws.title = (re.sub(r'[\\/*?:\[\]]', '', self.title) or 'RECORD')[:31]
        ncol = len(self.columns)
        note_col = 2 + ncol
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(note_col - 1, 2))
        t = E.put(ws, 1, 1, self.title)
        t.font, t.alignment = Font(name='Calibri', size=20, bold=True), E.CENTER
        ws.row_dimensions[1].height = 34
        heads = ['No.'] + self.columns + [f'{NOTE_HEAD}（不用贴进你们的表）']
        for i, h in enumerate(heads, start=1):
            cell = E.put(ws, 2, i, h)
            cell.fill, cell.font, cell.alignment, cell.border = E.HEAD_FILL, E.BOLD, E.CENTER, E.BOX
        ws.cell(row=2, column=note_col).font = Font(name='Calibri', size=10, color='808080')
        for r, (no, values, red, orange, note, whole) in enumerate(self.rows, start=3):
            c = E.put(ws, r, 1, no)
            c.alignment, c.border, c.font = E.CENTER, E.BOX, E.BOLD
            fill = {'red': E.RED_FILL, 'grey': E.GREY_FILL}.get(whole)
            for i, col in enumerate(self.columns, start=2):
                v = values.get(col, '')
                cell = E.put(ws, r, i, v if v != '' else None)
                cell.border, cell.font, cell.alignment = E.BOX, E.BASE_FONT, E.CENTER
                if fill is not None:
                    cell.fill = fill
                if col in orange:
                    cell.fill, cell.font = E.ORANGE_FILL, E.ORANGE_FONT
                if col in red:
                    cell.fill, cell.font = E.RED_FILL, E.RED_FONT
            if note:
                n = E.put(ws, r, note_col, note)
                n.font = E.RED_FONT if (red or whole == 'red') else (E.ORANGE_FONT if orange else E.BASE_FONT)
                n.alignment = E.LEFT
        ws.column_dimensions['A'].width = 6
        for i, col in enumerate(self.columns, start=2):
            longest = max([len(str(col))] + [len(str(row[1].get(col, ''))) for row in self.rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(40, max(8, longest * 1.2 + 2))
        ws.column_dimensions[get_column_letter(note_col)].width = 40
        last = max(2 + len(self.rows), 3)
        ws.freeze_panes = 'B3'
        ws.auto_filter.ref = f'A2:{get_column_letter(max(note_col - 1, 2))}{last}'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = '1:2'
        tmp = self.path + '.tmp'
        wb.save(tmp)
        for _ in range(5):
            try:
                os.replace(tmp, self.path)
                return
            except PermissionError:              # Excel 正开着这个文件
                time.sleep(2)
        alt = self.path.replace('.xlsx', f'_{datetime.now():%H%M%S}.xlsx')
        os.replace(tmp, alt)
        self.path = alt
