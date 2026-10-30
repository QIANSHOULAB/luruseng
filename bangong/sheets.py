"""读表：放入数据库认的几种格式，统一变成 [(页名, [一行一个元组])]。

- .xlsx / .xlsm：openpyxl（只读，公式取算好的值）
- .xls：工厂电脑上常见的老格式，用 xlrd 读，不用再另存为
- .csv：没有页，整本算一页（页名 = 文件名）；编码自己认（UTF-8、Excel 中文版存的 GBK、「Unicode 文本」的
  UTF-16 都行），逗号、分号、Tab 分隔都行

Excel 文件按里面的内容认，不光看后缀：别的系统「导出 Excel」出来的常是假的——其实是 .xlsx 的照读，
其实是老 .xls 的照读，其实是文字表（逗号、Tab 分隔）的按 CSV 读，其实是网页（HTML/XML）的读不了，
告诉人用 Excel 另存为 .xlsx。加了密码的、坏的，也说清楚怎么办。
"""
from __future__ import annotations

import codecs
import csv
import io
import os

EXTS = ('.xlsx', '.xlsm', '.xls', '.csv')
FILETYPES = [('表格（Excel、CSV）', '*.xlsx *.xlsm *.xls *.csv'), ('全部', '*.*')]
UNSUPPORTED = '只认 .xlsx、.xlsm、.xls 和 .csv（照片、PDF、Word 不吃：数据库只认人核对过的表）'
OLE2 = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'      # 老 .xls 的文件头（加了密码的 .xlsx 也是这个头）
NOT_EXCEL = '不是 Excel 文件，打不开：用 Excel 打开看看，能打开就另存为 .xlsx 再放'
WEB = '这个文件其实是网页导出的（HTML/XML），不是真 Excel：用 Excel 打开，另存为 .xlsx 再放'
LOCKED = '这本表加了密码，打不开：用 Excel 打开、去掉密码另存一份再放'
BROKEN = '这本表坏了，打不开：用 Excel 打开看看，能打开就另存为 .xlsx 再放'


class SheetError(ValueError):
    """读不了，但知道为什么、该怎么办：原话给人看，不带错误类型。"""


def why(e: Exception) -> str:
    return str(e) if isinstance(e, SheetError) else f'读不了：{type(e).__name__}：{e}'


def supported(path: str) -> bool:
    return path.lower().endswith(EXTS)


def read_sheets(path: str) -> list[tuple[str, list[tuple]]]:
    ext = os.path.splitext(path)[1].lower()
    if ext not in EXTS:
        raise SheetError(UNSUPPORTED)
    with open(path, 'rb') as f:
        raw = f.read()
    if ext == '.csv':                            # CSV 只当文字读（第一格是 PKG 之类的也不会认错）
        return _table(_decode(raw), path)
    if raw.startswith(b'PK'):                    # .xlsx / .xlsm（后缀写成 .xls 也一样）
        return _xlsx(raw)
    if raw.startswith(OLE2):                     # 老 .xls（后缀写成 .xlsx 也一样）
        return _xls(raw)
    text = _decode(raw, strict=True)             # 都不是：多半是别的系统导出的文字表
    if text is None or not text.strip():
        raise SheetError(NOT_EXCEL)
    if text.lstrip()[:1] == '<':
        raise SheetError(WEB)
    return _table(text, path)


def _xlsx(raw: bytes) -> list[tuple[str, list[tuple]]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return [(ws.title, [tuple(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def _xls(raw: bytes) -> list[tuple[str, list[tuple]]]:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=raw, on_demand=True)
    except Exception as e:
        msg = str(e).lower()
        # 老 .xls 加密码：「Workbook is encrypted」；新 .xlsx 加密码：外面包一层老格式，里面找不到表
        raise SheetError(LOCKED if 'encrypt' in msg or "can't find workbook" in msg else BROKEN) from e
    try:
        out = []
        for i in range(book.nsheets):
            sh = book.sheet_by_index(i)
            rows = [tuple(_xls_value(sh.cell(r, c), book.datemode) for c in range(sh.ncols)) for r in range(sh.nrows)]
            out.append((sh.name, rows))
            book.unload_sheet(i)
        return out
    finally:
        book.release_resources()


def _xls_value(cell, datemode):
    import xlrd
    t = cell.ctype
    if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if t == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        except Exception:
            return cell.value
    if t == xlrd.XL_CELL_NUMBER:
        v = cell.value
        return int(v) if float(v).is_integer() else v
    if t == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return cell.value


def _decode(raw: bytes, strict: bool = False) -> str | None:
    """Excel「Unicode 文本」是带 BOM 的 UTF-16；UTF-8（带不带 BOM）先试；不行就是 Excel 中文版的 GBK；
    再不行按西文。strict：看着不像文字（有 0 字节、哪种编码都不对）就给 None。"""
    if raw[:2] in (codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE):
        try:
            return raw.decode('utf-16')
        except UnicodeDecodeError:
            pass
    if not (strict and b'\x00' in raw[:4096]):
        for enc in ('utf-8-sig', 'gb18030'):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
    return None if strict else raw.decode('cp1252', 'replace')


def _table(text: str, path: str) -> list[tuple[str, list[tuple]]]:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    rows = [tuple(c if c.strip() else None for c in row) for row in csv.reader(io.StringIO(text), dialect)]
    return [(os.path.splitext(os.path.basename(path))[0], rows)]


def touch_ok(name: str) -> bool:
    """投喂文件夹里哪些文件该吃：认得的格式，跳过 Excel 开着时留的 ~$ 临时文件。"""
    return not name.startswith('~$') and supported(name)
